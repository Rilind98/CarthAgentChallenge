"""Score filled Arbeitsversion workbooks against ground-truth workbooks, cell by cell.

    .venv/bin/python bench/score_av.py <truth dir: Arbeitsversion_<id>.xlsx> <run dir: <id>/Arbeitsversion_<id>.xlsx> [--pdf-only]

Counted: every amount cell that holds a number in the truth (a formula means the client printed no value there).
Rows: Anlagenverzeichnis by Anlagen-Nr. (in order for repeats); Gitter sheets by an order-preserving alignment that
maximises equal amounts, within the same sheet.
Per truth cell: correct (equal to the cent) · blank (ours empty / formula / row missing) · incorrect (other number).
Also incorrect: our numbers where the matched truth row has none, and our numbers in rows the truth does not have.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

SHEETS = {"Anlagenverzeichnis": (2, [c for c in "W X Y Z AA AB AC AD AE AF AG AH AI AJ AK".split()]),
          "Gitter nach Anlagenklasse": (8, list("DEFGHIJKLMNOPQRSTU")),
          "Gitter nach Bilanzposition": (8, list("DEFGHIJKLMNOPQRSTU"))}


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def rows_of(ws, start, cols, key_col):
    out = []
    for r in range(start, ws.max_row + 1):
        vals = {c: _num(ws[f"{c}{r}"].value) for c in cols}
        if not any(v is not None for v in vals.values()):
            continue
        out.append({"row": r, "key": ws[f"{key_col}{r}"].value, "vals": vals})
    return out


def _same(a, b):
    return a is not None and b is not None and abs(a - b) < 0.005


def align(truth, ours):
    """Order-preserving alignment maximising equal amounts (DP)."""
    n, m = len(truth), len(ours)
    sim = [[sum(_same(t["vals"][c], o["vals"][c]) for c in t["vals"]) for o in ours] for t in truth]
    dp = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            dp[i][j] = max(dp[i + 1][j], dp[i][j + 1], (sim[i][j] + dp[i + 1][j + 1]) if sim[i][j] else -1)
    pairs, i, j = {}, 0, 0
    while i < n and j < m:
        if sim[i][j] and dp[i][j] == sim[i][j] + dp[i + 1][j + 1]:
            pairs[i] = j
            i, j = i + 1, j + 1
        elif dp[i + 1][j] >= dp[i][j + 1]:
            i += 1
        else:
            j += 1
    return pairs


def score_doc(truth_path: Path, ours_path: Path | None) -> Counter:
    c = Counter()
    tw = openpyxl.load_workbook(truth_path)
    ow = openpyxl.load_workbook(ours_path) if ours_path and ours_path.exists() else None
    for sheet, (start, cols) in SHEETS.items():
        truth = rows_of(tw[sheet], start, cols, "A" if sheet == "Anlagenverzeichnis" else "C")
        ours = rows_of(ow[sheet], start, cols, "A" if sheet == "Anlagenverzeichnis" else "C") if ow else []
        if sheet == "Anlagenverzeichnis":
            pool = {}
            for j, o in enumerate(ours):
                pool.setdefault(str(o["key"]), []).append(j)
            pairs = {}
            for i, t in enumerate(truth):
                js = pool.get(str(t["key"]))
                if js:
                    pairs[i] = js.pop(0)
        else:
            pairs = align(truth, ours)
        for i, t in enumerate(truth):
            o = ours[pairs[i]] if i in pairs else None
            for col, tv in t["vals"].items():
                ov = o["vals"][col] if o else None
                if tv is None:
                    if ov is not None:
                        c["incorrect_extra"] += 1
                    continue
                c["truth_values"] += 1
                if ov is None:
                    c["blank"] += 1
                elif _same(tv, ov):
                    c["correct"] += 1
                else:
                    c["incorrect"] += 1
        matched = set(pairs.values())
        for j, o in enumerate(ours):
            if j not in matched:
                c["incorrect_spurious"] += sum(v is not None for v in o["vals"].values())
    return c


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("truth", type=Path)
    ap.add_argument("run", type=Path)
    ap.add_argument("--pdf-only", action="store_true")
    ap.add_argument("--corpus", type=Path, default=Path(__file__).resolve().parents[2] / "A3-Hackathon-Corpus")
    a = ap.parse_args()
    total, per_doc = Counter(), {}
    for tp in sorted(a.truth.glob("Arbeitsversion_*.xlsx")):
        doc = tp.stem.removeprefix("Arbeitsversion_")
        fam = "_".join(doc.split("_")[:-1])
        src = next((a.corpus / fam).glob(f"{doc}.*[!n]"), None) if (a.corpus / fam).exists() else None
        if a.pdf_only and src is not None and src.suffix.lower() != ".pdf":
            continue
        c = score_doc(tp, a.run / doc / f"Arbeitsversion_{doc}.xlsx")
        per_doc[doc] = dict(c)
        total += c
    wrong = total["incorrect"] + total["incorrect_extra"] + total["incorrect_spurious"]
    print(json.dumps({"docs": len(per_doc), "truth_values": total["truth_values"], "correct": total["correct"],
                      "blank": total["blank"], "incorrect": wrong,
                      "incorrect_breakdown": {"different_number": total["incorrect"],
                                              "number_where_truth_has_none": total["incorrect_extra"],
                                              "row_not_in_truth": total["incorrect_spurious"]},
                      "score_+1_0_-3": total["correct"] - 3 * wrong}, indent=1))
    (a.run / "score_av.json").write_text(json.dumps(per_doc, indent=1))


if __name__ == "__main__":
    main()
