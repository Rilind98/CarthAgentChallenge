"""Step 0 - perception: every printed string with its exact position. No interpretation yet.

PDF with a text layer -> words from pdfplumber; rotated text is mapped into an upright reading frame.
xlsx -> one token per non-empty cell (cached formula values), with hidden/merged information.
PDF without a text layer (scan, glyph outlines) -> no tokens; the OCR engine plugs in here.
"""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.utils import get_column_letter

from .numbers import detect_style, fmt_de, split_glued, token_kind
from .schema import Token, Tokens


def _frame(b, direction, W, H):
    x0, top, x1, bottom = b
    if direction == "btt":  # text runs bottom to top (page turned 90 deg counter-clockwise)
        return [H - bottom, x0, H - top, x1]
    if direction == "ttb":
        return [top, W - x1, bottom, W - x0]
    return [x0, top, x1, bottom]


def _page_from_frame(f, direction, W, H):
    fx0, ftop, fx1, fbot = f
    if direction == "btt":
        return [ftop, H - fx1, fbot, H - fx0]
    if direction == "ttb":
        return [W - fbot, fx0, W - ftop, fx1]
    return [fx0, ftop, fx1, fbot]


def words_to_tokens(page_words: dict[int, list[dict]], pages: list[dict], source: str) -> tuple[list[Token], str]:
    """Shared by the text layer and OCR. page_words: page -> words {text, x0, top, x1, bottom[, conf]} in page
    points. The number style (1.234,56 or 1,234.56) is voted once per document, then used for every word."""
    style = detect_style(w["text"] for ws in page_words.values() for w in ws)
    info = {p["page"]: p for p in pages}
    toks = []
    for pno, words in sorted(page_words.items()):
        W, H, direction = info[pno]["width"], info[pno]["height"], info[pno]["direction"]
        for i, w in enumerate(words):
            bbox = [round(w["x0"], 2), round(w["top"], 2), round(w["x1"], 2), round(w["bottom"], 2)]
            frame = [round(v, 2) for v in _frame(bbox, direction, W, H)]
            tid, conf = f"p{pno}w{i}", w.get("conf", 1.0)
            parts = split_glued(w["text"], style)
            if parts:  # two amounts printed without a gap: split by character share
                n0, n = len(parts[0]), len(w["text"])
                cut = frame[0] + (frame[2] - frame[0]) * n0 / n
                for suffix, text, fr in (("a", parts[0], [frame[0], frame[1], cut, frame[3]]),
                                         ("b", parts[1], [cut, frame[1], frame[2], frame[3]])):
                    toks.append(Token(id=tid + suffix, text=text, kind="amount", source=source, page=pno, conf=conf,
                                      frame=[round(v, 2) for v in fr],
                                      bbox=[round(v, 2) for v in _page_from_frame(fr, direction, W, H)],
                                      glued_from=w["text"]))
                continue
            toks.append(Token(id=tid, text=w["text"], kind=token_kind(w["text"], style=style), source=source,
                              page=pno, bbox=bbox, frame=frame, conf=conf))
    return toks, style


def pdf_tokens(path: Path, doc_id: str) -> Tokens:
    pages, page_words = [], {}
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            W, H = float(page.width), float(page.height)
            rotated = [c for c in page.chars if not c.get("upright", True)]
            direction = "ltr"
            if len(rotated) > len(page.chars) / 2:
                direction = "btt" if sum(c["matrix"][1] > 0 for c in rotated) >= len(rotated) / 2 else "ttb"
            pages.append({"page": pno, "width": W, "height": H, "direction": direction, "chars": len(page.chars)})
            if page.chars:  # stream order keeps overlapping strings apart; rotated text needs the geometric order
                page_words[pno] = page.extract_words(keep_blank_chars=False, use_text_flow=direction == "ltr",
                                                     x_tolerance=1.5, char_dir_rotated="btt" if direction == "btt" else "ttb")
    toks, style = words_to_tokens(page_words, pages, "pdf-text")
    if not toks or os.environ.get("OCR_ALWAYS") == "1":  # no text layer (scan, glyph outlines): OCR it
        from .ocr import ocr_pdf
        pages, page_words = ocr_pdf(path)
        toks, style = words_to_tokens(page_words, pages, "ocr:tesseract")
        return Tokens(doc_id=doc_id, format="ocr-pdf" if toks else "image-pdf", pages=pages, tokens=toks,
                      number_style=style, source_path=str(path))
    return Tokens(doc_id=doc_id, format="pdf", pages=pages, tokens=toks, number_style=style)


def xlsx_tokens(path: Path, doc_id: str) -> Tokens:
    wb_f = openpyxl.load_workbook(path)  # formulas
    wb_v = openpyxl.load_workbook(path, data_only=True)  # cached values
    toks, sheets = [], []
    style = detect_style(c.value for ws in wb_v.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    for ws in wb_v.worksheets:
        if ws.sheet_state != "visible":
            continue
        sheets.append(ws.title)
        wf = wb_f[ws.title]
        hidden_cols = set()
        for dim in ws.column_dimensions.values():
            if dim.hidden and dim.min and dim.max:
                hidden_cols.update(range(dim.min, dim.max + 1))
        merged = {}
        for rng in ws.merged_cells.ranges:
            merged[(rng.min_row, rng.min_col)] = [rng.min_col, rng.max_col]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                f = wf.cell(c.row, c.column).value
                is_formula = isinstance(f, str) and f.startswith("=")
                if v is None and is_formula:
                    kind, text, num = "uncached", str(f), None
                elif v is None or (isinstance(v, str) and not v.strip()):
                    continue
                elif isinstance(v, (datetime, date)):
                    d = v.date() if isinstance(v, datetime) else v
                    kind, text, num = "date", d.strftime("%d.%m.%Y"), None
                elif isinstance(v, bool):
                    kind, text, num = "text", str(v), None
                elif isinstance(v, (int, float)):
                    cents = round(v * 100)
                    kind, text, num = "number", (str(int(v)) if float(v).is_integer() and abs(v) < 10000 and "0.00" not in c.number_format else fmt_de(cents)), float(v)
                else:
                    text = str(v).strip()
                    kind, num = token_kind(text, style=style), None
                toks.append(Token(id=f"{ws.title}!{get_column_letter(c.column)}{c.row}", text=text, kind=kind,
                                  source="xlsx", sheet=ws.title, row=c.row, col=c.column, number=num,
                                  merged_cols=merged.get((c.row, c.column)),
                                  hidden=c.column in hidden_cols or bool(ws.row_dimensions[c.row].hidden),
                                  hidden_col=c.column in hidden_cols))
    return Tokens(doc_id=doc_id, format="xlsx", sheets=sheets, tokens=toks, number_style=style)


def read_tokens(path: Path, doc_id: str) -> Tokens:
    return xlsx_tokens(path, doc_id) if path.suffix.lower() in (".xlsx", ".xlsm") else pdf_tokens(path, doc_id)
