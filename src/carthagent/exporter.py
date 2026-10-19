from __future__ import annotations

import re

from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
import json
import os
import fcntl
from pathlib import Path
import time
from typing import Any

from openpyxl import Workbook, load_workbook

from .models import Abstention, CoverageItem, OutputBinding, OutputIdentity, SourceObservation, ValidatedFact, jsonable
from .structure import ContextResolver


SCHEMA_VERSION = "carthagent.provisional.evidence.v1"


def canonical_digest(value: Any) -> str:
    return sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@dataclass(frozen=True, slots=True)
class Snapshot:
    id: str
    run_id: str
    document_id: str
    source_hash: str
    policy_digest: str
    facts: tuple[ValidatedFact, ...]
    abstentions: tuple[Abstention, ...]
    coverage: tuple[CoverageItem, ...]
    observations: tuple[dict[str, Any], ...]
    output_bindings: tuple[OutputBinding, ...]

    def core(self) -> dict[str, Any]:
        return {"run_id": self.run_id, "document_id": self.document_id, "source_hash": self.source_hash, "policy_digest": self.policy_digest, "facts": [fact_record(f) for f in self.facts], "abstentions": [jsonable(a) for a in self.abstentions], "coverage": [jsonable(c) for c in self.coverage], "observations": list(self.observations), "output_bindings": [jsonable(b) for b in self.output_bindings]}

    @property
    def digest(self) -> str:
        return canonical_digest(self.core())


def fact_record(fact: ValidatedFact) -> dict[str, Any]:
    return {
        "id": fact.id,
        "document_id": fact.document_id,
        "status": "validated",
        "field": fact.business.field,
        "asset_identity": {"entity": fact.business.entity, "book": fact.business.book, "asset_or_class": fact.business.asset_or_class, "source_name": fact.business.source_name, "asset_id": fact.business.asset_id, "asset_sub_id": fact.business.asset_sub_id, "class_code": fact.business.class_code, "class_name": fact.business.class_name, "account": fact.business.account, "balance_position": fact.business.balance_position, "granularity": fact.business.granularity, "row_kind": fact.business.row_kind, "period": fact.business.period, "temporal_role": fact.business.temporal_role},
        "amount": {"raw": fact.amount.raw_lexeme, "display_magnitude": str(fact.amount.displayed_value), "source_signed": str(fact.amount.source_signed_value), "normalized_base": str(fact.amount.normalized_value), "scale": str(fact.amount.scale), "precision": fact.amount.precision, "currency": fact.business.currency, "unit": fact.business.unit},
        "pointer": {"source_hash": fact.physical.source_hash, "occurrence": fact.physical.occurrence, **dict(fact.source_pointer)},
        "source_anomaly": fact.source_anomaly,
        "dependency_ids": list(fact.dependency_ids),
        "policy_versions": dict(fact.policy_versions),
    }


def evidence_payload(snapshot: Snapshot, run_status: str, diagnostics: tuple[str, ...]) -> dict[str, Any]:
    core = snapshot.core()
    return {"schema_version": SCHEMA_VERSION, "provisional": True, "contract_status": "contract_pending", "run_id": snapshot.run_id, "document": {"document_id": snapshot.document_id, "source_hash": snapshot.source_hash, "status": run_status}, "snapshot": {"id": snapshot.id, "digest": snapshot.digest, "policy_digest": snapshot.policy_digest}, "facts": core["facts"], "abstentions": core["abstentions"], "coverage": core["coverage"], "observations": core["observations"], "output_bindings": core["output_bindings"], "diagnostics": list(diagnostics)}


def _write_review(path: Path, snapshot: Snapshot) -> None:
    wb = Workbook()
    ws = wb.active; ws.title = "Provisional Review"
    headers = ["status", "fact_id", "asset_or_class", "field", "source_signed", "normalized_base", "raw", "source_pointer", "anomaly"]
    ws.append(headers)
    for fact in snapshot.facts:
        row = fact_record(fact)
        ws.append(["validated", safe_cell(fact.id), safe_cell(fact.business.asset_or_class), safe_cell(fact.business.field), str(fact.amount.source_signed_value), str(fact.amount.normalized_value), safe_cell(fact.amount.raw_lexeme), safe_cell(fact.physical.occurrence), safe_cell(fact.source_anomaly)])
    for abstention in snapshot.abstentions:
        ws.append(["abstained", abstention.id, None, None, None, None, None, safe_cell(",".join(abstention.candidate_observation_ids)), safe_cell(abstention.reason)])
    meta = wb.create_sheet("_carthagent_manifest")
    meta.sheet_state = "hidden"
    meta.append(["schema_version", SCHEMA_VERSION]); meta.append(["snapshot_id", snapshot.id]); meta.append(["snapshot_digest", snapshot.digest]); meta.append(["run_id", snapshot.run_id]); meta.append(["source_hash", snapshot.source_hash]); meta.append(["fact_count", len(snapshot.facts)]); meta.append(["abstention_count", len(snapshot.abstentions)])
    wb.save(path)


def _readback(evidence: Path, workbook: Path, snapshot: Snapshot, template_used: bool = False) -> None:
    snapshot_digest = snapshot.digest
    payload = json.loads(evidence.read_text())
    if payload["snapshot"]["digest"] != snapshot_digest or len(payload["facts"]) != len(snapshot.facts):
        raise RuntimeError("evidence readback mismatch")
    observations = {o["id"]: o for o in payload["observations"]}
    for fact in payload["facts"]:
        pointer = fact["pointer"]
        observation = observations.get(pointer.get("observation_id"))
        if observation is None or observation["locator"] != pointer["occurrence"]:
            raise RuntimeError("pointer/observation readback mismatch")
        if pointer["kind"] == "xlsx" and observation["locator"] != f"{pointer['sheet']}!{pointer['cell']}":
            raise RuntimeError("xlsx direct pointer mismatch")
        metadata = observation.get("metadata") or {}
        # A verified Azure cell points at its printed lexeme (the cell polygon stays in its metadata).
        verified_lexeme = metadata.get("lexeme_bbox") if observation.get("kind") == "cell" and observation.get("channel") == "pdf_ocr" and metadata.get("verified") is True else None
        if pointer["kind"] == "pdf" and list(verified_lexeme or observation.get("bbox") or []) != pointer["bbox"]:
            raise RuntimeError("pdf direct pointer mismatch")
    wb = load_workbook(workbook, read_only=not template_used, data_only=False)
    if template_used:
        facts = {f.id: f for f in snapshot.facts}
        for binding in snapshot.output_bindings:
            if binding.status != "bound" or not binding.destination.cell:
                continue
            sheet, cell = binding.destination.cell.rsplit("!", 1)
            expected = TemplateExportAdapter.target_value(facts[binding.record_id], binding)
            if Decimal(str(wb[sheet][cell].value)) != expected:
                raise RuntimeError("template binding readback mismatch")
        issues = [cell.value for row in wb["Offene Punkte"].iter_rows() for cell in row]
        if not any(snapshot_digest in str(value) for value in issues):
            raise RuntimeError("template snapshot audit missing")
    else:
        meta = {row[0].value: row[1].value for row in wb["_carthagent_manifest"].iter_rows()}
        if meta.get("snapshot_digest") != snapshot_digest or int(meta.get("fact_count", -1)) != len(snapshot.facts):
            raise RuntimeError("workbook readback mismatch")


def publish(snapshot: Snapshot, output_root: Path, run_status: str, diagnostics: tuple[str, ...], *, template: Path | None = None, fail_after_stage: bool = False) -> tuple[Path, Path, Path, Path]:
    output_root.mkdir(parents=True, exist_ok=True)
    lock_path = output_root / f".{snapshot.document_id}.publish.lock"
    with lock_path.open("a+b") as lock:
        fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
        return _publish_locked(snapshot, output_root, run_status, diagnostics, template=template, fail_after_stage=fail_after_stage)


def _publish_locked(snapshot: Snapshot, output_root: Path, run_status: str, diagnostics: tuple[str, ...], *, template: Path | None = None, fail_after_stage: bool = False) -> tuple[Path, Path, Path, Path]:
    # document_id is already content-namespaced by the pipeline, so it is both
    # collision-safe and the stable artifact directory name.
    final = output_root / snapshot.document_id
    if final.exists() and (final / ".complete").exists():
        try:
            existing_manifest = json.loads((final / "manifest.v1.json").read_text())
            hashes_valid = all((final / name).is_file() and sha256((final / name).read_bytes()).hexdigest() == digest for name, digest in existing_manifest.get("files", {}).items())
            existing_evidence = json.loads((final / "evidence.v1.json").read_text())
            identity_valid = (
                existing_manifest.get("run_id") == snapshot.run_id
                and existing_manifest.get("snapshot_digest") == snapshot.digest
                and existing_manifest.get("status") == run_status
                and existing_manifest.get("document_id") == snapshot.document_id
                and existing_manifest.get("source_hash") == snapshot.source_hash
                and existing_evidence.get("document", {}).get("status") == run_status
            )
            if identity_valid and hashes_valid:
                _readback(final / "evidence.v1.json", final / "review.v1.xlsx", snapshot, bool(template))
                return final, final / "evidence.v1.json", final / "review.v1.xlsx", final / "manifest.v1.json"
        except (OSError, KeyError, ValueError, RuntimeError, json.JSONDecodeError):
            pass
        os.replace(final, output_root / f"{final.name}.superseded-{time.time_ns()}")
    if final.exists():
        os.replace(final, output_root / f"{final.name}.incomplete-{time.time_ns()}")
    stage = output_root / f".{final.name}.staging-{time.time_ns()}"
    stage.mkdir()
    evidence = stage / "evidence.v1.json"; review = stage / "review.v1.xlsx"; manifest = stage / "manifest.v1.json"
    evidence.write_text(json.dumps(evidence_payload(snapshot, run_status, diagnostics), sort_keys=True, indent=2, ensure_ascii=False) + "\n")
    if template:
        TemplateExportAdapter().export_snapshot(template, review, snapshot)
    else:
        _write_review(review, snapshot)
    _readback(evidence, review, snapshot, bool(template))
    file_hashes = {p.name: sha256(p.read_bytes()).hexdigest() for p in (evidence, review)}
    manifest.write_text(json.dumps({"schema_version": SCHEMA_VERSION, "provisional": True, "official_contract": "contract_pending", "run_id": snapshot.run_id, "document_id": snapshot.document_id, "source_hash": snapshot.source_hash, "snapshot_id": snapshot.id, "snapshot_digest": snapshot.digest, "status": run_status, "files": file_hashes, "readback_verified": True}, sort_keys=True, indent=2) + "\n")
    if fail_after_stage:
        raise RuntimeError("simulated interruption after staging")
    (stage / ".complete").write_text(snapshot.digest + "\n")
    os.replace(stage, final)
    return final, final / evidence.name, final / review.name, final / manifest.name


class TemplateExportAdapter:
    """Fail-closed adapter for the supplied Arbeitsversion layout 20."""

    SUPPORTED_HASH = "0449e028926cbe6e32c32fabc6da27a2693c36147c581c8d3b008e9543239668"
    CONTRACT_VERSION = "A3_LAYOUT_VERSION=20"
    IMPLEMENTATION_VERSION = "3"
    REGISTER = "Anlagenverzeichnis"
    CLASS_GRID = "Gitter nach Anlagenklasse"
    BALANCE_GRID = "Gitter nach Bilanzposition"
    REGISTER_COLUMNS = {"ahk_open": "W", "zugang": "X", "abgang": "Y", "umb": "Z", "ahk_close": "AA", "kum_afa_open": "AB", "afa_plan": "AC", "extraordinary": "AD", "afa_abgang": "AE", "zuschreibung": "AF", "afa_umb": "AG", "kum_afa_close": "AH", "bw_close": "AI", "bw_prev": "AJ", "fk_zinsen": "AK"}
    GRID_COLUMNS = {"ahk_open": "D", "zugang": "E", "abgang": "F", "umb": "G", "net_ahk": "H", "ahk_close": "I", "kum_afa_open": "J", "afa_plan": "K", "total_afa_gj": "L", "extraordinary": "M", "afa_abgang": "N", "zuschreibung": "O", "afa_umb": "P", "net_afa": "Q", "kum_afa_close": "R", "bw_close": "S", "bw_prev": "T", "fk_zinsen": "U"}
    POSITIVE_TARGET_FIELDS = {"abgang", "kum_afa_open", "afa_plan", "total_afa_gj", "extraordinary", "afa_abgang", "kum_afa_close"}

    def verify(self, template: Path) -> None:
        if sha256(template.read_bytes()).hexdigest() != self.SUPPORTED_HASH:
            raise ValueError("unsupported template hash")
        wb = load_workbook(template, read_only=False, data_only=False)
        required = {self.REGISTER, self.CLASS_GRID, self.BALANCE_GRID, "Offene Punkte"}
        if len(wb.sheetnames) != 27 or not required <= set(wb.sheetnames):
            raise ValueError("template layout mismatch")

    @staticmethod
    def _projection(fact: ValidatedFact) -> tuple[str | None, str | None, str | None, int, int]:
        axes = (fact.business.entity or "", fact.business.book or "", fact.business.period or "")
        if fact.business.granularity == "asset" and fact.business.asset_id:
            row_key = "|".join((*axes, fact.business.asset_id, fact.business.asset_sub_id or ""))
            return TemplateExportAdapter.REGISTER, row_key, TemplateExportAdapter.REGISTER_COLUMNS.get(fact.business.field or ""), 2, 500
        if fact.business.granularity == "class" and (fact.business.class_code or fact.business.class_name):
            row_key = "|".join((*axes, fact.business.class_code or fact.business.class_name or ""))
            return TemplateExportAdapter.CLASS_GRID, row_key, TemplateExportAdapter.GRID_COLUMNS.get(fact.business.field or ""), 8, 20
        if fact.business.balance_position:
            row_key = "|".join((*axes, fact.business.balance_position))
            return TemplateExportAdapter.BALANCE_GRID, row_key, TemplateExportAdapter.GRID_COLUMNS.get(fact.business.field or ""), 8, 20
        return None, None, None, 0, 0

    def plan_bindings(self, snapshot_id: str, facts: tuple[ValidatedFact, ...]) -> tuple[OutputBinding, ...]:
        row_maps: dict[str, dict[str, int]] = {}
        bindings: list[OutputBinding] = []
        row_sources: dict[tuple[str, str], set[str]] = {}
        row_fields: dict[tuple[str, str], dict[str, list[str | None]]] = {}
        row_identities: dict[tuple[str, str], set[tuple[str | None, str | None, str | None]]] = {}
        for fact in facts:
            projected_sheet, projected_key, _, _, _ = self._projection(fact)
            # The source scope is the sheet, or the page of a PDF fact (item 6: a PDF occurrence carries no "!", so without the
            # page every field of one asset counted as its own source and no PDF fact could ever bind).
            page = fact.source_pointer.get("page_index")
            source_sheet = str(fact.source_pointer.get("sheet") or (f"page:{page}" if page is not None else fact.physical.occurrence.split("!", 1)[0]))
            if projected_sheet and projected_key:
                row_sources.setdefault((projected_sheet, projected_key), set()).add(source_sheet)
                row_fields.setdefault((projected_sheet, projected_key), {}).setdefault(source_sheet, []).append(fact.business.field)
                row_identities.setdefault((projected_sheet, projected_key), set()).add((fact.business.asset_id, fact.business.asset_sub_id, fact.business.class_code))

        def split_width(key: tuple[str, str]) -> bool:
            """R7-6: one record printed across consecutive PDF pages (AHK columns on one page, AfA/book values on the next):
            the same exact identity, field sets disjoint between the pages. Any shared field keeps the row ambiguous."""
            sources = row_fields.get(key, {})
            pages = sorted(int(name.split(":", 1)[1]) for name in sources if re.fullmatch(r"page:\d+", name))
            if len(pages) != len(sources) or pages != list(range(pages[0], pages[0] + len(pages))) or len(row_identities.get(key, ())) != 1:
                return False
            fields = [field for values in sources.values() for field in values]
            return None not in fields and len(fields) == len(set(fields))
        for index, fact in enumerate(facts):
            sheet, row_key, column, first_row, capacity = self._projection(fact)
            status = "bound"
            row = None
            if not sheet or not row_key or not column:
                status = "contract_pending"
            elif fact.business.entity is None and len(row_sources.get((sheet, row_key), ())) > 1 and not split_width((sheet, row_key)):
                status = "ambiguous_destination_identity"
            else:
                mapping = row_maps.setdefault(sheet, {})
                if row_key not in mapping:
                    mapping[row_key] = first_row + len(mapping)
                row = mapping[row_key]
                if row >= first_row + capacity:
                    status = "projection_overflow"
            transform = "positive_magnitude" if fact.business.field in self.POSITIVE_TARGET_FIELDS else "source_signed"
            destination = OutputIdentity(self.CONTRACT_VERSION, row_key, fact.business.field, f"{sheet}!{column}{row}" if status == "bound" else None)
            bindings.append(OutputBinding(f"{fact.document_id}:binding:{index}", snapshot_id, fact.id, destination, fact.source_pointer, status, {"polarity": transform, "source_unit": fact.business.unit, "target_unit": "EUR"}))
        # K1 defence in depth: two facts planned for one destination cell both refuse (the later would overwrite the earlier).
        from collections import Counter
        cells = Counter(binding.destination.cell for binding in bindings if binding.status == "bound" and binding.destination.cell)
        # R8-2: a row without a company code may be any company's row of the same remaining key: when such a row and another
        # row of that key are both planned, every binding of the key refuses (blank entity is a wildcard for collision).
        wildcard_keys: dict[tuple[str, str], set[str]] = {}
        for binding in bindings:
            if binding.status == "bound" and binding.destination.cell and binding.destination.row_key:
                sheet = binding.destination.cell.rsplit("!", 1)[0]
                wildcard_keys.setdefault((sheet, binding.destination.row_key.split("|", 1)[-1]), set()).add(binding.destination.row_key)
        def wildcard_collision(binding: OutputBinding) -> bool:
            if not (binding.destination.cell and binding.destination.row_key):
                return False
            keys = wildcard_keys.get((binding.destination.cell.rsplit("!", 1)[0], binding.destination.row_key.split("|", 1)[-1]), set())
            return len(keys) > 1 and any(key.split("|", 1)[0] == "" for key in keys)
        return tuple(OutputBinding(binding.id, binding.snapshot_id, binding.record_id, OutputIdentity(binding.destination.contract_version, binding.destination.row_key, binding.destination.field, None),
                                   binding.source_pointer, "destination_collision", binding.transformations)
                     if binding.status == "bound" and (cells[binding.destination.cell] > 1 or wildcard_collision(binding)) else binding for binding in bindings)

    @staticmethod
    def target_value(fact: ValidatedFact, binding: OutputBinding) -> Decimal:
        value = fact.amount.normalized_value
        return abs(value) if binding.transformations.get("polarity") == "positive_magnitude" else value

    def export_snapshot(self, template: Path, destination: Path, snapshot: Snapshot) -> None:
        self.verify(template)
        wb = load_workbook(template, data_only=False)
        facts = {f.id: f for f in snapshot.facts}
        written_identity: set[tuple[str, int]] = set()
        issues: list[tuple[str, str, str, str]] = []
        for binding in snapshot.output_bindings:
            fact = facts[binding.record_id]
            if binding.status != "bound" or not binding.destination.cell:
                issues.append((binding.destination.row_key or "Unmapped", fact.business.field or "unknown", f"Source fact {fact.id} has no safe template destination", binding.status))
                continue
            sheet_name, cell = binding.destination.cell.rsplit("!", 1); ws = wb[sheet_name]; row = ws[cell].row
            ws[cell] = self.target_value(fact, binding)
            key = (sheet_name, row)
            if key in written_identity: continue
            written_identity.add(key)
            if sheet_name == self.REGISTER:
                ws.cell(row, 1, safe_cell(fact.business.asset_id)); ws.cell(row, 2, safe_cell(fact.business.source_name)); ws.cell(row, 3, safe_cell(fact.business.class_name)); ws.cell(row, 4, safe_cell(fact.business.class_code)); ws.cell(row, 7, safe_cell(fact.business.balance_position)); ws.cell(row, 8, safe_cell(fact.business.account))
            else:
                ws.cell(row, 2, safe_cell(fact.business.class_code or fact.business.balance_position)); ws.cell(row, 3, safe_cell(fact.business.class_name or fact.business.source_name or fact.business.balance_position)); ws.cell(row, 22, safe_cell(fact.business.row_kind))
                if sheet_name == self.CLASS_GRID: ws.cell(row, 23, safe_cell(fact.business.balance_position))
        self._clear_unavailable_printed_fields(wb, snapshot, facts, issues)
        bound_cells = {binding.destination.cell for binding in snapshot.output_bindings if binding.status == "bound" and binding.destination.cell}
        register_rows = {wb[self.REGISTER][binding.destination.cell.rsplit("!", 1)[1]].row for binding in snapshot.output_bindings if binding.status == "bound" and binding.destination.cell and binding.destination.cell.startswith(self.REGISTER + "!")}
        for row in register_rows:
            for column in ("AA", "AH", "AI", "AJ"):
                destination_cell = f"{self.REGISTER}!{column}{row}"
                if destination_cell not in bound_cells and wb[self.REGISTER][f"{column}{row}"].data_type == "f":
                    wb[self.REGISTER][f"{column}{row}"] = None
                    issues.append((self.REGISTER, f"{column}{row}", "Template-derived monetary value withheld because no source or explicit derived binding exists", "source_only_export"))
        issues.append(("CarthAgent", "JSON/scorer", f"Snapshot {snapshot.digest}; official JSON schema and scorer remain unavailable", "contract_pending"))
        issue_ws = wb["Offene Punkte"]
        for row_no, issue in enumerate(issues, 4):
            for col, value in enumerate(issue, 1): issue_ws.cell(row_no, col, safe_cell(value))
        wb.save(destination)

    def _clear_unavailable_printed_fields(self, wb, snapshot: Snapshot, facts: dict[str, ValidatedFact], issues: list[tuple[str, str, str, str]]) -> None:
        observations = tuple(SourceObservation(item["id"], item["document_id"], item["source_hash"], item["kind"], item["locator"], item["channel"], item.get("text"), tuple(item["bbox"]) if item.get("bbox") else None, item.get("metadata", {})) for item in snapshot.observations)
        resolver = ContextResolver()
        bound_rows: dict[tuple[object, ...], int] = {}
        bound_cells = {binding.destination.cell for binding in snapshot.output_bindings if binding.status == "bound" and binding.destination.cell}

        def identity_key(sheet: str, business) -> tuple[object, ...]:
            return (
                sheet, business.entity, business.book, business.period,
                business.asset_id, business.asset_sub_id, business.class_code,
                business.balance_position, business.granularity,
            )

        for binding in snapshot.output_bindings:
            if binding.status != "bound" or not binding.destination.cell:
                continue
            fact = facts[binding.record_id]
            sheet, cell = binding.destination.cell.rsplit("!", 1)
            row = wb[sheet][cell].row
            bound_rows[identity_key(sheet, fact.business)] = row
        for observation in observations:
            if observation.channel != "xlsx_native" or observation.kind != "cell" or not observation.metadata.get("formula") or observation.metadata.get("cached_value") is not None:
                continue
            context = resolver.resolve(observation, observations)
            business = context.business
            if business is None or not business.field:
                continue
            if business.asset_id:
                sheet = self.REGISTER; column = self.REGISTER_COLUMNS.get(business.field)
            elif business.class_code:
                sheet = self.CLASS_GRID; column = self.GRID_COLUMNS.get(business.field)
            elif business.balance_position:
                sheet = self.BALANCE_GRID; column = self.GRID_COLUMNS.get(business.field)
            else:
                continue
            row = bound_rows.get(identity_key(sheet, business))
            destination_cell = f"{sheet}!{column}{row}" if row is not None and column else None
            if destination_cell and destination_cell not in bound_cells:
                wb[sheet][f"{column}{row}"] = None
                issues.append((sheet, f"{column}{row}", f"Printed source field {observation.locator} is unavailable", "source_unreadable_cleared"))
