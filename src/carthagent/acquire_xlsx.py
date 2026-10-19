from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import time

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .models import CoverageDisposition, CoverageItem, SourceObservation
from .numeric import NumericParseError, parse_amount
from .policies import ResourceBounds


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def normalize_ooxml_member(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    if target.startswith("xl/"):
        return target
    return "xl/" + target.lstrip("/")


class ResourceLimitError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class AcquisitionResult:
    observations: tuple[SourceObservation, ...]
    coverage: tuple[CoverageItem, ...]
    diagnostics: tuple[str, ...]
    partial: bool = False


def _preflight_archive(path: Path, bounds: ResourceBounds) -> None:
    with ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > bounds.max_cells:
            raise ResourceLimitError(f"archive_member_limit:{len(infos)}>{bounds.max_cells}")
        total = 0
        for info in infos:
            if info.file_size > bounds.max_input_bytes:
                raise ResourceLimitError(f"archive_member_uncompressed_limit:{info.filename}:{info.file_size}>{bounds.max_input_bytes}")
            total += info.file_size
            if total > bounds.max_input_bytes:
                raise ResourceLimitError(f"archive_uncompressed_limit:{total}>{bounds.max_input_bytes}")
            # The configured unpacked-byte ceiling also bounds compression-ratio
            # amplification without introducing an unapproved second numeric limit.
            if info.compress_size == 0 and info.file_size:
                raise ResourceLimitError(f"archive_invalid_compression:{info.filename}")


def _xml_cells(path: Path, bounds: ResourceBounds) -> dict[str, dict[str, str | None]]:
    _preflight_archive(path, bounds)
    result: dict[str, dict[str, str | None]] = {}
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {n.attrib["Id"]: n.attrib["Target"] for n in rels}
        sheets = workbook.find("m:sheets", NS)
        for sheet in [] if sheets is None else sheets:
            name = sheet.attrib["name"]
            target = rel_map[sheet.attrib[f"{{{NS['r']}}}id"]]
            member = normalize_ooxml_member(target)
            root = ET.fromstring(archive.read(member))
            for cell in root.findall(".//m:c", NS):
                if len(result) >= bounds.max_cells:
                    raise ResourceLimitError(f"cell_limit:{len(result) + 1}>{bounds.max_cells}")
                ref = cell.attrib["r"]
                formula = cell.findtext("m:f", default=None, namespaces=NS)
                cached = cell.findtext("m:v", default=None, namespaces=NS)
                result[f"{name}!{ref}"] = {"formula_xml": formula, "cached_xml": cached, "xml_type": cell.attrib.get("t"), "member": member}
    return result


def _amount_text(text: str) -> bool:
    try:
        parse_amount(text); return True
    except NumericParseError:
        return False


def _record_blocks(rows: list[tuple[int, str, str, str]], document_id: str, source_hash: str, sheet: str) -> list[SourceObservation]:
    """An identity-text row followed by consecutive numeric-only rows; blank, text or mixed rows end the block. Evidence only."""
    blocks: list[SourceObservation] = []
    for index, (number, row_id, shape, text) in enumerate(rows):
        if shape != "identity":
            continue
        members = [(number, row_id, text)]
        cursor = index + 1
        while cursor < len(rows) and rows[cursor][2] == "numeric" and rows[cursor][0] == members[-1][0] + 1:
            members.append((rows[cursor][0], rows[cursor][1], rows[cursor][3])); cursor += 1
        if len(members) < 2:
            continue
        marker = bool(re.match(r"\s*(?:\*+(?:\s|\||$)|(?:summe|gesamt|total)\b)", text, re.I))
        blocks.append(SourceObservation(f"{document_id}:xlsx:{sheet}:block:{number}", document_id, source_hash, "record_block", f"{sheet}!block:{number}", "xlsx_structure", " || ".join(item[2] for item in members), metadata={
            "member_ids": [item[1] for item in members], "identity_row_id": row_id, "measure_row_ids": [item[1] for item in members[1:]],
            "sheet": sheet, "row": number, "subtotal_marker": marker, "structure_role": "record_block"}))
    return blocks


def acquire_xlsx(path: Path, document_id: str, source_hash: str, bounds: ResourceBounds, *, deadline: float | None = None) -> AcquisitionResult:
    if deadline is not None and time.monotonic() >= deadline:
        raise ResourceLimitError("document_timeout:before_xlsx_acquisition")
    xml = _xml_cells(path, bounds)
    with path.open("rb") as stream:
        wb_formula = load_workbook(stream, data_only=False, read_only=False, keep_links=True)
    if deadline is not None and time.monotonic() >= deadline:
        raise ResourceLimitError("document_timeout:xlsx_formula_load")
    with path.open("rb") as stream:
        wb_cached = load_workbook(stream, data_only=True, read_only=False, keep_links=True)
    if len(wb_formula.worksheets) > bounds.max_worksheets:
        raise ResourceLimitError(f"worksheet_limit:{len(wb_formula.worksheets)}>{bounds.max_worksheets}")
    observations: list[SourceObservation] = []
    coverage: list[CoverageItem] = [
        CoverageItem(
            f"{document_id}:coverage:sheet:{ws.title}", document_id, "worksheet", ws.title,
            CoverageDisposition.UNRESOLVED, reason="unprocessed_document_timeout",
        )
        for ws in wb_formula.worksheets
    ]
    worksheet_coverage_index = {ws.title: index for index, ws in enumerate(wb_formula.worksheets)}
    diagnostics: list[str] = []
    count = 0
    for ws in wb_formula.worksheets:
        if deadline is not None and time.monotonic() >= deadline:
            diagnostics.append(f"document_timeout:worksheet:{ws.title}")
            return AcquisitionResult(tuple(observations), tuple(coverage), tuple(diagnostics), partial=True)
        cached_ws = wb_cached[ws.title]
        sheet_meta = {"visibility": ws.sheet_state, "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden], "hidden_columns": [i for i, d in ws.column_dimensions.items() if d.hidden], "auto_filter": str(ws.auto_filter.ref or "")}
        merged_by_cell: dict[str, tuple[str, bool]] = {}
        for merged in ws.merged_cells.ranges:
            merge_area = (merged.max_row - merged.min_row + 1) * (merged.max_col - merged.min_col + 1)
            if merge_area > bounds.max_cells:
                raise ResourceLimitError(f"merge_area_limit:{merged}:{merge_area}>{bounds.max_cells}")
            anchor = ws.cell(merged.min_row, merged.min_col).coordinate
            for row in ws.iter_rows(min_row=merged.min_row, max_row=merged.max_row, min_col=merged.min_col, max_col=merged.max_col):
                for cell in row:
                    merged_by_cell[cell.coordinate] = (str(merged), cell.coordinate == anchor)
        min_col, min_row, max_col, max_row = range_boundaries(ws.calculate_dimension())
        # OOXML dimensions are anchored to worksheet coordinates; a lone styled
        # cell at XFD1048576 must not disguise a 17-billion-cell declared extent.
        declared_area = max_row * max_col
        if declared_area > bounds.max_cells:
            raise ResourceLimitError(f"declared_range_limit:{ws.title}:{declared_area}>{bounds.max_cells}")
        sparse_rows: dict[int, list] = {}
        sheet_row_ids: list[str] = []
        row_shapes: list[tuple[int, str, str, str]] = []
        for cell in ws._cells.values():
            sparse_rows.setdefault(cell.row, []).append(cell)
        for row_number in sorted(sparse_rows):
            if deadline is not None and time.monotonic() >= deadline:
                diagnostics.append(f"document_timeout:worksheet:{ws.title}:row:{row_number}")
                return AcquisitionResult(tuple(observations), tuple(coverage), tuple(diagnostics), partial=True)
            existing = sparse_rows[row_number]
            # Once a row and bounded used region are established, materialize
            # its absent coordinates as explicit empty observations.  This is
            # the positive evidence needed for blank-means-zero; it remains
            # distinguishable from omitted rows, unreadable cells, formulas
            # with missing caches, and merged followers.
            if row_number > min_row and any(c.value is not None for c in existing):
                row = [ws.cell(row_number, column) for column in range(min_col, max_col + 1)]
            else:
                row = sorted(existing, key=lambda c: c.column)
            row_member_ids: list[str] = []
            row_texts: list[str] = []
            row_numbers = row_words = 0
            for cell in row:
                raw = cell.value
                cache = cached_ws[cell.coordinate].value
                explicit_empty = raw is None and cell.coordinate not in merged_by_cell and row_number > min_row
                if raw is None and not explicit_empty and cell.coordinate not in merged_by_cell:
                    continue
                count += 1
                if count > bounds.max_cells:
                    diagnostics.append(f"cell_limit:{count}>{bounds.max_cells}")
                    return AcquisitionResult(tuple(observations), tuple(coverage), tuple(diagnostics), partial=True)
                locator = f"{ws.title}!{cell.coordinate}"
                formula = raw if cell.data_type == "f" else None
                xml_meta = xml.get(locator, {})
                merge = merged_by_cell.get(cell.coordinate)
                metadata = {**sheet_meta, "coordinate": cell.coordinate, "raw_value": raw, "formula": formula, "cached_value": cache, "display_value": None if cache is None and formula else str(cache if formula else raw), "number_format": cell.number_format, "style_id": cell.style_id, "data_type": cell.data_type, "merged_range": merge[0] if merge else None, "merged_anchor": merge[1] if merge else None, "confirmed_empty": explicit_empty, "full_region_present": explicit_empty, **xml_meta}
                observation_text = cache if formula and cache is not None else raw
                observation_id = f"{document_id}:xlsx:{locator}"
                observations.append(SourceObservation(observation_id, document_id, source_hash, "empty_cell" if explicit_empty else "cell", locator, "xlsx_native", None if observation_text is None else str(observation_text), metadata=metadata))
                row_member_ids.append(observation_id)
                if observation_text is not None:
                    row_texts.append(str(observation_text))
                    if isinstance(observation_text, (int, float)) and not isinstance(observation_text, bool):
                        row_numbers += 1
                    elif not (isinstance(observation_text, str) and _amount_text(observation_text)):
                        row_words += 1
                coverage.append(CoverageItem(f"{document_id}:coverage:cell:{locator}", document_id, "cell", locator, CoverageDisposition.UNRESOLVED, reason="awaiting_semantic_disposition"))
            if row_member_ids:
                row_number = row[0].row
                row_id = f"{document_id}:xlsx:{ws.title}:row:{row_number}"
                observations.append(SourceObservation(row_id, document_id, source_hash, "logical_row", f"{ws.title}!row:{row_number}", "xlsx_structure", " | ".join(row_texts), metadata={"member_ids": row_member_ids, "sheet": ws.title, "row": row_number}))
                sheet_row_ids.append(row_id)
                shape = "identity" if row_words and not row_numbers else ("numeric" if not row_words and row_texts else "other")
                row_shapes.append((row_number, row_id, shape, " | ".join(row_texts)))
        record_blocks = _record_blocks(row_shapes, document_id, source_hash, ws.title)
        observations.extend(record_blocks)
        observations.append(SourceObservation(f"{document_id}:xlsx:{ws.title}:region", document_id, source_hash, "table_region", f"{ws.title}!region", "xlsx_structure", ws.title, metadata={"member_ids": sheet_row_ids, "sheet": ws.title, "structure_role": "worksheet_region"}))
        coverage[worksheet_coverage_index[ws.title]] = CoverageItem(
            f"{document_id}:coverage:sheet:{ws.title}", document_id, "worksheet", ws.title,
            CoverageDisposition.UNRESOLVED if ws.sheet_state != "visible" else CoverageDisposition.OUT_OF_SCOPE,
            reason="hidden_sheet_scope_unknown" if ws.sheet_state != "visible" else "container_inventory",
        )
        archive_members: list[str] = []
        with ZipFile(path) as archive:
            archive_members = [n for n in archive.namelist() if n.startswith(("xl/drawings/", "xl/media/", "xl/embeddings/")) and not n.endswith("/")]
        for member in archive_members:
            coverage.append(CoverageItem(f"{document_id}:coverage:object:{member}", document_id, "container_object", member, CoverageDisposition.UNRESOLVED, reason="embedded_or_drawing_scope_unknown"))
    return AcquisitionResult(tuple(observations), tuple(coverage), tuple(diagnostics), partial=any(c.disposition == CoverageDisposition.UNRESOLVED and c.kind in {"worksheet", "container_object"} for c in coverage))

