from __future__ import annotations

from dataclasses import FrozenInstanceError, asdict, replace
from datetime import datetime
from decimal import Decimal
import json
from pathlib import Path
from hashlib import sha256
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
import carthagent.pipeline as pipeline_module
import carthagent.acquire_xlsx as acquire_xlsx_module

from carthagent.acquire_xlsx import ResourceLimitError, acquire_xlsx, normalize_ooxml_member
from carthagent.acquire_pdf import _mistral_recognition_candidates
from carthagent.exporter import Snapshot, TemplateExportAdapter, publish, safe_cell
from carthagent.geometry import PageTransform
from carthagent.models import AssociationClaim, AssociationKind, BusinessIdentity, CoverageDisposition, CoverageItem, HypothesisSet, OutputBinding, OutputIdentity, PhysicalIdentity, SourceObservation, ValidatedFact, frozen_map, jsonable
from carthagent.numeric import NumericParseError, parse_amount
from carthagent.pipeline import RunConfig, _observation_chunks, run_input
from carthagent.policies import DEFAULT_BOUNDS, DEFAULT_POLICIES, EmptyZeroPolicy, ExcelValuePolicy, PointerPolicy, SignPolarityPolicy
from carthagent.semantics import AssociationEngine, interpret, prove_empty_zero
from carthagent.store import Ledger


def observation(doc="d", suffix="A1", text="100", **kwargs):
    return SourceObservation(f"{doc}:o:{suffix}", doc, "a" * 64, kwargs.pop("kind", "cell"), f"S!{suffix}", kwargs.pop("channel", "xlsx_native"), text, metadata=kwargs.pop("metadata", {}), **kwargs)


def test_recursive_immutability_and_serialization():
    obs = observation(metadata={"nested": {"items": [1, {"x": 2}]}})
    with pytest.raises(TypeError): obs.metadata["nested"]["x"] = 3
    with pytest.raises((AttributeError, TypeError)): obs.metadata["nested"]["items"].append(4)
    assert jsonable(obs)["metadata"]["nested"]["items"][1]["x"] == 2
    with pytest.raises(FrozenInstanceError): obs.text = "changed"


def test_document_local_id_and_open_world_hypothesis():
    with pytest.raises(ValueError): observation(doc="d", suffix="x").__class__("foreign:o", "d", "h", "cell", "S!A1", "xlsx_native")
    with pytest.raises(ValueError): HypothesisSet("h", "d", "d:o", (AssociationClaim(AssociationKind.OTHER, ("d:r",)),))


@pytest.mark.parametrize(("raw", "expected", "notation"), [("1.234,56", "1234.56", "leading_plus_or_unsigned"), ("184.782,83-", "-184782.83", "trailing_minus"), ("(12.50)", "-12.50", "parentheses"), ("12 CR", "-12", "cr")])
def test_numeric_replay_sign_scale(raw, expected, notation):
    parsed = parse_amount(raw, scale=Decimal(1000) if raw == "12 CR" else Decimal(1), cr_dr=True)
    source_expected = Decimal(expected)
    assert parsed.source_signed_value == source_expected
    assert parsed.sign_notation == notation
    assert parse_amount(parsed.raw_lexeme, scale=parsed.scale, cr_dr=True).normalized_value == parsed.normalized_value


def test_numeric_rejects_dirty_and_sign_policy_is_explicit():
    with pytest.raises(NumericParseError): parse_amount("12x")
    policy = SignPolarityPolicy("SignPolarityPolicy", "1")
    assert policy.target(-12, None) is None
    assert policy.target(-12, "positive_magnitude") == 12


def test_association_positive_and_ambiguous_controls():
    amount = observation()
    row = SourceObservation("d:row:1", "d", "a" * 64, "logical_row", "S!row:1", "xlsx_structure", "Asset A | 100", metadata={"member_ids": [amount.id]})
    unknown = AssociationClaim(AssociationKind.UNKNOWN_UNMODELED, ())
    engine = AssociationEngine()
    positive = HypothesisSet("d:h1", "d", amount.id, (AssociationClaim(AssociationKind.OTHER, (row.id,), (row.id,), (), True), unknown))
    assert engine.resolve(positive).target_observation_ids == (row.id,)
    ambiguous = HypothesisSet("d:h2", "d", amount.id, (AssociationClaim(AssociationKind.PREVIOUS, ("d:x",), ("d:x",), (), True), AssociationClaim(AssociationKind.FOLLOWING, ("d:y",), ("d:y",), (), True), unknown))
    assert engine.resolve(ambiguous) is None


def test_empty_zero_requires_positive_evidence():
    assert prove_empty_zero(row_established=True, amount_column_established=True, full_region_present=True, merged_follower=False, unreadable_or_omitted=False)
    assert not EmptyZeroPolicy("EmptyZeroPolicy", "1").accepts(row_established=True, amount_column_established=True, full_region_present=True, merged_follower=True, unreadable_or_omitted=False)


def test_pdf_coordinate_roundtrip():
    transform = PageTransform(crop_x=11, crop_y=7, scale=4.166)
    box = (10.25, 20.5, 30.75, 40.0)
    assert transform.roundtrip_ok(box)
    assert PointerPolicy("PointerPolicy", "1").validate({"kind": "pdf", "bbox": box, "observation_id": "d:o"})
    for rotation in (90, 180, 270):
        assert PageTransform(rotation=rotation, scale=3, page_width=600, page_height=800).roundtrip_ok(box)


def test_dependency_invalidation_is_per_axis(tmp_path):
    ledger = Ledger(tmp_path / "l.db"); ledger.depend("fact:a", "header:unit", "unit"); ledger.depend("fact:b", "header:unit", "period")
    assert ledger.invalidate("header:unit", "unit") == {"fact:a"}
    ledger.close()


def test_policy_version_changes_cache_key():
    a = Ledger.cache_key("h", "r", {}, {"PointerPolicy": "1"}); b = Ledger.cache_key("h", "r", {}, {"PointerPolicy": "2"})
    assert a != b


def test_content_cache_roundtrip(tmp_path):
    ledger = Ledger(tmp_path / "cache.db"); ledger.cache_put("key", {"v": [1, 2]})
    assert ledger.cache_get("key") == {"v": [1, 2]}
    ledger.close()


def test_xlsx_preserves_formula_cache_merge_style_hidden(tmp_path):
    path = tmp_path / "book.xlsx"; wb = Workbook(); wb.active.title = "Visible"; ws = wb.create_sheet("Data"); ws.sheet_state = "hidden"; ws["A1"] = "Asset ID"; ws["B1"] = "Amount EUR"; ws["A2"] = "A-1"; ws["B2"] = "=1+1"; ws["B2"].number_format = "#,##0.00"; ws.merge_cells("C2:D2"); ws["C2"] = "merged"; wb.save(path)
    result = acquire_xlsx(path, "d", "a" * 64, DEFAULT_BOUNDS); values = {o.locator: o for o in result.observations if o.kind == "cell"}
    assert values["Data!B2"].metadata["formula"] == "=1+1" and values["Data!B2"].metadata["cached_value"] is None
    assert values["Data!C2"].metadata["merged_anchor"] is True
    assert any(c.reason == "hidden_sheet_scope_unknown" for c in result.coverage)
    assert ExcelValuePolicy("ExcelValuePolicy", "1").select(raw="=1+1", formula="=1+1", cached=None, display=None)[0] == "formula_result_unavailable"


def test_xlsx_datetime_serializes_and_evidenced_blank_becomes_zero(tmp_path):
    path = tmp_path / "blank.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["Asset ID", "Opening cost EUR 2025", "Additions EUR 2025", "As of"]); ws.append(["A-1", None, 1, datetime(2025, 12, 31)])
    wb.save(path)
    acquired = acquire_xlsx(path, "d", "a" * 64, DEFAULT_BOUNDS)
    blank = next(o for o in acquired.observations if o.locator == "Data!B2")
    assert blank.kind == "empty_cell" and blank.metadata["confirmed_empty"] is True
    assert jsonable(next(o for o in acquired.observations if o.locator == "Data!D2"))["metadata"]["raw_value"].startswith("2025-12-31")
    _, facts, abstentions, coverage = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    zero = next(f for f in facts if f.source_observation_id == blank.id)
    assert zero.amount.source_signed_value == 0 and zero.amount.raw_lexeme == "" and zero.amount.sign_notation == "implicit_zero"
    assert not [a for a in abstentions if blank.id in a.candidate_observation_ids]
    assert next(c for c in coverage if c.locator == "Data!B2").disposition == CoverageDisposition.FACT


def test_buchwert_header_is_amount_not_book_axis(tmp_path):
    path = tmp_path / "book-value.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["Anlagen-Nr", "Buchwert 31.12.2025 EUR"]); ws.append(["A-1", 800]); wb.save(path)
    acquired = acquire_xlsx(path, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, abstentions, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    assert not abstentions and len(facts) == 1
    assert facts[0].business.field == "bw_close" and facts[0].business.book is None


def test_ooxml_relationship_targets_absolute_and_relative():
    assert normalize_ooxml_member("/xl/worksheets/sheet1.xml") == "xl/worksheets/sheet1.xml"
    assert normalize_ooxml_member("worksheets/sheet1.xml") == "xl/worksheets/sheet1.xml"
    assert normalize_ooxml_member("xl/worksheets/sheet1.xml") == "xl/worksheets/sheet1.xml"


def test_giant_declared_range_fails_fast(tmp_path):
    path = tmp_path / "giant.xlsx"; wb = Workbook(); wb.active["XFD1048576"] = "styled edge"; wb.save(path)
    with pytest.raises(ResourceLimitError, match="declared_range_limit"):
        acquire_xlsx(path, "d", "a" * 64, DEFAULT_BOUNDS)


def make_fact():
    amount = parse_amount("1.234,56")
    return ValidatedFact("d:f", "d", PhysicalIdentity("a" * 64, "d", "S!B2"), BusinessIdentity(asset_or_class="A", asset_id="A", field="amount", currency="EUR", unit="EUR", granularity="asset", row_kind="detail"), amount, "d:o:B2", {"kind": "xlsx", "sheet": "S", "cell": "B2", "observation_id": "d:o:B2"}, ("d:o:B2",), {"PointerPolicy": "1"})


def test_no_hypotheses_in_matched_snapshot_and_interrupted_recovery(tmp_path):
    fact = make_fact(); obs = observation(suffix="B2", text="1.234,56"); obs = SourceObservation("d:o:B2", "d", "a" * 64, "cell", "S!B2", "xlsx_native", "1.234,56")
    cov = CoverageItem("d:c", "d", "cell", "S!B2", CoverageDisposition.FACT, fact.id)
    binding = OutputBinding("d:b", "snap", fact.id, OutputIdentity("pending"), fact.source_pointer)
    snap0 = Snapshot("snap", "run", "d", "a" * 64, "p", (fact,), (), (cov,), (jsonable(obs),), (binding,))
    with pytest.raises(RuntimeError): publish(snap0, tmp_path, "partial", (), fail_after_stage=True)
    final, evidence, review, manifest = publish(snap0, tmp_path, "partial", ())
    payload = json.loads(evidence.read_text()); assert "hypotheses" not in payload and payload["snapshot"]["digest"] == snap0.digest
    assert load_workbook(review, read_only=True)["_carthagent_manifest"]["B3"].value == snap0.digest
    assert json.loads(manifest.read_text())["readback_verified"] is True and (final / ".complete").exists()


def test_formula_injection_safe_and_unsupported_corrupt_limits(tmp_path):
    assert safe_cell("=cmd") == "'=cmd"
    doc = tmp_path / "x.docx"; doc.write_bytes(b"x")
    assert run_input(doc, tmp_path / "out").status.value == "unsupported_input"
    corrupt = tmp_path / "x.pdf"; corrupt.write_bytes(b"bad")
    assert run_input(corrupt, tmp_path / "out").status.value == "failed"
    huge = tmp_path / "x.xlsx"; huge.write_bytes(b"123")
    tiny = RunConfig(bounds=DEFAULT_BOUNDS.__class__(**{**asdict(DEFAULT_BOUNDS), "max_input_bytes": 1}))
    assert run_input(huge, tmp_path / "out", tiny).diagnostics[0].startswith("ResourceLimitError:input_size_limit")


def test_template_adapter_preserves_layout_and_audits_bindings(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    template = tmp_path / "template.xlsx"; wb = Workbook(); wb.active.title = "Anlagenverzeichnis"
    for name in ["Gitter nach Anlagenklasse", "Gitter nach Bilanzposition", "Offene Punkte"] + [f"Other {i}" for i in range(23)]: wb.create_sheet(name)
    wb["Gitter nach Anlagenklasse"].merge_cells("D5:I5"); wb["Gitter nach Anlagenklasse"]["Z1"] = "=1+1"; wb["Offene Punkte"]["A4"] = "Keine offenen Punkte"; wb.save(template)
    monkeypatch.setattr(TemplateExportAdapter, "SUPPORTED_HASH", sha256(template.read_bytes()).hexdigest())
    base = make_fact(); business = replace(base.business, field="ahk_open", source_name="Asset A")
    fact = replace(base, business=business); adapter = TemplateExportAdapter(); bindings = adapter.plan_bindings("snap", (fact,))
    snapshot = Snapshot("snap", "run", "d", "a" * 64, "p", (fact,), (), (), (), bindings)
    output = tmp_path / "out.xlsx"; before = sha256(template.read_bytes()).hexdigest(); adapter.export_snapshot(template, output, snapshot)
    assert sha256(template.read_bytes()).hexdigest() == before
    check = load_workbook(output, data_only=False)
    assert len(check.sheetnames) == 27 and str(check["Gitter nach Anlagenklasse"]["Z1"].value).startswith("=")
    assert str(check["Anlagenverzeichnis"]["W2"].value) == "1234.56" and check["Anlagenverzeichnis"]["A2"].value == "A"
    assert snapshot.digest in check["Offene Punkte"]["C4"].value
    assert not list(tmp_path.glob("*!*")), "cell locators must never become filesystem paths"


def test_verified_publication_resumes(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    first = run_input(source, tmp_path / "out"); second = run_input(source, tmp_path / "out")
    assert first.status.value in {"complete", "partial"} and second.output_dir == first.output_dir
    assert second.diagnostics == ("resumed_verified_publication",)


def test_manifest_status_tamper_cannot_turn_partial_into_resumed_complete(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    first = run_input(source, tmp_path / "out")
    manifest = json.loads(first.manifest_path.read_text())
    manifest["status"] = "complete" if manifest["status"] == "partial" else "partial"
    first.manifest_path.write_text(json.dumps(manifest, sort_keys=True))
    rebuilt = run_input(source, tmp_path / "out")
    assert rebuilt.diagnostics != ("resumed_verified_publication",)
    evidence = json.loads(rebuilt.evidence_path.read_text())
    assert json.loads(rebuilt.manifest_path.read_text())["status"] == evidence["document"]["status"]


def test_tampered_evidence_rebuild_uses_new_immutable_export_record(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    output = tmp_path / "out"; first = run_input(source, output)
    altered = json.loads(first.evidence_path.read_text())
    altered["facts"][0]["amount"]["normalized_base"] = "999999999"
    first.evidence_path.write_text(json.dumps(altered))
    rebuilt = run_input(source, output)
    assert rebuilt.status.value in {"complete", "partial"}
    assert rebuilt.diagnostics != ("resumed_verified_publication",)
    repaired = json.loads(rebuilt.evidence_path.read_text())
    assert all(fact["amount"]["normalized_base"] != "999999999" for fact in repaired["facts"])
    ledger = Ledger(output / ".carthagent" / "ledger.sqlite3")
    export_records = ledger.db.execute("SELECT id FROM records WHERE kind='export'").fetchall()
    checkpoint = ledger.get_checkpoint(rebuilt.document_id, "publication_verified")
    ledger.close()
    assert export_records and checkpoint["export_record_id"] in {row[0] for row in export_records}
    assert not any("immutable record collision" in diagnostic for diagnostic in rebuilt.diagnostics)


def test_concurrent_same_document_publication_converges_on_verified_pair(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    output = tmp_path / "out"
    with ThreadPoolExecutor(max_workers=2) as pool:
        left, right = tuple(pool.map(lambda _: run_input(source, output), range(2)))
    assert left.status.value in {"complete", "partial"} and right.status.value in {"complete", "partial"}
    assert left.output_dir == right.output_dir
    manifest = json.loads(left.manifest_path.read_text())
    assert manifest["readback_verified"] is True and left.output_dir.joinpath(".complete").read_text().strip() == manifest["snapshot_digest"]


def test_concurrent_identical_bytes_under_distinct_names_share_source_safely(tmp_path):
    first_source = tmp_path / "alpha.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(first_source)
    second_source = tmp_path / "beta.xlsx"; second_source.write_bytes(first_source.read_bytes())
    output = tmp_path / "out"
    with ThreadPoolExecutor(max_workers=2) as pool:
        left, right = tuple(pool.map(lambda source: run_input(source, output), (first_source, second_source)))
    assert {left.status.value, right.status.value} <= {"complete", "partial"}
    assert left.source_hash == right.source_hash and left.document_id != right.document_id
    assert left.output_dir != right.output_dir


def test_component_version_change_rejects_same_directory_stale_resume(tmp_path, monkeypatch):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    output = tmp_path / "out"; first = run_input(source, output)
    assert run_input(source, output).diagnostics == ("resumed_verified_publication",)
    monkeypatch.setattr(pipeline_module, "PIPELINE_VERSION", "0.8.0-next-test")
    rebuilt = run_input(source, output)
    assert rebuilt.run_id != first.run_id and rebuilt.diagnostics != ("resumed_verified_publication",)


def test_legacy_source_record_with_absolute_path_does_not_collide_on_replay(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    output = tmp_path / "out"; digest = sha256(source.read_bytes()).hexdigest()
    ledger = Ledger(output / ".carthagent" / "ledger.sqlite3")
    ledger.append(f"source:{digest}", "legacy", "source_bytes", {"source_hash": digest, "stored_path": "/old/run/local/path", "size": source.stat().st_size})
    ledger.close()
    result = run_input(source, output)
    assert result.status.value in {"complete", "partial"} and not any("immutable record collision" in item for item in result.diagnostics)


def test_source_bytes_are_immutable_and_corrupt_artifact_is_rebuilt(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    first = run_input(source, tmp_path / "out")
    stored = tmp_path / "out" / ".carthagent" / "sources" / first.source_hash
    assert stored.read_bytes() == source.read_bytes() and stored.stat().st_mode & 0o222 == 0
    assert first.review_path is not None
    with first.review_path.open("ab") as stream:
        stream.write(b"tamper")
    second = run_input(source, tmp_path / "out")
    assert second.status.value in {"complete", "partial"} and second.diagnostics != ("resumed_verified_publication",)
    assert load_workbook(second.review_path, read_only=True)["_carthagent_manifest"]["B3"].value


def test_invalidated_fact_closure_prevents_stale_publication_resume(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Opening cost EUR 2025"]); ws.append(["A-1", 12]); wb.save(source)
    output = tmp_path / "out"; first = run_input(source, output)
    evidence = json.loads(first.evidence_path.read_text())
    observation_id = evidence["facts"][0]["pointer"]["observation_id"]
    ledger = Ledger(output / ".carthagent" / "ledger.sqlite3")
    invalidated = ledger.invalidate(observation_id, "semantic")
    assert any("binding" in record_id for record_id in invalidated) and any("export" in record_id for record_id in invalidated)
    ledger.close()
    second = run_input(source, output)
    assert second.status.value in {"complete", "partial"}
    assert second.diagnostics != ("resumed_verified_publication",)


def test_same_basename_sources_have_distinct_document_namespaces(tmp_path):
    left = tmp_path / "left"; right = tmp_path / "right"; left.mkdir(); right.mkdir()
    for folder, value in ((left, 1), (right, 2)):
        wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Asset ID", "Amount EUR"]); ws.append(["A-1", value]); wb.save(folder / "same.xlsx")
    first = run_input(left / "same.xlsx", tmp_path / "out"); second = run_input(right / "same.xlsx", tmp_path / "out")
    assert first.status.value in {"complete", "partial"} and second.status.value in {"complete", "partial"}
    assert first.document_id != second.document_id and first.output_dir != second.output_dir


def test_complete_rollforward_flags_only_explicit_inconsistency(tmp_path):
    source = tmp_path / "source-error.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["Jahresabschluss per 31.12.2025"])
    ws.append(["Asset ID", "AHK 01.01 EUR", "Zugänge EUR", "Abgänge EUR", "Umbuchungen EUR", "AHK 31.12 EUR"])
    ws.append(["A-1", 100, 60, 0, 0, 100]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, _, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    assert len(facts) == 5 and all(f.source_anomaly == "arithmetic_inconsistency:ahk_rollforward" for f in facts)


def test_chunking_honors_configured_page_and_row_bounds():
    pdf = tuple(SourceObservation(f"d:p:{i}", "d", "a" * 64, "word", f"page:{i + 1}:word:0", "pdf_native", "x", metadata={"page_index": i}) for i in range(45))
    chunks = _observation_chunks(pdf, DEFAULT_BOUNDS)
    assert [len(chunk) for chunk in chunks] == [20, 20, 9]
    assert len({o.id for o in chunks[0]} & {o.id for o in chunks[1]}) == DEFAULT_BOUNDS.page_overlap
    rows = tuple(SourceObservation(f"d:r:{i}", "d", "a" * 64, "logical_row", f"S!row:{i}", "xlsx_structure", "x", metadata={"row": i}) for i in range(1, 601))
    row_chunks = _observation_chunks(rows, DEFAULT_BOUNDS)
    assert max(len(chunk) for chunk in row_chunks) <= DEFAULT_BOUNDS.row_chunk_size + DEFAULT_BOUNDS.row_overlap


def test_zero_document_timeout_is_explicit_failure(tmp_path):
    source = tmp_path / "clear.xlsx"; wb = Workbook(); ws = wb.active; ws.append(["Asset ID", "Amount EUR"]); ws.append(["A-1", 1]); wb.save(source)
    config = RunConfig(bounds=replace(DEFAULT_BOUNDS, document_timeout_seconds=0))
    result = run_input(source, tmp_path / "out", config)
    assert result.status.value == "failed" and "document_timeout" in result.diagnostics[0]


def test_xlsx_deadline_retains_unprocessed_sheet_coverage(tmp_path, monkeypatch):
    source = tmp_path / "two-sheets.xlsx"; wb = Workbook(); wb.active.title = "First"; wb.active.append(["Asset ID", "Opening cost EUR 2025"]); wb.active.append(["A-1", 1]); wb.create_sheet("Second").append(["Asset ID", "Opening cost EUR 2025"]); wb["Second"].append(["B-1", 2]); wb.save(source)
    ticks = iter((0.0, 0.0, 0.0, 2.0))
    monkeypatch.setattr(acquire_xlsx_module.time, "monotonic", lambda: next(ticks, 2.0))
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS, deadline=1.0)
    worksheets = {item.locator: item for item in acquired.coverage if item.kind == "worksheet"}
    assert acquired.partial and set(worksheets) == {"First", "Second"}
    assert all(item.disposition == CoverageDisposition.UNRESOLVED for item in worksheets.values())
    assert all(item.reason == "unprocessed_document_timeout" for item in worksheets.values())


def test_merged_header_tree_periods_and_section_row_zero_gate(tmp_path):
    source = tmp_path / "hierarchy.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Schedule"
    ws.append(["Example entity"]); ws.append(["Fiscal period 01.07.2024 to 30.06.2025 — EUR"]); ws.append([]); ws.append([])
    for span, label in (("C5:G5", "Anschaffungs- und Herstellungskosten"), ("H5:L5", "Kumulierte Abschreibungen"), ("M5:N5", "Buchwerte")):
        ws.merge_cells(span); ws[span.split(":")[0]] = label
    headers = ["Anl.-Nr.", "Bezeichnung", "01.07.2024", "Zugänge", "Abgänge", "Umbuchungen", "30.06.2025", "01.07.2024", "Geschäftsjahr", "Abgänge", "Zuschreib.", "30.06.2025", "30.06.2025", "30.06.2024"]
    for column, value in enumerate(headers, 1): ws.cell(6, column, value)
    ws.append(["0077", "Software section"])
    ws.append(["077001", "Scheduler", 1000, 50, 20, 0, 1030, 500, 100, 10, 0, 590, 440, 500])
    ws.append(["077002", "Dispatch", 600, None, 0, 0, 600, 200, 60, 0, 0, 260, 340, 400]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, _, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    by_cell = {fact.source_pointer["cell"]: fact for fact in facts}
    expected = {"C8": "ahk_open", "D8": "zugang", "E8": "abgang", "G8": "ahk_close", "H8": "kum_afa_open", "J8": "afa_abgang", "L8": "kum_afa_close", "M8": "bw_close", "N8": "bw_prev"}
    assert all(by_cell[cell].business.field == field and by_cell[cell].business.period == "2025" for cell, field in expected.items())
    assert "D9" in by_cell and by_cell["D9"].amount.normalized_value == 0
    assert not any(cell.endswith("7") for cell in by_cell)


def test_column_local_unit_and_parallel_period_override(tmp_path):
    source = tmp_path / "axes.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Schedule"
    ws.append(["Asset schedule — EUR"]); ws.append(["Asset ID", "Name", "Additions 2024", "Additions 2025 (TEUR)"]); ws.append(["A-1", "Machine", 111, 2]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, _, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    by_cell = {fact.source_pointer["cell"]: fact for fact in facts}
    assert by_cell["C3"].business.period == "2024" and by_cell["C3"].amount.normalized_value == 111
    assert by_cell["D3"].business.period == "2025" and by_cell["D3"].amount.normalized_value == 2000


def test_asset_subnumber_is_a_separate_business_identity_axis(tmp_path):
    source = tmp_path / "subnumbers.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Register"
    ws.append(["Asset ID", "Subnumber", "Name", "Opening cost EUR 2025"])
    ws.append(["000101", "0", "Machine base", 100]); ws.append(["000101", "1", "Machine extension", 25]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    facts = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)[1]
    assert [(fact.business.asset_id, fact.business.asset_sub_id, fact.amount.normalized_value) for fact in facts] == [
        ("000101", "0", Decimal("100")), ("000101", "1", Decimal("25")),
    ]


def test_pdf_column_local_unit_does_not_scale_neighbor_column():
    def obs(suffix, kind, text, bbox, metadata=None, channel="pdf_native"):
        return SourceObservation(f"d:{suffix}", "d", "a" * 64, kind, f"page:1:{suffix}", channel, text, bbox, metadata or {"page_index": 0})
    asset_header = obs("asset-header", "logical_row", "Asset ID", (10, 20, 70, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_structure")
    first_header = obs("first-header", "logical_row", "Additions 2025", (100, 20, 180, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_structure")
    second_header = obs("second-header", "logical_row", "Additions 2025", (200, 20, 280, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_structure")
    eur = obs("eur", "word", "EUR", (140, 35, 160, 45)); teur = obs("teur", "word", "TEUR", (235, 35, 265, 45))
    asset = obs("asset", "word", "A-1", (15, 80, 45, 90)); first = obs("first", "word", "3", (130, 80, 150, 90)); second = obs("second", "word", "4", (230, 80, 250, 90))
    row = obs("row", "logical_row", "A-1 3 4", (10, 80, 250, 90), {"page_index": 0, "structure_role": "visual_row", "member_ids": [asset.id, first.id, second.id]}, "pdf_structure")
    observations = (asset_header, first_header, second_header, eur, teur, asset, first, second, row)
    coverage = tuple(CoverageItem(f"d:c:{item.id}", "d", "pdf_word", item.locator, CoverageDisposition.UNRESOLVED) for item in (first, second))
    facts = interpret(observations, coverage, DEFAULT_POLICIES)[1]
    values = {fact.source_observation_id: fact.amount.normalized_value for fact in facts}
    assert values[first.id] == 3 and values[second.id] == 4000


def _pdf_observation(suffix, kind, text, bbox, metadata=None, channel="pdf_ocr"):
    return SourceObservation(f"d:{suffix}", "d", "a" * 64, kind, f"page:1:{suffix}", channel, text, bbox, metadata or {"page_index": 0})


def test_pdf_ocr_section_id_shape_disagreement_withholds_entire_row():
    asset_header = _pdf_observation("asset-header", "word", "Asset ID", (10, 20, 60, 30))
    amount_header = _pdf_observation("amount-header", "word", "Additions 2025", (100, 20, 160, 30))
    header = _pdf_observation("header", "logical_row", "Asset ID Additions", (10, 20, 160, 30), {"page_index": 0, "structure_role": "visual_row", "member_ids": [asset_header.id, amount_header.id]}, "pdf_ocr_structure")
    unit = _pdf_observation("unit", "word", "EUR", (120, 35, 145, 45))
    observations = [asset_header, amount_header, header, unit]
    amounts = []
    for index, asset_id in enumerate(("0210", "210001", "210002", "2100038")):
        top = 60 + index * 20
        asset = _pdf_observation(f"asset-{index}", "word", asset_id, (10, top, 55, top + 10))
        members = [asset.id]
        row_text = asset_id
        if index:
            amount = _pdf_observation(f"amount-{index}", "word", "100", (110, top, 140, top + 10)); observations.append(amount); amounts.append(amount); members.append(amount.id); row_text += " 100"
        observations.extend((asset, _pdf_observation(f"row-{index}", "logical_row", row_text, (10, top, 140, top + 10), {"page_index": 0, "structure_role": "visual_row", "member_ids": members}, "pdf_ocr_structure")))
    coverage = tuple(CoverageItem(f"d:c:{item.id}", "d", "pdf_word", item.locator, CoverageDisposition.UNRESOLVED) for item in amounts)
    _, facts, abstentions, _ = interpret(tuple(observations), coverage, DEFAULT_POLICIES)
    assert {fact.source_observation_id for fact in facts} == {amounts[0].id, amounts[1].id}
    assert any(amounts[2].id in abstention.candidate_observation_ids and "ocr_identity_disagreement" in abstention.reason for abstention in abstentions)


def test_pdf_header_span_bounds_unknown_columns_and_same_slot_duplicates():
    asset_header = _pdf_observation("asset-header", "logical_row", "Asset ID", (10, 20, 60, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    field_header = _pdf_observation("field-header", "logical_row", "Write-ups 2025", (100, 20, 160, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    unit = _pdf_observation("unit", "word", "EUR", (120, 35, 145, 45))
    asset = _pdf_observation("asset", "word", "A-1", (10, 80, 45, 90)); supported = _pdf_observation("supported", "word", "10", (125, 80, 145, 90)); unknown = _pdf_observation("unknown", "word", "20", (200, 80, 225, 90))
    row = _pdf_observation("row", "logical_row", "A-1 10 20", (10, 80, 225, 90), {"page_index": 0, "structure_role": "visual_row", "member_ids": [asset.id, supported.id, unknown.id]}, "pdf_ocr_structure")
    observations = (asset_header, field_header, unit, asset, supported, unknown, row)
    coverage = tuple(CoverageItem(f"d:c:{item.id}", "d", "pdf_word", item.locator, CoverageDisposition.UNRESOLVED) for item in (supported, unknown))
    _, facts, abstentions, _ = interpret(observations, coverage, DEFAULT_POLICIES)
    assert [fact.source_observation_id for fact in facts] == [supported.id]
    assert any(unknown.id in abstention.candidate_observation_ids for abstention in abstentions)

    second_header = _pdf_observation("second-header", "logical_row", "Additions EUR 2025", (200, 20, 260, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    first_header = replace(field_header, text="Additions EUR 2025")
    first = replace(supported, text="100"); second = replace(unknown, text="100", bbox=(220, 80, 240, 90))
    duplicate_observations = (asset_header, first_header, second_header, unit, asset, first, second, replace(row, text="A-1 100 100", metadata={"page_index": 0, "structure_role": "visual_row", "member_ids": [asset.id, first.id, second.id]}))
    _, duplicate_facts, duplicate_abstentions, _ = interpret(duplicate_observations, coverage, DEFAULT_POLICIES)
    assert duplicate_facts == () and sum(item.reason == "duplicate_same_row_field_ambiguous" for item in duplicate_abstentions) == 2


def test_pdf_column_polarity_header_never_rewrites_source_sign():
    asset_header = _pdf_observation("asset-header-sign", "logical_row", "Asset ID", (10, 20, 60, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    field_header = _pdf_observation("field-header-sign", "logical_row", "Disposals 2025", (100, 20, 160, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    unit = _pdf_observation("unit-sign", "word", "(-)EUR", (115, 35, 150, 45))
    asset = _pdf_observation("asset-sign", "word", "A-1", (10, 80, 45, 90)); amount = _pdf_observation("amount-sign", "word", "100", (120, 80, 145, 90))
    row = _pdf_observation("row-sign", "logical_row", "A-1 100", (10, 80, 145, 90), {"page_index": 0, "structure_role": "visual_row", "member_ids": [asset.id, amount.id]}, "pdf_ocr_structure")
    fact = interpret((asset_header, field_header, unit, asset, amount, row), (CoverageItem("d:c:sign", "d", "pdf_word", amount.locator, CoverageDisposition.UNRESOLVED),), DEFAULT_POLICIES)[1][0]
    assert fact.amount.raw_lexeme == "100" and fact.amount.source_signed_value == Decimal("100") and fact.amount.sign_notation == "leading_plus_or_unsigned"


def test_inventory_number_identity_and_compound_movement_use_source_basis(tmp_path):
    source = tmp_path / "inventory.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Register"
    ws.append(["Jahresabschluss per 30.06.2025 — Historische Anschaffungs- und Herstellungskosten — EUR"])
    ws.append(["Inventar- Nummer", "Inventar- Bezeichnung", "Abgänge / Minderungen (-) EUR"])
    ws.append(["77", "CAD Software", 25]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, abstentions, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    assert len(facts) == 1 and not abstentions
    fact = facts[0]
    assert fact.business.asset_id == "77" and fact.business.source_name == "CAD Software"
    assert fact.business.field == "abgang" and fact.amount.source_signed_value == Decimal("25")


def test_mistral_recognition_is_candidate_only_and_consumed_through_two_sided_anchor():
    before = _pdf_observation("anchor-before", "word", "Asset", (1, 10, 8, 20))
    wrong = _pdf_observation("anchor-id", "word", "2100038", (10, 10, 30, 20))
    after = _pdf_observation("anchor-after", "word", "100", (40, 10, 50, 20))
    evidence = (
        SimpleNamespace(text="Asset", confidence=0.99, start_index=0),
        SimpleNamespace(text="210003", confidence=0.98, start_index=6),
        SimpleNamespace(text="100", confidence=0.99, start_index=13),
    )
    table = SimpleNamespace(id="tbl-0.html", recognition_evidence=evidence)
    page = SimpleNamespace(recognition_evidence=(), tables=(table,))
    ownership_row = _pdf_observation("candidate-owner", "logical_row", "Asset 2100038 100", (1, 10, 50, 20), {"page_index": 0, "structure_role": "visual_row", "member_ids": [before.id, wrong.id, after.id]}, "pdf_ocr_structure")
    candidates = _mistral_recognition_candidates(page, [before, wrong, after, ownership_row], "d", "a" * 64, 0)
    assert len(candidates) == 1
    candidate = candidates[0]
    assert candidate.text == "210003" and candidate.bbox == wrong.bbox
    assert candidate.metadata["anchor_observation_id"] == wrong.id
    assert candidate.metadata["provider_table_ids"] == ("tbl-0.html",)
    assert candidate.metadata["authoritative_row_id"] == ownership_row.id
    assert candidate.metadata["ownership_closed"] is True
    assert candidate.metadata["candidate_only"] is True and candidate.metadata["cell_structure_authoritative"] is False

    asset_header = _pdf_observation("candidate-asset-header", "logical_row", "Asset ID", (10, 20, 60, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    field_header = _pdf_observation("candidate-field-header", "logical_row", "Additions 2025", (100, 20, 160, 30), {"page_index": 0, "structure_role": "native_line"}, "pdf_ocr_structure")
    unit = _pdf_observation("candidate-unit", "word", "EUR", (120, 35, 145, 45))
    anchored_asset = replace(wrong, bbox=(10, 80, 55, 90))
    amount = _pdf_observation("candidate-amount", "word", "100", (120, 80, 145, 90))
    row = _pdf_observation("candidate-row", "logical_row", "2100038 100", (10, 80, 145, 90), {"page_index": 0, "structure_role": "visual_row", "member_ids": [anchored_asset.id, amount.id]}, "pdf_ocr_structure")
    candidate = replace(candidate, bbox=anchored_asset.bbox, metadata={**dict(candidate.metadata), "authoritative_row_id": row.id})
    fact = interpret((asset_header, field_header, unit, anchored_asset, amount, row, candidate), (CoverageItem("d:c:candidate", "d", "pdf_word", amount.locator, CoverageDisposition.UNRESOLVED),), DEFAULT_POLICIES)[1][0]
    assert fact.business.asset_id == "210003" and candidate.id in fact.dependency_ids


def test_mistral_table_candidate_with_ambiguous_physical_ownership_never_overrides_identity():
    before_a = _pdf_observation("before-a", "word", "Asset", (1, 10, 8, 20))
    wrong_a = _pdf_observation("wrong-a", "word", "2100038", (10, 10, 30, 20))
    after_a = _pdf_observation("after-a", "word", "100", (40, 10, 50, 20))
    before_b = _pdf_observation("before-b", "word", "Asset", (1, 30, 8, 40))
    wrong_b = _pdf_observation("wrong-b", "word", "2100038", (10, 30, 30, 40))
    after_b = _pdf_observation("after-b", "word", "100", (40, 30, 50, 40))
    row_a = _pdf_observation("row-a", "logical_row", "Asset 2100038 100", (1, 10, 50, 20), {"page_index": 0, "structure_role": "visual_row", "member_ids": [before_a.id, wrong_a.id, after_a.id]}, "pdf_ocr_structure")
    row_b = _pdf_observation("row-b", "logical_row", "Asset 2100038 100", (1, 30, 50, 40), {"page_index": 0, "structure_role": "visual_row", "member_ids": [before_b.id, wrong_b.id, after_b.id]}, "pdf_ocr_structure")
    evidence = tuple(SimpleNamespace(text=text, confidence=0.99, start_index=index) for index, text in enumerate(("Asset", "210003", "100")))
    page = SimpleNamespace(recognition_evidence=(), tables=(SimpleNamespace(id="tbl-0.html", recognition_evidence=evidence),))
    candidates = _mistral_recognition_candidates(page, [before_a, wrong_a, after_a, before_b, wrong_b, after_b, row_a, row_b], "d", "a" * 64, 0)
    assert candidates == []


def test_unknown_explicit_identifier_column_blocks_name_fallback(tmp_path):
    source = tmp_path / "unknown-id.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Register"
    ws.append(["Register 2025 — EUR"])
    ws.append(["Record Number", "Description", "Additions EUR"])
    ws.append(["77", "CAD Software", 25]); wb.save(source)
    acquired = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS)
    _, facts, abstentions, _ = interpret(acquired.observations, acquired.coverage, DEFAULT_POLICIES)
    assert facts == ()
    assert any("unresolved_identifier" in item.reason for item in abstentions)
