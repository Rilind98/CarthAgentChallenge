"""Round 4 (cost): skeleton requests + block patterns (R4-1), skipping regions the model cannot help with (R4-2)."""
from __future__ import annotations

from decimal import Decimal
import importlib.util
import json
from pathlib import Path

import pytest

from carthagent.acquire_xlsx import acquire_xlsx
from carthagent.mapper import MAPPER_PROMPT_VERSION, MAPPER_SCHEMA_VERSION, MapperConfig, OpenAICompatibleMapper
from carthagent.pipeline import RunConfig, run_input
from carthagent.policies import DEFAULT_BOUNDS

ROOT = Path(__file__).resolve().parents[2]


def _load(name: str):
    spec = importlib.util.spec_from_file_location(f"r4_{name}", ROOT / "tests" / "fixtures" / f"{name}.py")
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


fixtures = _load("build_option_b"); stub = _load("option_b_stub")
S = "Anlagengitter"


def _body(observations, **config):
    mapper = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "m", 1, 1, **config))
    data, body = mapper._request_body("d", observations)
    return data, len(json.dumps(body, separators=(",", ":")).encode())


def test_skeleton_sends_header_band_and_three_sample_blocks_only(tmp_path):
    source, _ = fixtures.sap_block_many(tmp_path / "many.xlsx", blocks=30)
    observations = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS).observations
    full, full_bytes = _body(observations, skeleton_sample_blocks=None)
    skeleton, skeleton_bytes = _body(observations)
    assert len(full) >= 5 * len(skeleton) and full_bytes >= 5 * skeleton_bytes, (len(full), len(skeleton), full_bytes, skeleton_bytes)
    rows = {item["row"] for item in skeleton if item["kind"] == "logical_row"}
    assert {12, 15, 16, 17, 18} <= rows, "the header band travels in full"
    blocks = [item for item in skeleton if item["kind"] == "record_block"]
    sampled = [b for b in blocks if b.get("sampled") is True]
    assert len(sampled) == 3 and len(blocks) == 30, "first/middle/last in full, every other block as a summary"
    assert {b["identity_row_id"] for b in sampled} == {f"d:xlsx:{S}:row:{r}" for r in (19, 19 + 15 * 5, 19 + 29 * 5)}
    assert all(set(b) >= {"id", "identity_row_id", "member_count", "subtotal_marker"} for b in blocks)
    cells = {(item["row"], item["column"]) for item in skeleton if item["kind"] in {"cell", "empty_cell"}}
    assert (19 + 5 * 5 + 1, 2) not in cells, "unsampled rows send nothing about their cells"
    assert MAPPER_SCHEMA_VERSION == "semantic_region_plan_v9" and MAPPER_PROMPT_VERSION == "compact-region-plan-v12"


def test_odd_shaped_block_is_an_anchor_in_the_skeleton(tmp_path):
    source, _ = fixtures.sap_block_many(tmp_path / "odd.xlsx", blocks=30, odd_at=7)
    skeleton, _ = _body(acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS).observations)
    sampled = {b["identity_row_id"] for b in skeleton if b["kind"] == "record_block" and b.get("sampled")}
    assert f"d:xlsx:{S}:row:{19 + 7 * 5}" in sampled and len(sampled) == 4


def pattern_plan(rows=((1, 16, "ahk"), (2, 17, "cum_depreciation"), (3, 18, "book_value"))):
    plan = stub.load_plan("sap_block_unlabelled"); q = plan["plans"][0]
    q["record_blocks"] = []
    q["block_pattern"] = {"rows": [{"ordinal": o, "header_row_id": f"@row:{S}:{h}", "measure": m} for o, h, m in rows], "role": "data", "applies_to": "all_blocks_of_same_shape"}
    q["unit_contexts"] = [{"observation_id": f"@cell:{S}!H19", "unit": "EUR", "scope": "block"}]
    return plan


def _facts(tmp_path, source, plan):
    with stub.recorded_plan_server(plan) as endpoint:
        result = run_input(source, tmp_path / "out", RunConfig(mapper=MapperConfig(True, endpoint, "local-stub", 2, 256)))
    assert result.evidence_path is not None, result.diagnostics
    return json.loads(Path(result.evidence_path).read_text())["facts"], len(stub.RecordedPlanHandler.requests)


def _rows(facts):
    return {f["pointer"]["cell"]: (f["field"], f["asset_identity"]["class_code"], Decimal(f["amount"]["normalized_base"])) for f in facts}


def test_pattern_plan_binds_every_block_of_the_sampled_shape(tmp_path):
    source, truth = fixtures.sap_block_many(tmp_path / "many.xlsx", blocks=30)
    facts, requests = _facts(tmp_path, source, pattern_plan())
    assert requests == 1
    assert _rows(facts) == {cell: (field, klass, Decimal(str(value))) for cell, (field, klass, value) in truth.items()}
    assert {f["asset_identity"]["period"] for f in facts} == {"2025"} and {f["amount"]["unit"] for f in facts} == {"EUR"}


def test_pattern_never_expands_onto_a_block_of_another_shape(tmp_path):
    source, truth = fixtures.sap_block_many(tmp_path / "odd.xlsx", blocks=30, odd_at=7)
    facts, _ = _facts(tmp_path, source, pattern_plan())
    odd_rows = {19 + 7 * 5 + k for k in (1, 2)}
    assert not [cell for cell in _rows(facts) if int(cell[1:]) in odd_rows]
    expected = {cell: (field, klass, Decimal(str(value))) for cell, (field, klass, value) in truth.items() if int(cell[1:]) not in odd_rows}
    assert _rows(facts) == expected


def test_swapped_pattern_binds_nothing(tmp_path):
    source, _ = fixtures.sap_block_many(tmp_path / "many.xlsx", blocks=30)
    facts, _ = _facts(tmp_path, source, pattern_plan(((1, 17, "cum_depreciation"), (2, 16, "ahk"), (3, 18, "book_value"))))
    assert not facts


def test_flat_register_skeleton_keeps_header_and_samples_and_field_headers_bind_all_rows(tmp_path):
    source, truth = fixtures.flat_register_many(tmp_path / "reg.xlsx", rows=40)
    observations = acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS).observations
    skeleton, skeleton_bytes = _body(observations)
    _, full_bytes = _body(observations, skeleton_sample_blocks=None)
    data_rows = sorted(item["row"] for item in skeleton if item["kind"] == "logical_row" and item["row"] >= 14)
    assert data_rows == [*range(14, 26), 51, 52, 53] and full_bytes >= 2 * skeleton_bytes
    assert {item["row"] for item in skeleton if item["kind"] == "logical_row"} >= {4, 13}
    region = next(item for item in skeleton if item["kind"] == "table_region")
    assert region["row_count"] == 42 and region["data_row_count"] == 40
    plan = stub.load_plan("flat_sap_register"); q = plan["plans"][0]
    q["record_blocks"] = []; q["unit_contexts"] = []
    facts, _ = _facts(tmp_path, source, plan)
    assert {f["pointer"]["cell"]: (f["field"], f["asset_identity"]["asset_id"], Decimal(f["amount"]["normalized_base"])) for f in facts} == \
        {cell: (field, asset, Decimal(str(value))) for cell, (field, asset, value) in truth.items()}


# --- R4-2 ---

def test_regions_without_residual_values_are_not_sent():
    from carthagent.models import SourceObservation
    from carthagent.pipeline import _regions_needing_mapper
    def obs(identifier, kind, text=None, metadata=None):
        return SourceObservation(f"d:{identifier}", "d", "a" * 64, kind, identifier, "xlsx_native" if kind == "cell" else "xlsx_structure", text, metadata=metadata or {})
    a1, b1 = obs("A!B2", "cell", "1"), obs("B!B2", "cell", "2")
    ra, rb = obs("A:row:2", "logical_row", metadata={"member_ids": [a1.id]}), obs("B:row:2", "logical_row", metadata={"member_ids": [b1.id]})
    rega, regb = obs("A:region", "table_region", metadata={"member_ids": [ra.id]}), obs("B:region", "table_region", metadata={"member_ids": [rb.id]})
    kept = _regions_needing_mapper((a1, b1, ra, rb, rega, regb), {b1.id})
    assert {o.id for o in kept} == {b1.id, rb.id, regb.id}
    assert _regions_needing_mapper((a1, ra, rega), set()) == ()


def test_document_resolved_locally_makes_no_model_call(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / "clear.xlsx"; workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"
    sheet.append(["Fixed asset schedule fiscal 2025 EUR"]); sheet.append(["Asset ID", "Opening cost EUR 2025"]); sheet.append(["A-1", 12]); workbook.save(source)
    facts, requests = _facts(tmp_path, source, {"plans": []})
    assert facts and requests == 0


def test_request_cap_defaults_to_six_and_keeps_local_facts(tmp_path):
    from openpyxl import Workbook
    assert DEFAULT_BOUNDS.max_mapper_requests == 6
    source = tmp_path / "many_sheets.xlsx"; workbook = Workbook(); first = workbook.active; first.title = "Clear"
    first.append(["Fixed asset schedule fiscal 2025 EUR"]); first.append(["Asset ID", "Opening cost EUR 2025"]); first.append(["A-1", 12])
    for index in range(8):
        sheet = workbook.create_sheet(f"Open{index}"); sheet.append(["Mystery", "Thing"]); sheet.append(["x", 5])
    workbook.save(source)
    with stub.recorded_plan_server({"plans": []}) as endpoint:
        result = run_input(source, tmp_path / "out", RunConfig(mapper=MapperConfig(True, endpoint, "local-stub", 2, 256)))
    facts = json.loads(Path(result.evidence_path).read_text())["facts"]
    assert len(stub.RecordedPlanHandler.requests) <= 6 and any(item.startswith("mapper_budget_exhausted") for item in result.diagnostics)
    assert [f for f in facts if f["pointer"]["sheet"] == "Clear"]
