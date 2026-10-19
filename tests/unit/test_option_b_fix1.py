"""Fix round 1 majors/minors: each test reproduces a reviewed defect (M1-M6, m1, m2, m4, cache replay)."""
from __future__ import annotations

import importlib.util
from pathlib import Path
import time

from openpyxl import load_workbook
import pytest

from carthagent.acquire_xlsx import acquire_xlsx
from carthagent.mapper import MapperConfig, MapperUnavailableError, MapperValidationError, OpenAICompatibleMapper
from carthagent.models import SourceObservation, jsonable
from carthagent.pipeline import _cached_hypotheses
from carthagent.policies import DEFAULT_BOUNDS, DEFAULT_POLICIES
from carthagent.semantics import interpret

ROOT = Path(__file__).resolve().parents[2]


def _load(name: str):
    spec = importlib.util.spec_from_file_location(f"fix1_{name}", ROOT / "tests" / "fixtures" / f"{name}.py")
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


fixtures = _load("build_option_b"); stub = _load("option_b_stub")


def _map(source, plan, keep=lambda o: True):
    observations = tuple(o for o in acquire_xlsx(source, "d", "a" * 64, DEFAULT_BOUNDS).observations if keep(o))
    with stub.recorded_plan_server(plan) as endpoint:
        mapper = OpenAICompatibleMapper(MapperConfig(True, endpoint, "local-stub", 2, 256))
        mapped = mapper.map("d", observations)
    return observations, mapped, mapper


def _blocked(mapped):
    return [h for h in mapped if h.candidates[0].record_identity_row_id is not None]


@pytest.mark.parametrize("where", ["identity_row_id", "row_id", "span_header", "field_header"])
def test_m1_unhashable_ids_reject_the_plan_instead_of_crashing(tmp_path, where):
    plan = stub.load_plan("sap_block_unlabelled"); q = plan["plans"][0]
    if where == "identity_row_id":
        q["record_blocks"][0]["identity_row_id"] = ["@row:Anlagengitter:19"]
    elif where == "row_id":
        q["record_blocks"][0]["rows"][0]["row_id"] = {"id": "x"}
    elif where == "span_header":
        q["column_spans"] = [{"header_observation_id": ["x"], "first_column": 2, "last_column": 2}]
    else:
        q["field_headers"][0]["observation_id"] = ["x"]
    with pytest.raises(MapperValidationError):
        _map(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), plan)


def test_m2_cache_fingerprint_names_the_verifier_version():
    fingerprint = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "m", 1, 1)).fingerprint()
    assert fingerprint["verifier_version"] == "option-b-verifier-4"


def test_m3_subtotal_guard_holds_without_acquisition_record_blocks(tmp_path):
    plan = stub.load_plan("sap_block_unlabelled"); plan["plans"][0]["record_blocks"][2]["role"] = "data"
    _, mapped, mapper = _map(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), plan, keep=lambda o: o.kind != "record_block")
    assert not [h for h in _blocked(mapped) if h.amount_observation_id.endswith(("30", "31", "32"))]
    assert any(item["rule"] == "R4" for item in mapper.last_verification)


def test_m4_column_measure_needs_a_declared_header_row_above(tmp_path):
    plan = stub.load_plan("two_year_two_sheets")
    for q in plan["plans"]:
        q["header_rows"] = [h for h in q["header_rows"] if not h.endswith(":4")]
    _, mapped, _ = _map(fixtures.two_year_two_sheets(tmp_path / "s.xlsx"), plan)
    assert not _blocked(mapped)


def test_m5_verification_honours_the_deadline_and_record_blocks_are_linear(tmp_path):
    observations = acquire_xlsx(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), "d", "a" * 64, DEFAULT_BOUNDS).observations
    from carthagent.mapper import _proposal_context, _proposal_observation
    context = _proposal_context(observations, {o.id for o in observations})
    payload = [_proposal_observation(o, context) for o in observations]
    plan = stub.resolve(stub.load_plan("sap_block_unlabelled"), payload)
    mapper = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "m", 1, 1))
    with pytest.raises(MapperUnavailableError, match="mapper_document_timeout"):
        mapper._validate("d", observations, plan, deadline=time.monotonic() - 1)
    from carthagent.acquire_xlsx import _record_blocks
    rows = [(n, f"r{n}", "identity", "t") for n in range(1, 80_001)]
    started = time.monotonic(); _record_blocks(rows, "d", "a" * 64, "S")
    assert time.monotonic() - started < 2.0


@pytest.mark.parametrize(("text", "unit"), [("in Tsd. EUR", "TEUR"), ("Mio. EUR", None), ("Angaben in T€", "TEUR"), ("TEUR", "TEUR"), ("EUR", "EUR"), ("Mrd. EUR", None), ("Betrag", None)])
def test_m6_unit_text_never_endorses_eur_for_a_scaled_header(text, unit):
    from carthagent.structure import unit_from_text
    assert unit_from_text(text) == unit


def test_m6_block_unit_in_millions_abstains(tmp_path):
    source = fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx")
    workbook = load_workbook(source); ws = workbook.active
    for cell in ("H19", "H24"):
        ws[cell] = "Mio. EUR"
    workbook.save(source)
    observations, mapped, _ = _map(source, stub.load_plan("sap_block_unlabelled"))
    assert not interpret(observations, (), DEFAULT_POLICIES, mapped)[1]


def test_m1b_dropped_connection_keeps_local_facts(tmp_path):
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Thread

    class Hangup(BaseHTTPRequestHandler):
        def do_POST(self):
            self.rfile.read(int(self.headers["Content-Length"])); self.close_connection = True

        def log_message(self, *args):
            return

    server = ThreadingHTTPServer(("127.0.0.1", 0), Hangup); thread = Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        mapper = OpenAICompatibleMapper(MapperConfig(True, f"http://127.0.0.1:{server.server_port}/v1/responses", "m", 2, 256))
        with pytest.raises(MapperUnavailableError):
            mapper.map("d", acquire_xlsx(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), "d", "a" * 64, DEFAULT_BOUNDS).observations)
    finally:
        server.shutdown(); thread.join()


def test_minor_data_block_declared_subtotal_is_logged_and_float_spans_are_coerced(tmp_path):
    plan = stub.load_plan("sap_block_unlabelled"); q = plan["plans"][0]
    q["record_blocks"][0]["role"] = "subtotal"
    q["column_spans"] = [{"header_observation_id": "@cell:Anlagengitter!B16", "first_column": 2.0, "last_column": 2.0},
                         {"header_observation_id": "@cell:Anlagengitter!C16", "first_column": 3.5, "last_column": 3.5}]
    _, mapped, mapper = _map(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), plan)
    rules = [item["rule"] for item in mapper.last_verification]
    assert "role" in rules and rules.count("column_span") == 1 and _blocked(mapped)


def test_minor_flat_path_never_takes_a_typed_number_as_identity():
    def cell(ref, text, data_type="s"):
        return SourceObservation(f"d:x:{ref}", "d", "a" * 64, "cell", f"Data!{ref}", "xlsx_native", text, metadata={"data_type": data_type})
    header, amount_header, identity, amount = cell("A1", "Anlage"), cell("B1", "Zugang 2025 EUR"), cell("A2", "1234", "n"), cell("B2", "55", "n")
    row = SourceObservation("d:x:row:2", "d", "a" * 64, "logical_row", "Data!row:2", "xlsx_structure", "1234 | 55", metadata={"member_ids": [identity.id, amount.id], "sheet": "Data", "row": 2})
    region = SourceObservation("d:x:region", "d", "a" * 64, "table_region", "Data!region", "xlsx_structure", "Data", metadata={"member_ids": [row.id], "sheet": "Data"})
    plan = {"plans": [{"region_observation_id": region.id, "field_headers": [{"observation_id": amount_header.id, "semantic_field": "zugang", "parent_header_ids": []}],
                       "identity_headers": [{"observation_id": header.id, "identity_role": "asset_id"}], "measure_axes": [], "movement_axes": [], "row_roles": [],
                       "period_contexts": [], "unit_contexts": []}]}
    mapper = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "m", 1, 1))
    hypotheses = mapper._validate("d", (header, amount_header, identity, amount, row, region), plan)
    assert all(not h.candidates[0].semantic_identity_assignments for h in hypotheses)


def test_block_claims_survive_cache_replay(tmp_path):
    _, mapped, _ = _map(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), stub.load_plan("sap_block_unlabelled"))
    assert _blocked(mapped) and _cached_hypotheses([jsonable(h) for h in mapped]) == tuple(mapped)


def test_m3_split_requests_keep_whole_record_blocks(tmp_path, monkeypatch):
    from carthagent.pipeline import _mapper_bounded_chunks
    observations = acquire_xlsx(fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), "d", "a" * 64, DEFAULT_BOUNDS).observations
    mapper = OpenAICompatibleMapper(MapperConfig())
    real_fits = mapper.request_fits
    monkeypatch.setattr(mapper, "request_fits", lambda document_id, chunk: len(chunk) <= 140 and real_fits(document_id, chunk))
    chunks = _mapper_bounded_chunks(mapper, "d", observations, 3)
    blocks = {o.id for o in observations if o.kind == "record_block"}
    assert len(chunks) > 1 and blocks and blocks == {o.id for chunk in chunks for o in chunk if o.kind == "record_block"}
    for chunk in chunks:
        ids = {o.id for o in chunk}; region = next(o for o in chunk if o.kind == "table_region")
        for block in (o for o in chunk if o.kind == "record_block"):
            assert set(block.metadata["member_ids"]) <= ids and set(block.metadata["member_ids"]) <= set(region.metadata["member_ids"])
    assert all(mapper.request_fits("d", chunk) for chunk in chunks)
