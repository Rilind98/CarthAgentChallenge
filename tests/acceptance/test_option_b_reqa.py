"""Fix round 2: re-QA probes (qa_probes/test_reqa_*.py) and recorded LIVE Claude plans as repo tests.

Truth comes from fixture constants. Live plans (tests/fixtures/option_b_plans/live_*.json) are replayed against the
corpus files they were recorded on; those tests skip when the private corpus is not present.
"""
from __future__ import annotations

import collections
import copy
from decimal import Decimal
import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from carthagent.acquire_xlsx import acquire_xlsx
from carthagent.mapper import MapperConfig, OpenAICompatibleMapper
from carthagent.policies import DEFAULT_BOUNDS, DEFAULT_POLICIES
from carthagent.semantics import interpret

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location("reqa_adversarial", Path(__file__).with_name("test_option_b_adversarial.py"))
adv = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(adv)
fixtures, stub, run, key, wrong, plan, blk, sap = adv.fixtures, adv.stub, adv.run, adv.key, adv.wrong, adv.plan, adv.blk, adv.sap
CORPUS = ROOT / "hackathon-local" / "ElabHackathon" / "A3-Hackathon-Corpus"


# --- R2-4 / R2-6: identity headers and periods on the register shape (qa_probes/test_reqa_b2_b4.py) ---

def reg_go(tmp_path, mutate=None, idents=None, periods=None):
    recorded = plan("flat_sap_register"); q = recorded["plans"][0]
    if idents is not None:
        q["identity_headers"] = [{"observation_id": f"@cell:Sheet1!{c}", "identity_role": r} for c, r in idents]
    if periods is not None:
        q["period_contexts"] = periods
    _, facts = run(tmp_path, adv.register(tmp_path, mutate), recorded)
    keys = collections.Counter((f["asset_identity"].get("asset_or_class"), f["asset_identity"].get("asset_sub_id"), f["field"], f["asset_identity"].get("period")) for f in facts)
    return facts, wrong(facts, adv.reg_truth()), [k for k, n in keys.items() if n > 1]


def test_register_baseline_still_binds_every_value(tmp_path):
    facts, bad, collisions = reg_go(tmp_path)
    assert len(facts) == 36 and not bad and not collisions


def _remarks(ws):
    for row, text in ((14, "Umbuchung aus 100001002"), (15, "siehe Anlage 100001001"), (16, "Nachrüstung")):
        ws[f"V{row}"] = text


IDENTITY_PROBES = {
    "subnumber_column_as_asset_id": (None, (("B13", "asset_id"), ("D13", "source_name"))),
    "bemerkung_text_header_as_asset_id": (lambda ws: (ws.__setitem__("A13", "ID"), ws.__setitem__("V13", "Bemerkung"), _remarks(ws)), (("V13", "asset_id"), ("B13", "asset_sub_id"), ("D13", "source_name"))),
    "pos_header_as_asset_id": (lambda ws: (ws.__setitem__("C13", "Pos."), [ws.__setitem__(f"C{r}", t) for r, t in ((14, "10"), (15, "20"), (16, "30"))]), (("C13", "asset_id"), ("B13", "asset_sub_id"), ("D13", "source_name"))),
    "empty_header_source_name_as_key": (lambda ws: (ws.__setitem__("B13", "Kz"), _remarks(ws)), (("V13", "source_name"),)),
    "split_identifier_header": (lambda ws: [ws.__setitem__(c, t) for c, t in (("A12", "Anlagen-"), ("A13", "Nummer"), ("E12", "Währ-"), ("E13", "ung"))], (("B13", "asset_id"), ("D13", "source_name"))),
}


@pytest.mark.parametrize("name", list(IDENTITY_PROBES))
def test_key_identity_roles_need_positive_header_support(tmp_path, name):
    """What the mapper ADDS over the local-only run must be correct and collision-free (local wrong-key paths on renamed
    headers are pre-existing, R2-8)."""
    mutate, idents = IDENTITY_PROBES[name]
    _, local = run(tmp_path, adv.register(tmp_path, mutate), {"plans": []}, "local")
    facts, _, _ = reg_go(tmp_path, mutate, idents)
    added = [f for f in facts if (f["pointer"]["cell"], key(f)) not in {(g["pointer"]["cell"], key(g)) for g in local}]
    keys = collections.Counter((f["asset_identity"].get("asset_or_class"), f["asset_identity"].get("asset_sub_id"), f["field"]) for f in added)
    assert not wrong(added, adv.reg_truth()) and not [k for k, n in keys.items() if n > 1], wrong(added, adv.reg_truth())[:3]


@pytest.mark.parametrize("text", ["Unter-nummer", "UNr", "Sub", "Teil", "Lfd.", "Unterpos.", "SNr", "U-Nr."])
def test_undeclared_sub_number_header_variants_skip_the_block(tmp_path, text):
    """Option B must add nothing here. "UNr"/"U-Nr." also trip a PRE-EXISTING local wrong-value path (R2-8, separate
    ticket), so the assertion is on what the mapper adds over the local-only run."""
    mutate = lambda ws: ws.__setitem__("B13", text)  # noqa: E731
    _, local = run(tmp_path, adv.register(tmp_path, mutate), {"plans": []}, "local")
    facts, _, _ = reg_go(tmp_path, mutate, (("A13", "asset_id"), ("D13", "source_name")))
    added = [f for f in facts if (f["pointer"]["cell"], key(f)) not in {(g["pointer"]["cell"], key(g)) for g in local}]
    keys = collections.Counter((f["asset_identity"].get("asset_or_class"), f["asset_identity"].get("asset_sub_id"), f["field"]) for f in added)
    assert not wrong(added, adv.reg_truth()) and not [k for k, n in keys.items() if n > 1]


@pytest.mark.parametrize("name", ["year_inside_asset_name", "identity_label_two_years", "header_row_two_years", "preamble_prior_year"])
def test_period_comes_only_from_a_recognised_period_context(tmp_path, name):
    def mutate(ws):
        ws["A4"] = "Zeitraum 01.01.2024 - 31.12.2025" if name == "header_row_two_years" else "Jahresabschluss"
        if name == "year_inside_asset_name":
            ws["D14"] = "Erweiterung Halle 2019"; ws["D15"] = "CAD-Software Lizenz 2021"; ws["D16"] = "Nachrüstung 2022"
        if name == "identity_label_two_years":
            ws["D14"] = "Geschäftsjahr 2024/2025"
        if name == "preamble_prior_year":
            ws["A6"] = "Vergleich Vorjahr 2024"
    periods = {"year_inside_asset_name": [{"observation_id": f"@cell:Sheet1!D{r}", "temporal_role": "closing", "scope": "block"} for r in (14, 15, 16)],
               "identity_label_two_years": [{"observation_id": "@cell:Sheet1!D14", "temporal_role": "closing", "scope": "block"}],
               "header_row_two_years": [{"observation_id": "@cell:Sheet1!A4", "temporal_role": "closing", "scope": "sheet"}],
               "preamble_prior_year": [{"observation_id": "@cell:Sheet1!A6", "temporal_role": "closing", "scope": "sheet"}]}[name]
    _, local = run(tmp_path, adv.register(tmp_path, mutate), {"plans": []}, "local")
    facts, _, _ = reg_go(tmp_path, mutate, periods=periods)
    added = [f for f in facts if (f["pointer"]["cell"], key(f)) not in {(g["pointer"]["cell"], key(g)) for g in local}]
    assert not [f for f in added if f["asset_identity"].get("period") not in (None, "2025")]


# --- R2-5: mirrored AHK<->AfA swap on a column Spiegel with positive AfA (qa_probes/test_reqa_b1.py) ---

RECS = {8: ("Software", (1000, 200, 0, 0, 1200), (600, 100, 0, 700), (500, 400)),
        9: ("Maschinen", (3000, 500, 100, 50, 3450), (1000, 300, 80, 1220), (2230, 2000)),
        10: ("Grundstücke", (5000, 0, 0, 0, 5000), (0, 0, 0, 0), (5000, 5000))}
AHK_COLS, AFA_COLS, BV_COLS = "CDEFG", "IJKL", "NO"
AV_TRUTH = {}
for _row, (_name, _ahk, _afa, _bv) in RECS.items():
    for _cols, _vals, _fields in ((AHK_COLS, _ahk, ("ahk_open", "zugang", "abgang", "umb", "ahk_close")), (AFA_COLS, _afa, ("kum_afa_open", "total_afa_gj", "afa_abgang", "kum_afa_close")), (BV_COLS, _bv, ("bw_close", "bw_prev"))):
        for _c, _v, _f in zip(_cols, _vals, _fields):
            AV_TRUTH[("AV", f"{_c}{_row}")] = (_f, _name, "2025", Decimal(_v))


def build_av(path, movement_texts=("Anfangsbestand", "Zug.", "Abg.", "Umb.", "Endbestand"), groups=("Bruttowerte", "Wertberichtigungen", "Restbuchwerte")):
    workbook = Workbook(); ws = workbook.active; ws.title = "AV"
    ws["B1"] = "Anlagenspiegel"; ws["B2"] = "Geschäftsjahr 2025"
    for (start, end), text in zip((("C", "G"), ("I", "L"), ("N", "O")), groups):
        ws[f"{start}4"] = text; ws.merge_cells(f"{start}4:{end}4")
    ws["B5"] = "Bezeichnung"
    o, z, a, u, e = movement_texts
    for c, t in zip(AHK_COLS, (o, z, a, u, e)):
        ws[f"{c}5"] = t
    for c, t in zip(AFA_COLS, (o, z, a, e)):
        ws[f"{c}5"] = t
    ws["N5"] = "31.12.2025"; ws["O5"] = "31.12.2024"
    for c in AHK_COLS + AFA_COLS + BV_COLS:
        ws[f"{c}6"] = "EUR"
    for row, (name, ahk, afa, bv) in RECS.items():
        ws[f"B{row}"] = name
        for cols, values in ((AHK_COLS, ahk), (AFA_COLS, afa), (BV_COLS, bv)):
            for c, v in zip(cols, values):
                ws[f"{c}{row}"] = v
    workbook.save(path); return path


def av_plan(swap=False, label_bv=True):
    C = lambda ref: f"@cell:AV!{ref}"; R = lambda n: f"@row:AV:{n}"  # noqa: E731
    ahk, afa = ("cum_depreciation", "ahk") if swap else ("ahk", "cum_depreciation")
    movements = [("C5", "opening"), ("D5", "depreciation" if swap else "addition"), ("E5", "disposal"), ("F5", "transfer"), ("G5", "closing"),
                 ("I5", "opening"), ("J5", "addition" if swap else "depreciation"), ("K5", "disposal"), ("L5", "closing")]
    measures = [{"observation_id": C("C4"), "semantic_measure": ahk}, {"observation_id": C("I4"), "semantic_measure": afa}]
    spans = [{"header_observation_id": C("C4"), "first_column": 3, "last_column": 7}, {"header_observation_id": C("I4"), "first_column": 9, "last_column": 12}]
    if label_bv:
        movements += [("N5", "closing"), ("O5", "opening")]; measures.append({"observation_id": C("N4"), "semantic_measure": "book_value"})
        spans.append({"header_observation_id": C("N4"), "first_column": 14, "last_column": 15})
    return {"plans": [{"region_observation_id": "@region:AV", "table_type": "spiegel_columns", "header_rows": [R(4), R(5), R(6)],
        "record_blocks": [{"identity_row_id": R(r), "rows": [{"row_id": R(r), "header_row_id": R(5), "measure": None}], "role": "data"} for r in RECS],
        "column_spans": spans, "field_headers": [], "identity_headers": [{"observation_id": C("B5"), "identity_role": "source_name"}],
        "measure_axes": measures, "movement_axes": [{"observation_id": C(c), "movement_axis": m, "parent_header_ids": []} for c, m in movements],
        "row_roles": [], "period_contexts": [{"observation_id": C("B2"), "temporal_role": "current_period", "scope": "sheet"}],
        "unit_contexts": [{"observation_id": C("C6"), "unit": "EUR", "scope": "sheet"}]}]}


STANDARD_WORDS = ("01.01.2025", "Zugänge", "Abgänge", "Umbuchungen", "31.12.2025")


@pytest.mark.parametrize(("name", "swap", "label_bv", "words"), [
    ("control", False, True, None), ("control_standard_words", False, True, STANDARD_WORDS),
    ("mirrored_swap_bv_labelled", True, True, None), ("mirrored_swap_bv_unlabelled", True, False, None),
    ("mirrored_swap_bv_unlabelled_standard_words", True, False, STANDARD_WORDS)])
def test_positive_afa_column_spiegel_swap_never_binds_a_wrong_value(tmp_path, name, swap, label_bv, words):
    source = build_av(tmp_path / "av.xlsx", **({"movement_texts": words} if words else {}))
    _, facts = run(tmp_path, source, av_plan(swap, label_bv))
    assert not wrong(facts, AV_TRUTH)
    if name.startswith("control"):
        assert facts, "the correct plan must still bind"


def test_column_scope_periods_on_merged_and_prior_dated_headers(tmp_path):
    source = build_av(tmp_path / "av.xlsx", groups=("Bruttowerte 2025", "Wertberichtigungen 2025", "Restbuchwerte"))
    workbook = load_workbook(source); workbook["AV"]["B2"] = "Anlagenspiegel"; workbook.save(source)
    recorded = av_plan()
    recorded["plans"][0]["period_contexts"] = [{"observation_id": f"@cell:AV!{c}", "temporal_role": "current_period", "scope": "column"} for c in ("C4", "I4", "N5")] + \
        [{"observation_id": "@cell:AV!O5", "temporal_role": "prior_closing", "scope": "column"}]
    _, facts = run(tmp_path, source, recorded)
    assert not wrong(facts, AV_TRUTH)


SAP_B1 = {
    "2rows_correct_on_3header": ([(20, 16, "ahk"), (21, 17, "cum_depreciation")], None),
    "2rows_mirrored_swap": ([(20, 17, "cum_depreciation"), (21, 16, "ahk")], None),
    "2rows_shift_down_17_18": ([(20, 17, "cum_depreciation"), (21, 18, "book_value")], None),
    "one_block_out_of_order": ([(20, 16, "ahk"), (21, 17, "cum_depreciation"), (22, 18, "book_value")], [(25, 17, "cum_depreciation"), (26, 16, "ahk"), (27, 18, "book_value")]),
    "3rows_ahk_bv_swapped_measures": ([(20, 16, "book_value"), (21, 17, "cum_depreciation"), (22, 18, "ahk")], None),
}


@pytest.mark.parametrize("name", list(SAP_B1))
def test_sap_grid_measure_variants(tmp_path, name):
    first, second = SAP_B1[name]
    recorded = plan("sap_block_unlabelled"); q = recorded["plans"][0]
    q["record_blocks"] = [blk(19, first), blk(24, second or [(r + 5, h, m) for r, h, m in first]), q["record_blocks"][2]]
    _, bad = sap(tmp_path, recorded)
    assert not bad


# --- R2-1: group-aware header text (f08 "Zugänge" under the cumulative-AfA group is current-year depreciation) ---

def spiegel_field_label_plan():
    """The live f08 plan's style on the synthetic fixture: field labels on row 5, merged spans, no measure axes."""
    recorded = plan("two_year_two_sheets")
    labels = ("ahk_open", "zugang", "abgang", "umb", "ahk_close", "kum_afa_open", "total_afa_gj", "afa_abgang", "kum_afa_close", "bw_close", "bw_prev")
    for q, year in zip(recorded["plans"], (2025, 2024)):
        sheet = "Tabelle1" if year == 2025 else "Tabelle2"
        q["field_headers"] = [{"observation_id": f"@cell:{sheet}!{fixtures.spiegel_column(c, year)}5", "semantic_field": f, "parent_header_ids": []}
                              for c, f in zip(fixtures.SPIEGEL_VALUE_COLUMNS, labels)]
        q["measure_axes"] = []; q["movement_axes"] = []
    return recorded


def test_zugaenge_under_cumulative_afa_group_binds_as_current_year_depreciation(tmp_path):
    _, facts = run(tmp_path, fixtures.two_year_two_sheets(tmp_path / "s.xlsx"), spiegel_field_label_plan())
    assert not wrong(facts, adv.spiegel_truth())
    k_column = [f for f in facts if f["pointer"]["cell"][0] == ("K" if f["pointer"]["sheet"] == "Tabelle1" else "J")]
    assert len(k_column) == 6 and {f["field"] for f in k_column} == {"total_afa_gj"}


def test_zugaenge_label_under_cumulative_afa_group_is_refused(tmp_path):
    recorded = spiegel_field_label_plan()
    for q in recorded["plans"]:
        for header in q["field_headers"]:
            if header["semantic_field"] == "total_afa_gj":
                header["semantic_field"] = "zugang"
    _, facts = run(tmp_path, fixtures.two_year_two_sheets(tmp_path / "s.xlsx"), recorded)
    assert not wrong(facts, adv.spiegel_truth())


def _live(name):
    return json.loads((ROOT / "tests" / "fixtures" / "option_b_plans" / f"{name}.json").read_text())


def _replay(live_names, source, role_fix=None):
    records = [_live(name) for name in live_names]
    document_id = records[0]["document_id"]
    recorded = {"plans": [p for record in records for p in copy.deepcopy(record["plans"])]}
    for q in recorded["plans"]:
        if role_fix:
            role_fix(q)
    # Replay the recorded model output through the same verifier the pipeline uses (no HTTP, no request-size split:
    # the plans were recorded against the live run's own chunking).
    observations = acquire_xlsx(source, document_id, "a" * 64, DEFAULT_BOUNDS).observations
    mapper = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "replay", 1, 1))
    mapper.last_verification = []
    mapped = mapper._validate(document_id, observations, recorded)
    return interpret(observations, (), DEFAULT_POLICIES, mapped)[1], mapper


F08 = CORPUS / "f08_excel_spiegel_2y" / "f08_excel_spiegel_2y_01.xlsx"
F05 = CORPUS / "f05_sap_class_" "gitter" / ("f05_sap_class_" "gitter" "_02.xlsx")


@pytest.mark.skipif(not F08.exists(), reason="private corpus unavailable")
def test_live_f08_plan_no_longer_disputes_the_correct_afa_label():
    facts, mapper = _replay(["live_f08_tabelle1", "live_f08_tabelle2"], F08)
    assert not [v for v in mapper.last_verification if v["detail"].startswith("header_text_names_other_field")]
    assert any(v["detail"].startswith("empty_header_key_role") for v in mapper.last_verification), "the empty A5 class_code stays refused"

    def label_column_is_source_name(q):
        # The R2-2 prompt rule (a): a label column under an empty header is source_name. Everything else is the model's.
        q["identity_headers"] = [{**h, "identity_role": "source_name"} for h in q["identity_headers"] if h["identity_role"] in {"class_name", "balance_position"}]
    facts, _ = _replay(["live_f08_tabelle1", "live_f08_tabelle2"], F08, label_column_is_source_name)
    k_cells = [f for f in facts if f.physical.occurrence.split("!")[1][0] == ("K" if f.physical.occurrence.startswith("Tabelle1") else "J") and int(f.physical.occurrence.split("!")[1][1:]) >= 7]
    assert any(f.business.field == "total_afa_gj" for f in k_cells) and not [f for f in k_cells if f.business.field == "zugang"]


@pytest.mark.skipif(not F05.exists(), reason="private corpus unavailable")
def test_live_f05_plan_replays_without_error():
    facts, mapper = _replay(["live_f05_02"], F05)
    assert all(f.business.period == "2025" for f in facts)
