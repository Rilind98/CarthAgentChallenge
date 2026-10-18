"""Black-box acceptance tests based on independently authored documents."""
from __future__ import annotations

from decimal import Decimal
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

import openpyxl
import pytest


ROOT = Path(__file__).resolve().parents[2]
GENERATOR = ROOT / "tests" / "fixtures" / "build_cases.py"


@pytest.fixture(scope="module")
def documents(tmp_path_factory):
    directory = tmp_path_factory.mktemp("independent_sources")
    spec = importlib.util.spec_from_file_location("acceptance_fixtures", GENERATOR)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.generate(directory)
    return directory


def run_cli(source, output, config=None, template=None, timeout=90):
    command = [sys.executable, "-m", "carthagent.cli", "run", str(source), "--output", str(output)]
    if config:
        command += ["--config", str(config)]
    if template:
        command += ["--template", str(template)]
    result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, timeout=timeout)
    lines = []
    for line in result.stdout.splitlines():
        if line.startswith("{"):
            lines.append(json.loads(line))
    assert lines, f"CLI must emit explicit results: exit={result.returncode}, stderr={result.stderr}, stdout={result.stdout}"
    return result, lines


def evidence_for(source, output, **kwargs):
    completed, records = run_cli(source, output, **kwargs)
    assert len(records) == 1
    record = records[0]
    assert record["status"] in {"complete", "partial"}, record
    assert completed.returncode == 0, completed.stderr
    return record, json.loads(Path(record["evidence_path"]).read_text())


def amount(fact):
    return Decimal(fact["amount"]["normalized_base"])


def at_cell(data, address, sheet=None):
    return [f for f in data["facts"] if f["pointer"].get("cell") == address and (sheet is None or f["pointer"].get("sheet") == sheet)]


def assert_pair(record, data):
    directory = Path(record["output_dir"])
    manifest = json.loads(Path(record["manifest_path"]).read_text())
    assert (directory / ".complete").read_text().strip() == data["snapshot"]["digest"]
    assert manifest["snapshot_digest"] == data["snapshot"]["digest"]
    for name, digest in manifest["files"].items():
        assert hashlib.sha256((directory / name).read_bytes()).hexdigest() == digest
    assert manifest["readback_verified"] is True
    assert data["contract_status"] != "official_scorer_verified"


def test_positive_xlsx_requires_correct_fields_and_row_identity(documents, tmp_path):
    record, data = evidence_for(documents / "clear.xlsx", tmp_path)
    expected = {
        "C3": ("000101", "ahk_open", "1000"), "D3": ("000101", "zugang", "120"),
        "E3": ("000101", "abgang", "20"), "F3": ("000101", "ahk_close", "1100"),
        "G3": ("000101", "bw_close", "800"), "C4": ("000202", "ahk_open", "2000"),
        "D4": ("000202", "zugang", "340"), "E4": ("000202", "abgang", "40"),
        "F4": ("000202", "ahk_close", "2300"), "G4": ("000202", "bw_close", "1600"),
    }
    assert len(data["facts"]) == 10, "Every monetary cell in this simple native table has an unambiguous header and row"
    assert len(data["facts"]) <= 10, "Identifiers and dates must not become monetary facts"
    for fact in data["facts"]:
        cell = fact["pointer"]["cell"]
        assert cell in expected, fact
        asset, field, value = expected[cell]
        assert fact["asset_identity"]["asset_or_class"] == asset
        assert fact["asset_identity"]["book"] is None, "Buchwert is a value, not an accounting-book identity"
        assert fact["asset_identity"]["period"] == "2025", "Opening balances and movements belong to the same reporting period"
        assert fact["field"] == field
        assert amount(fact) == Decimal(value)
        assert fact["pointer"]["sheet"] == "Anlagen"
        assert fact["dependency_ids"]
    assert_pair(record, data)


def test_stacked_schedules_keep_local_headers_periods_and_units(documents, tmp_path):
    _, data = evidence_for(documents / "stacked_schedules.xlsx", tmp_path)
    expected = {}
    first_fields = ["ahk_open", "zugang", "abgang", "ahk_close", "bw_close"]
    for row, asset, values in [(3, "000101", [1000, 120, 20, 1100, 800]), (4, "000202", [2000, 340, 40, 2300, 1600])]:
        for column, field, value in zip("CDEFG", first_fields, values):
            expected[f"{column}{row}"] = (asset, "2024", field, Decimal(value))
    for column, field, value in zip("CDEFG", ["ahk_open", "abgang", "zugang", "ahk_close", "bw_close"], [10000, 1000, 2000, 11000, 8000]):
        expected[f"{column}9"] = ("000303", "2025", field, Decimal(value))
    actual = {fact["pointer"].get("cell"): (fact["asset_identity"]["asset_or_class"], fact["asset_identity"]["period"], fact["field"], amount(fact)) for fact in data["facts"]}
    assert len(data["facts"]) == 15, "Both explicit tables should retain all five monetary columns for each asset"
    assert actual == expected, "Later table headers, years and TEUR scaling must not leak into the earlier EUR table"


@pytest.mark.parametrize("label,prefix", [("equals", "="), ("plus", "+"), ("minus", "-"), ("at", "@")])
def test_formula_prefixed_source_sheet_locator_is_escaped(documents, tmp_path, label, prefix):
    record, data = evidence_for(documents / f"formula_sheet_{label}.xlsx", tmp_path)
    assert len(data["facts"]) == 10
    assert all(fact["pointer"]["sheet"] == prefix + "2+2" for fact in data["facts"])
    workbook = openpyxl.load_workbook(record["review_path"], data_only=False)
    try:
        locator_cells = [cell for sheet in workbook for row in sheet for cell in row if isinstance(cell.value, str) and prefix + "2+2!" in cell.value]
        assert locator_cells, "The human review workbook must retain source locators"
        assert all(cell.data_type != "f" and cell.value.startswith("'") for cell in locator_cells), "An untrusted sheet name must never become an executable workbook formula"
    finally:
        workbook.close()


@pytest.mark.parametrize("name", ["clear.pdf", "rotated_90.pdf", "rotated_180.pdf", "rotated_270.pdf", "cropped.pdf"])
def test_pdf_positive_and_original_page_pointer(documents, tmp_path, name):
    import pymupdf
    record, data = evidence_for(documents / name, tmp_path)
    assert len(data["facts"]) >= 6, (name, data["diagnostics"])
    allowed = {"000101": {"ahk_open": 1000, "zugang": 120, "abgang": 20, "ahk_close": 1100, "bw_close": 800},
               "000202": {"ahk_open": 2000, "zugang": 340, "abgang": 40, "ahk_close": 2300, "bw_close": 1600}}
    with pymupdf.open(documents / name) as doc:
        for fact in data["facts"]:
            identity = fact["asset_identity"]["asset_or_class"]
            assert identity in allowed
            assert fact["field"] in allowed[identity]
            assert amount(fact) == Decimal(allowed[identity][fact["field"]])
            pointer = fact["pointer"]
            assert pointer["kind"] == "pdf"
            assert len(pointer["bbox"]) == 4
            page = doc[pointer["page_index"]]
            crop = page.get_textbox(pymupdf.Rect(pointer["bbox"]))
            # A second source read must find the actual lexeme in this exact box.
            assert fact["amount"]["raw"].replace(" ", "") in crop.replace(" ", "").replace("\n", ""), (pointer, crop)
    assert_pair(record, data)


def test_repeated_pdf_amounts_require_the_correct_record_and_column_pointer(documents, tmp_path):
    import pymupdf
    _, data = evidence_for(documents / "repeated_values.pdf", tmp_path)
    assert len(data["facts"]) >= 6
    fields = ["ahk_open", "zugang", "abgang", "ahk_close", "bw_close"]
    with pymupdf.open(documents / "repeated_values.pdf") as doc:
        words = doc[0].get_text("words")
    expected = {}
    for identity in ("000101", "000202"):
        label = next(w for w in words if w[4] == identity)
        center_y = (label[1] + label[3]) / 2
        values = sorted((w for w in words if w[4] == "1.000,00" and abs((w[1] + w[3]) / 2 - center_y) < 2), key=lambda w: w[0])
        assert len(values) == 5
        expected.update({(identity, field): value[:4] for field, value in zip(fields, values)})
    for fact in data["facts"]:
        key = (fact["asset_identity"]["asset_or_class"], fact["field"])
        assert key in expected
        assert amount(fact) == 1000
        x0, y0, x1, y1 = expected[key]
        box = fact["pointer"]["bbox"]
        assert x0 <= (box[0] + box[2]) / 2 <= x1
        assert y0 <= (box[1] + box[3]) / 2 <= y1


@pytest.mark.parametrize("name", ["ambiguous.xlsx", "ambiguous.pdf"])
def test_sparse_unknown_owner_is_not_guessed(documents, tmp_path, name):
    _, data = evidence_for(documents / name, tmp_path)
    assert not any(amount(f) == Decimal("137.42") for f in data["facts"])
    assert not any(f["asset_identity"].get("asset_or_class") in {"X", "Y"} for f in data["facts"]), "Unresolved sparse records cannot acquire invented zero balances"
    assert data["abstentions"] or any(c["disposition"] == "unresolved_unsupported" for c in data["coverage"])


@pytest.mark.parametrize("name", ["measure_blocks.xlsx", "measure_blocks.pdf"])
def test_measure_rows_and_balance_columns_form_one_asset_block(documents, tmp_path, name):
    _, data = evidence_for(documents / name, tmp_path)
    fields = [("ahk_open", "ahk_close"), ("kum_afa_open", "kum_afa_close"), ("bw_prev", "bw_close")]
    values = [(1000, 1100), (400, 500), (600, 600)]
    expected = {(asset, field): Decimal(value) for asset in ("BL-001", "BL-002")
                for row_fields, row_values in zip(fields, values) for field, value in zip(row_fields, row_values)}
    assert len(data["facts"]) == 12, "All twelve printed balances have explicit row-measure, column-date and asset-block evidence"
    actual = {(f["asset_identity"]["asset_or_class"], f["field"]): amount(f) for f in data["facts"]}
    assert actual == expected
    for fact in data["facts"]:
        assert fact["asset_identity"]["period"] == "2025"
        block = 0 if fact["asset_identity"]["asset_or_class"] == "BL-001" else 1
        measure = next(i for i, pair in enumerate(fields) if fact["field"] in pair)
        column = fields[measure].index(fact["field"])
        if name.endswith(".xlsx"):
            assert fact["pointer"]["cell"] == f"{'CD'[column]}{3 + block * 3 + measure}"
            assert fact["pointer"]["sheet"] == "Block schedule"
        else:
            import pymupdf
            with pymupdf.open(documents / name) as doc:
                words = doc[0].get_text("words")
                baseline = doc[0].rect.height - (490 - (block * 3 + measure) * 26)
            matches = [w for w in words if abs(w[3] - baseline) < 5 and 450 + column * 200 <= w[0] < 600 + column * 200]
            assert len(matches) == 1
            pointer = fact["pointer"]
            assert pointer["page_index"] == 0
            box = pointer["bbox"]
            x0, y0, x1, y1 = matches[0][:4]
            assert x0 <= (box[0] + box[2]) / 2 <= x1
            assert y0 <= (box[1] + box[3]) / 2 <= y1


def test_confirmed_empty_and_na_are_different(documents, tmp_path):
    _, data = evidence_for(documents / "blank_and_na.xlsx", tmp_path)
    empty = at_cell(data, "D3")
    assert len(empty) == 1 and amount(empty[0]) == 0, "A true native amount-grid blank must be represented as zero"
    assert not at_cell(data, "D4"), "N/A must not silently become zero"


def test_acquisition_dates_do_not_supply_a_missing_reporting_period(documents, tmp_path):
    record, data = evidence_for(documents / "missing_reporting_period.xlsx", tmp_path)
    assert record["status"] == "partial"
    assert not data["facts"], "A financial fact without a reporting period must remain an unresolved candidate"
    assert data["abstentions"]


def test_zero_subnumber_is_preserved_and_distinguishes_same_asset_id(documents, tmp_path):
    _, data = evidence_for(documents / "asset_subnumbers.xlsx", tmp_path)
    assert len(data["facts"]) == 10
    for fact in data["facts"]:
        row = fact["pointer"]["cell"][-1]
        assert row in {"3", "4"}
        assert fact["asset_identity"]["asset_or_class"] == "000101"
        assert fact["asset_identity"]["asset_sub_id"] == {"3": "0", "4": "1"}[row]
        assert fact["asset_identity"]["period"] == "2025"
    assert not at_cell(data, "H3") and not at_cell(data, "H4"), "Asset subnumbers are identifiers, not amounts"


@pytest.mark.parametrize("name", ["later_print_date.xlsx", "later_print_date.pdf"])
def test_print_date_does_not_override_reporting_period(documents, tmp_path, name):
    _, data = evidence_for(documents / name, tmp_path)
    assert len(data["facts"]) >= 6
    assert all(f["asset_identity"]["period"] == "2025" for f in data["facts"])


def test_missing_formula_cache_never_recalculated(documents, tmp_path):
    _, data = evidence_for(documents / "formula_missing_cache.xlsx", tmp_path)
    assert not at_cell(data, "F3")
    assert any("formula" in json.dumps(item).lower() or "cache" in json.dumps(item).lower() for item in data["observations"] + data["coverage"] + data["abstentions"])


def test_hidden_scope_remains_explicit(documents, tmp_path):
    record, data = evidence_for(documents / "hidden_content.xlsx", tmp_path)
    assert record["status"] == "partial"
    assert "Hidden schedule" in json.dumps(data)
    assert "hidden" in json.dumps(data["coverage"]).lower() or "hidden" in json.dumps(data["diagnostics"]).lower()


def test_source_arithmetic_error_preserved(documents, tmp_path):
    _, data = evidence_for(documents / "source_error.xlsx", tmp_path)
    closing = at_cell(data, "F3")
    assert closing and amount(closing[0]) == 100
    assert not any(amount(f) == 160 for f in data["facts"])
    assert any(f["source_anomaly"] != "none" for f in data["facts"]) or any("inconsisten" in str(x).lower() for x in data["diagnostics"])


def test_signs_and_scale_replay(documents, tmp_path):
    _, signs = evidence_for(documents / "signs.xlsx", tmp_path / "signs")
    for cell, expected in {"C3": "1234.56", "E3": "-234.56", "G3": "-700"}.items():
        facts = at_cell(signs, cell)
        assert len(facts) == 1 and Decimal(facts[0]["amount"]["source_signed"]) == Decimal(expected)
    _, scaled = evidence_for(documents / "scaled.xlsx", tmp_path / "scale")
    opening = at_cell(scaled, "C3")
    assert opening and amount(opening[0]) == 12000
    assert Decimal(opening[0]["amount"]["scale"]) == 1000


def test_duplicate_identity_context_is_not_collapsed(documents, tmp_path):
    _, data = evidence_for(documents / "duplicate_ids.xlsx", tmp_path)
    first, second = at_cell(data, "D3", "Anlagen"), at_cell(data, "D3", "Entity B")
    assert first and second
    assert amount(first[0]) == 120 and amount(second[0]) == 789
    assert first[0]["id"] != second[0]["id"]


def test_typed_dates_and_useful_life_are_not_monetary_values(documents, tmp_path):
    _, data = evidence_for(documents / "typed_metadata.xlsx", tmp_path)
    assert len(data["facts"]) >= 6
    assert not any(f["pointer"].get("cell") in {"H3", "H4", "I3", "I4"} for f in data["facts"])
    assert "2023-07-15" in json.dumps(data["observations"])


def test_corrupt_input_does_not_stop_batch_and_gold_is_excluded(documents, tmp_path):
    source = tmp_path / "batch"
    source.mkdir()
    for name in ["clear.xlsx", "broken.pdf", "broken.xlsx", "unsupported.docx", "poison.key.json"]:
        shutil.copy(documents / name, source / name)
    shutil.copytree(documents / "_generator", source / "_generator")
    _, results = run_cli(source, tmp_path / "out")
    assert len(results) == 4, results
    statuses = {r["document_id"]: r["status"] for r in results}
    assert any(r["document_id"].startswith("clear-") and r["status"] in {"partial", "complete"} for r in results)
    assert sum(r["status"] == "failed" for r in results) == 2
    assert "unsupported_input" in statuses.values()
    assert not any("forbidden" in r["document_id"] or "poison" in r["document_id"] for r in results)


def test_different_sources_with_same_basename_do_not_collide(documents, tmp_path):
    source = tmp_path / "in"
    (source / "a").mkdir(parents=True)
    (source / "b").mkdir()
    shutil.copy(documents / "clear.xlsx", source / "a" / "same.xlsx")
    shutil.copy(documents / "scaled.xlsx", source / "b" / "same.xlsx")
    _, results = run_cli(source, tmp_path / "out")
    assert len(results) == 2
    assert all(r["status"] in {"complete", "partial"} for r in results), results
    assert len({r["output_dir"] for r in results}) == 2


def test_repeat_run_verifies_and_reuses_consistent_pair(documents, tmp_path):
    first, a = evidence_for(documents / "clear.xlsx", tmp_path)
    second, b = evidence_for(documents / "clear.xlsx", tmp_path)
    assert a["snapshot"]["digest"] == b["snapshot"]["digest"]
    assert first["output_dir"] == second["output_dir"]
    assert_pair(second, b)


def test_modified_artifact_is_not_trusted_on_resume(documents, tmp_path):
    first, _ = evidence_for(documents / "clear.xlsx", tmp_path)
    path = Path(first["evidence_path"])
    altered = json.loads(path.read_text())
    altered["facts"][0]["amount"]["normalized_base"] = "999999999"
    path.write_text(json.dumps(altered))
    _, results = run_cli(documents / "clear.xlsx", tmp_path)
    if results[0]["status"] in {"complete", "partial"}:
        repaired = json.loads(Path(results[0]["evidence_path"]).read_text())
        assert_pair(results[0], repaired)
        assert not any(amount(f) == 999999999 for f in repaired["facts"])
    else:
        assert results[0]["status"] == "failed"
        assert any(word in str(results[0]["diagnostics"]).lower() for word in ("hash", "integrity", "publication", "readback", "existing"))


def test_identical_bytes_at_different_names_keep_document_provenance(documents, tmp_path):
    source = tmp_path / "in"
    source.mkdir()
    for name in ("alpha.xlsx", "beta.xlsx"):
        shutil.copy(documents / "clear.xlsx", source / name)
    _, results = run_cli(source, tmp_path / "out")
    assert len(results) == 2
    for result in results:
        assert result["status"] in {"complete", "partial"}, result
        data = json.loads(Path(result["evidence_path"]).read_text())
        assert data["document"]["document_id"] == result["document_id"]
        assert all(f["document_id"] == result["document_id"] for f in data["facts"])


def test_outputs_inside_input_are_not_reingested(documents, tmp_path):
    source = tmp_path / "batch"
    source.mkdir()
    shutil.copy(documents / "clear.xlsx", source / "clear.xlsx")
    output = source / "results"
    run_cli(source, output)
    _, second = run_cli(source, output)
    assert len(second) == 1, "Published review workbooks must not become new inputs"


def test_grouped_headers_and_section_rows_have_distinct_meanings(documents, tmp_path):
    _, data = evidence_for(documents / "hierarchical_headers.xlsx", tmp_path)
    assert not any(f["pointer"].get("cell", "").endswith("7") for f in data["facts"]), "A section heading has no monetary slots to fill with zero"
    expected = {"C8": ("ahk_open", 1000), "D8": ("zugang", 50), "E8": ("abgang", 20),
                "G8": ("ahk_close", 1030), "H8": ("kum_afa_open", 500), "J8": ("afa_abgang", 10),
                "L8": ("kum_afa_close", 590), "M8": ("bw_close", 440), "N8": ("bw_prev", 500)}
    for cell, (field, value) in expected.items():
        fact = at_cell(data, cell)
        assert len(fact) == 1, (cell, data["abstentions"][:2])
        assert fact[0]["field"] == field
        assert amount(fact[0]) == value
        assert fact[0]["asset_identity"]["asset_id"] == "077001"
        assert fact[0]["asset_identity"]["period"] == "2025"
    blank = at_cell(data, "D9")
    assert len(blank) == 1 and amount(blank[0]) == 0


def test_parallel_year_columns_use_local_period_evidence(documents, tmp_path):
    _, data = evidence_for(documents / "parallel_years.xlsx", tmp_path)
    for cell, year, value in [("C3", "2024", 111), ("D3", "2025", 222)]:
        facts = at_cell(data, cell)
        assert len(facts) == 1
        assert facts[0]["field"] == "zugang"
        assert amount(facts[0]) == value
        assert facts[0]["asset_identity"]["period"] == year


def test_column_unit_overrides_document_unit_only_for_that_column(documents, tmp_path):
    _, data = evidence_for(documents / "local_unit_override.xlsx", tmp_path)
    for cell, value in [("C3", 1000), ("D3", 2000)]:
        facts = at_cell(data, cell)
        assert len(facts) == 1
        assert amount(facts[0]) == value


def test_pdf_column_unit_does_not_scale_unrelated_columns(documents, tmp_path):
    _, data = evidence_for(documents / "local_unit_override.pdf", tmp_path)
    expected = {"ahk_open": 1000, "zugang": 2000, "abgang": 20, "ahk_close": 2980, "bw_close": 800}
    assert len(data["facts"]) >= 3
    for fact in data["facts"]:
        assert fact["asset_identity"]["asset_or_class"] == "000101"
        assert fact["field"] in expected
        assert amount(fact) == expected[fact["field"]], fact
    # The clearly labelled column is recoverable; blanket abstention is not a fix.
    assert any(f["field"] == "zugang" for f in data["facts"])


@pytest.mark.parametrize("name", ["scanned.pdf", "scanned_rotated_90.pdf", "scanned_rotated_180.pdf", "scanned_rotated_270.pdf", "mixed_native_scan.pdf"])
def test_scan_recovers_values_and_preserves_original_coordinates(documents, tmp_path, name):
    import pymupdf
    if not shutil.which("tesseract"):
        pytest.skip("Tesseract is required for local OCR acceptance")
    _, data = evidence_for(documents / name, tmp_path)
    assert len(data["facts"]) >= 6, (name, data["diagnostics"])
    with pymupdf.open(documents / "clear.pdf") as original:
        words = original[0].get_text("words")
    allowed = {"000101": {"ahk_open": 1000, "zugang": 120, "abgang": 20, "ahk_close": 1100, "bw_close": 800},
               "000202": {"ahk_open": 2000, "zugang": 340, "abgang": 40, "ahk_close": 2300, "bw_close": 1600}}
    for fact in data["facts"]:
        value = amount(fact)
        identity = fact["asset_identity"]["asset_or_class"]
        assert identity in allowed, (name, fact)
        assert fact["field"] in allowed[identity], (name, fact)
        assert value == Decimal(allowed[identity][fact["field"]]), (name, fact)
        display = f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
        candidates = [w for w in words if w[4] == display]
        assert candidates, (name, fact)
        box = fact["pointer"]["bbox"]
        x, y = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        assert any(w[0] - 4 <= x <= w[2] + 4 and w[1] - 4 <= y <= w[3] + 4 for w in candidates), (name, box, candidates)


def test_large_styled_range_has_bounded_truthful_outcome(documents, tmp_path):
    _, results = run_cli(documents / "giant_styled_range.xlsx", tmp_path, timeout=20)
    assert results[0]["status"] in {"partial", "failed"}
    assert any("limit" in str(d).lower() or "range" in str(d).lower() for d in results[0]["diagnostics"])


def test_no_formula_injection_in_provisional_output(documents, tmp_path):
    record, data = evidence_for(documents / "injection.xlsx", tmp_path)
    assert not any(amount(f) == 999999 for f in data["facts"])
    wb = openpyxl.load_workbook(record["review_path"], data_only=False)
    for ws in wb:
        for row in ws:
            assert not any(c.data_type == "f" for c in row), "Input text must not become an executable Excel formula"


def test_template_preserves_unrelated_content_and_maps_printed_values(documents, tmp_path):
    template_text = os.environ.get("CARTHAGENT_TEST_TEMPLATE")
    if not template_text:
        pytest.skip("Private template not distributed; set CARTHAGENT_TEST_TEMPLATE to test it locally")
    template = Path(template_text)
    source_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    record, data = evidence_for(documents / "clear.xlsx", tmp_path, template=template)
    original = openpyxl.load_workbook(template, data_only=False)
    actual = openpyxl.load_workbook(record["review_path"], data_only=False)
    assert original.sheetnames == actual.sheetnames, "Preserve all official template tabs and their order"
    assert hashlib.sha256(template.read_bytes()).hexdigest() == source_hash
    for sheet in ["Zugänge", "Abgänge", "Belege", "Summen- und Saldenliste"]:
        assert original[sheet].max_row == actual[sheet].max_row
        for row in original[sheet]:
            for cell in row:
                copied = actual[sheet][cell.coordinate]
                if isinstance(cell.value, openpyxl.worksheet.formula.ArrayFormula):
                    assert copied.value.text == cell.value.text
                else:
                    assert copied.value == cell.value
    register = actual["Anlagenverzeichnis"]
    by_asset = {str(register.cell(r, 1).value): r for r in range(2, 502) if register.cell(r, 1).value is not None}
    assert {"000101", "000202"} <= set(by_asset)
    r = by_asset["000101"]
    assert Decimal(str(register[f"W{r}"].value)) == 1000
    assert Decimal(str(register[f"X{r}"].value)) == 120
    assert Decimal(str(register[f"AA{r}"].value)) == 1100
    assert register[f"AA{r}"].data_type != "f", "Printed closing balance takes precedence over template calculation"
    assert_pair(record, data)


def test_template_never_fills_an_unavailable_printed_balance(documents, tmp_path):
    template = os.environ.get("CARTHAGENT_TEST_TEMPLATE")
    if not template:
        pytest.skip("Private template unavailable")
    record, data = evidence_for(documents / "formula_missing_cache.xlsx", tmp_path, template=template)
    wb = openpyxl.load_workbook(record["review_path"], data_only=False)
    ws = wb["Anlagenverzeichnis"]
    row = next(r for r in range(2, 502) if str(ws.cell(r, 1).value) == "000101")
    assert ws.cell(row, 27).value is None, "An unavailable source closing balance cannot be replaced by the template formula"
    assert record["status"] == "partial"
    assert_pair(record, data)

    # An absent source column also cannot become an unlabelled calculated fact.
    source = tmp_path / "no_closing_column.xlsx"
    native = openpyxl.load_workbook(documents / "clear.xlsx")
    for cell in ("F2", "F3", "F4"):
        native.active[cell] = None
    native.save(source)
    record, data = evidence_for(source, tmp_path / "absent-output", template=template)
    ws = openpyxl.load_workbook(record["review_path"], data_only=False)["Anlagenverzeichnis"]
    row = next(r for r in range(2, 502) if str(ws.cell(r, 1).value) == "000101")
    assert ws.cell(row, 27).value is None, "A derived balance needs a separate, explicitly enabled lineage contract"
    assert not any(f["field"] == "ahk_close" for f in data["facts"])


def test_total_depreciation_does_not_become_planned_depreciation(documents, tmp_path):
    template = os.environ.get("CARTHAGENT_TEST_TEMPLATE")
    if not template:
        pytest.skip("Private template unavailable")
    source = tmp_path / "total_depreciation.xlsx"
    wb = openpyxl.load_workbook(documents / "clear.xlsx")
    wb.active["H2"] = "AfA GJ"
    wb.active["H3"] = 95
    wb.save(source)
    record, data = evidence_for(source, tmp_path / "out", template=template)
    assert any(f["field"] == "total_afa_gj" and amount(f) == 95 for f in data["facts"])
    out = openpyxl.load_workbook(record["review_path"], data_only=False)["Anlagenverzeichnis"]
    row = next(r for r in range(2, 502) if str(out.cell(r, 1).value) == "000101")
    assert out.cell(row, 29).value is None
    assert record["status"] == "partial"
    assert any(b["status"] != "bound" for b in data["output_bindings"])


def test_template_grid_overflow_is_explicit_and_keeps_all_source_facts(tmp_path):
    template = os.environ.get("CARTHAGENT_TEST_TEMPLATE")
    if not template:
        pytest.skip("Private template unavailable")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Anlagenspiegel 2025 EUR"])
    ws.append(["Anlagenklasse", "Bezeichnung", "AHK 01.01.2025", "Zugänge", "Abgänge", "AHK 31.12.2025", "Buchwert 31.12.2025"])
    for index in range(21):
        ws.append([str(1000 + index), f"Class {index}", 1000, 10, 0, 1010, 800])
    source = tmp_path / "classes.xlsx"
    wb.save(source)
    record, data = evidence_for(source, tmp_path / "out", template=template)
    assert len(data["facts"]) == 105
    statuses = [b["status"] for b in data["output_bindings"]]
    if "projection_overflow" in statuses:
        assert record["status"] == "partial"
        assert statuses.count("projection_overflow") == 5
    else:
        # A future adapter may safely expand every dependent template range.
        assert all(status == "bound" for status in statuses)
        assert len(data["output_bindings"]) == 105
    assert_pair(record, data)
