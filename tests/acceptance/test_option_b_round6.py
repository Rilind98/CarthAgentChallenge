"""Round 6: regressions seen in live stage-1 on 935f331, replayed offline from the recorded live plans.

R6-1 f00_demo_01 local-only must equal the 779e1a4 facts (no new abstentions).
R6-2 the live f05 plan (F16/G16 labelled bew_ahk under AHK) must bind through run_input with 0 wrong vs the sheet.
R6-3 the live f08 plans must bind at least 25 main-table cells with 0 wrong and 0 reconciliation rows.
"""

from __future__ import annotations

from collections import Counter
from decimal import Decimal
import importlib.util
import json
import re
from pathlib import Path

import pytest

from carthagent.mapper import MapperConfig
from carthagent.pipeline import RunConfig, run_input

ROOT = Path(__file__).resolve().parents[2]
CORPUS = ROOT / "hackathon-local" / "ElabHackathon" / "A3-Hackathon-Corpus"
PLANS = ROOT / "tests" / "fixtures" / "option_b_plans"
F00 = CORPUS / "f00_demo" / "f00_demo_01.xlsx"
F05 = CORPUS / ("f05_sap_class_" "gitter") / ("f05_sap_class_" "gitter_02.xlsx")


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


stub = _load(ROOT / "tests" / "fixtures" / "option_b_stub.py", "r6_stub")


def replay(source: Path, name: str, out: Path) -> dict:
    """Run the real pipeline with both recorded live requests served by the loopback stub (plan IDs untouched)."""
    plans = [plan for n in (1, 2) for plan in json.loads((PLANS / f"live_r5_{name}_req{n}.json").read_text())["plans"]]
    with stub.recorded_plan_server({"plans": plans}) as endpoint:
        result = run_input(source, out, RunConfig(mapper=MapperConfig(True, endpoint, "replay", 2, 256)))
    assert not [d for d in result.diagnostics if "mapper" in d and "error" in d.lower()], result.diagnostics
    return json.loads(Path(result.evidence_path).read_text())


F00_EXPECTED = ROOT / "hackathon-local" / "test-fixtures" / "f00_demo_01_local_779e1a4.json"


@pytest.mark.skipif(not (F00.exists() and F00_EXPECTED.exists()), reason="private corpus or its recorded local facts unavailable")
def test_r6_1_f00_local_only_equals_779e1a4():
    from carthagent.acquire_xlsx import acquire_xlsx
    from carthagent.policies import DEFAULT_BOUNDS, DEFAULT_POLICIES
    from carthagent.semantics import interpret
    expected = {tuple(row) for row in json.loads(F00_EXPECTED.read_text())["facts"]}
    observations = acquire_xlsx(F00, "d", "a" * 64, DEFAULT_BOUNDS).observations
    _, facts, abstentions, _ = interpret(observations, (), DEFAULT_POLICIES)
    got = {(f.physical.occurrence, f.business.field, str(f.amount.normalized_value), f.business.asset_or_class, f.business.period) for f in facts}
    assert len(expected) == 415
    assert got == expected, (sorted(expected - got)[:5], sorted(got - expected)[:5])
    assert not abstentions, Counter(a.reason for a in abstentions)


@pytest.mark.skipif(not F05.exists(), reason="private corpus unavailable")
def test_r6_2_live_f05_plan_binds_through_run_input_with_zero_wrong(tmp_path):
    qa4 = _load(Path(__file__).with_name("test_option_b_round5_qa4.py"), "r6_qa4")
    evidence = replay(F05, "f05_sap_class_gitter_02", tmp_path / "out")
    truth = qa4.sheet_truth(F05)
    got = {f["pointer"]["cell"]: (f["field"], f["asset_identity"].get("class_code"), Decimal(str(f["amount"]["normalized_base"])), f["asset_identity"].get("period"), f["amount"]["unit"])
           for f in evidence["facts"] if f["pointer"]["sheet"] == "Anlagengitter"}
    assert got, "the live plan binds"
    wrong = [(cell, value) for cell, value in got.items() if truth.get(cell) != value[:3] or value[3:] != ("2025", "EUR")]
    assert not wrong, wrong[:5]
    assert got["B20"] == ("ahk_open", "1000", Decimal("208034.23"), "2025", "EUR")
    assert got["B21"][:2] == ("kum_afa_open", "1000") and got["H22"][:2] == ("bw_close", "1000")
    assert not {cell for cell in got if cell[0] in "FG" and int(cell[1:]) % 5 == 0}, "Nachaktivg / Invest.Foerderg (labelled bew_ahk under AHK) abstain"


F08 = CORPUS / "f08_excel_spiegel_2y" / "f08_excel_spiegel_2y_01.xlsx"
SPIEGEL_FIELDS = ("ahk_open", "zugang", "abgang", "umb", "ahk_close", "kum_afa_open", "total_afa_gj", "afa_abgang", "kum_afa_close", "bw_close", "bw_prev")
F08_COLUMNS = {"Tabelle1": dict(zip("CDEFHJKLMOP", SPIEGEL_FIELDS)), "Tabelle2": dict(zip("BCDEGIJKLNO", SPIEGEL_FIELDS))}


def f08_sheet_truth(path):
    """Truth read from the sheet: numbered class rows ("2." in A or leading the name), field from the column, empty = 0."""
    from openpyxl import load_workbook
    truth = {}
    workbook = load_workbook(path, data_only=True)
    for sheet, (name_column, period) in {"Tabelle1": ("B", "2025"), "Tabelle2": ("A", "2024")}.items():
        ws = workbook[sheet]
        for row in range(7, 27):
            if not re.match(r"\d+\.", str(ws[f"A{row}"].value or "").strip()) or not ws[f"{name_column}{row}"].value:
                continue
            for column, field in F08_COLUMNS[sheet].items():
                value = ws[f"{column}{row}"].value
                truth[(sheet, f"{column}{row}")] = (field, Decimal(str(value)) if isinstance(value, (int, float)) else Decimal(0), period, "EUR")
    return truth


@pytest.mark.skipif(not F08.exists(), reason="private corpus unavailable")
def test_r6_3_live_f08_plans_bind_main_table_with_units_and_zero_wrong(tmp_path):
    evidence = replay(F08, "f08_excel_spiegel_2y_01", tmp_path / "out")
    truth = f08_sheet_truth(F08)
    got = {(f["pointer"]["sheet"], f["pointer"]["cell"]): (f["field"], Decimal(str(f["amount"]["normalized_base"])), f["asset_identity"].get("period"), f["amount"]["unit"]) for f in evidence["facts"]}
    assert not [key for key in got if int(key[1][1:]) >= 27], "reconciliation rows below the table never bind"
    wrong = [(key, value, truth.get(key)) for key, value in got.items() if truth.get(key) != value]
    assert not wrong, wrong[:5]
    correct = sum(1 for key, value in got.items() if truth.get(key) == value)
    assert correct >= 25, (correct, Counter(a["reason"] for a in evidence["abstentions"]).most_common(4))
    assert got[("Tabelle1", "C10")] == ("ahk_open", Decimal("144589.68"), "2025", "EUR")


# --- Item 4: K2 ends a table only on a one-cell title whose year CONTRADICTS the sheet period ---

reqa = _load(Path(__file__).with_name("test_option_b_reqa.py"), "r6_reqa")
adv = _load(Path(__file__).with_name("test_option_b_adversarial.py"), "r6_adv")
BGA = ("BGA", (800, 100, 0, 0, 900), (300, 100, 0, 400), (500, 500))
SPIEGEL_AV = (("ahk_open", "zugang", "abgang", "umb", "ahk_close"), ("kum_afa_open", "total_afa_gj", "afa_abgang", "kum_afa_close"), ("bw_close", "bw_prev"))


def av_with_title(path, title):
    from openpyxl import load_workbook
    reqa.build_av(path)
    workbook = load_workbook(path); ws = workbook["AV"]
    ws["B11"] = title; ws["B12"] = BGA[0]
    for cols, values in zip((reqa.AHK_COLS, reqa.AFA_COLS, reqa.BV_COLS), BGA[1:]):
        for column, value in zip(cols, values):
            ws[f"{column}12"] = value
    workbook.save(path); return path


def av_plans(form):
    if form == "local":
        return {"plans": []}
    recorded = reqa.av_plan(); plan = recorded["plans"][0]
    plan["record_blocks"].append({"identity_row_id": "@row:AV:12", "rows": [{"row_id": "@row:AV:12", "header_row_id": "@row:AV:5", "measure": None}], "role": "data"})
    if form == "v11":
        plan.pop("record_blocks"); plan.pop("column_spans")
        plan["block_pattern"] = [{"rows": [{"ordinal": 0, "header_row_id": "@row:AV:5", "measure": "none"}], "role": "data", "row_count": 1, "applies_to": "all_blocks_of_same_shape"}]
    return recorded


@pytest.mark.parametrize("form", ["local", "legacy", "v11"])
@pytest.mark.parametrize(("title", "row_12_binds"), [("Sachanlagen 2025", True), ("Geschäftsjahr 2025", True), ("Stand 31.12.2025", True), ("Entwicklung Vorjahr 2024", False), ("Vorjahr 2024", False)])
def test_item4_k2_same_year_title_keeps_the_table(tmp_path, form, title, row_12_binds):
    source = av_with_title(tmp_path / "av.xlsx", title)
    truth = dict(reqa.AV_TRUTH)
    for cols, values, fields in zip((reqa.AHK_COLS, reqa.AFA_COLS, reqa.BV_COLS), BGA[1:], SPIEGEL_AV):
        for column, value, field in zip(cols, values, fields):
            truth[("AV", f"{column}12")] = (field, "BGA", "2025", Decimal(value))
    _, facts = adv.run(tmp_path, source, av_plans(form), form)
    assert not adv.wrong(facts, truth)
    # Local rules read only the two book-value columns of this abbreviated-header fixture (2 per record, pre-existing).
    per_record = 2 if form == "local" else 11
    row_12 = [f for f in facts if f["pointer"]["cell"][1:] == "12"]
    assert len(row_12) == (per_record if row_12_binds else 0), (form, title, len(row_12))
    assert len([f for f in facts if f["pointer"]["cell"][1:] in ("8", "9", "10")]) == 3 * per_record


# --- Item 5: a refused plan period falls back to a local period read from a reporting line; it never suppresses it ---

@pytest.mark.parametrize("stamp", ["Stand: 15.03.2026", "Datum 26.04.2026", "Seite 1 per 26.04.2026", "Abfrage zum 01.02.2026", "Gültig ab 01.01.2019"])
@pytest.mark.parametrize("reporting", ["Jahresabschluss per 30.06.2025", "Jahresabschluss"])
def test_item5_refused_sheet_stamp_period_falls_back_to_the_local_reporting_period(tmp_path, stamp, reporting):
    from openpyxl import load_workbook
    source = adv.fixtures.flat_sap_register(tmp_path / "reg.xlsx")
    workbook = load_workbook(source); ws = workbook["Sheet1"]; ws["A4"] = reporting; ws["A6"] = stamp; workbook.save(source)
    _, local = adv.run(tmp_path, source, {"plans": []}, "local")
    recorded = adv.plan("flat_sap_register")
    recorded["plans"][0]["period_contexts"] = [{"observation_id": "@cell:Sheet1!A6", "temporal_role": "closing", "scope": "sheet"}]
    _, facts = adv.run(tmp_path, source, recorded, "mapped")
    def fact_key(f):
        return (f["pointer"]["cell"], f["field"], f["asset_identity"].get("asset_or_class"), f["asset_identity"].get("period"), f["amount"]["normalized_base"])
    assert not [f for f in facts if f["asset_identity"].get("period") != "2025"], "the stamp never becomes the period"
    if reporting.endswith("2025"):
        assert local and {fact_key(f) for f in local} <= {fact_key(f) for f in facts}, "a refused plan period must not remove a local fact"


# --- Item 7: the year's opening (01.01.Y / 31.12.Y-1) or closing (31.12.Y) column date agrees with the sheet period Y ---

@pytest.mark.skipif(not F08.exists(), reason="private corpus unavailable")
def test_item7_f08_prior_year_sheet_boundary_dates_agree_with_the_sheet_period(tmp_path):
    evidence = replay(F08, "f08_excel_spiegel_2y_01", tmp_path / "out")
    truth = f08_sheet_truth(F08)
    got = {(f["pointer"]["sheet"], f["pointer"]["cell"]): (f["field"], Decimal(str(f["amount"]["normalized_base"])), f["asset_identity"].get("period"), f["amount"]["unit"]) for f in evidence["facts"]}
    assert not [(key, value) for key, value in got.items() if truth.get(key) != value]
    for cell in ("B16", "G16", "I16", "L16", "N16", "O16"):
        assert got.get(("Tabelle2", cell)) == truth[("Tabelle2", cell)], cell


@pytest.mark.skipif(not F08.exists(), reason="private corpus unavailable")
def test_item7_a_column_date_from_another_year_still_abstains(tmp_path):
    import datetime
    from openpyxl import load_workbook
    source = tmp_path / "f08_conflict.xlsx"
    workbook = load_workbook(F08, data_only=True); workbook["Tabelle2"]["O5"] = datetime.datetime(2021, 12, 31); workbook.save(source)
    plans = [plan for n in (1, 2) for plan in json.loads((PLANS / f"live_r5_f08_excel_spiegel_2y_01_req{n}.json").read_text())["plans"]]
    document_id = plans[0]["region_observation_id"].split(":xlsx:")[0]
    def placeholders(text):  # the edited copy has another document id: re-address the live plan by sheet coordinates
        text = re.sub(re.escape(document_id) + r":xlsx:([^\"!:]+)!([A-Z]+\d+)", r"@cell:\1!\2", text)
        text = re.sub(re.escape(document_id) + r":xlsx:([^\"!:]+):row:(\d+)", r"@row:\1:\2", text)
        return re.sub(re.escape(document_id) + r":xlsx:([^\"!:]+):region", r"@region:\1", text)
    with stub.recorded_plan_server({"plans": [json.loads(placeholders(json.dumps(plan))) for plan in plans]}) as endpoint:
        result = run_input(source, tmp_path / "out", RunConfig(mapper=MapperConfig(True, endpoint, "replay", 2, 256)))
    facts = json.loads(Path(result.evidence_path).read_text())["facts"]
    assert len([f for f in facts if f["pointer"]["sheet"] == "Tabelle2"]) >= 15, (collections_counter(facts), result.diagnostics, len(stub.RecordedPlanHandler.requests))
    assert not [f for f in facts if f["pointer"]["sheet"] == "Tabelle2" and f["pointer"]["cell"][0] == "O"], "a genuine year conflict abstains"


def collections_counter(facts):
    return Counter(f["pointer"]["sheet"] for f in facts)
