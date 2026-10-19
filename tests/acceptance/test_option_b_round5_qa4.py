"""Fix round 5, QA round 4 (779e1a4) probes: G2' (local field drops only when inadmissible for the row's measure) and the
coverage items (per-block rejection, unmarked subtotal, TEUR in an identity row, local "Abgänge AfA")."""
from __future__ import annotations

import copy
from decimal import Decimal
import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from carthagent.pipeline import RunConfig, run_input

ROOT = Path(__file__).resolve().parents[2]


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


adv = _load(Path(__file__).with_name("test_option_b_adversarial.py"), "qa4_adversarial")
reqa = _load(Path(__file__).with_name("test_option_b_reqa.py"), "qa4_reqa")
S = "Anlagengitter"
R = lambda n: f"@row:{S}:{n}"  # noqa: E731
C = lambda ref: f"@cell:{S}!{ref}"  # noqa: E731


def many(path, blocks=30, mutate=None):
    source, truth = adv.fixtures.sap_block_many(path, blocks=blocks)
    if mutate:
        workbook = load_workbook(source); mutate(workbook[S], truth); workbook.save(source)
    return source, truth


def v11_plan(rows=((1, 16, "ahk"), (2, 17, "cum_depreciation"), (3, 18, "book_value"))):
    recorded = adv.plan("sap_block_unlabelled"); q = recorded["plans"][0]
    q.pop("record_blocks"); q.pop("column_spans")
    q["block_pattern"] = [{"rows": [{"ordinal": o, "header_row_id": R(h), "measure": m} for o, h, m in rows], "role": "data", "row_count": 4, "applies_to": "all_blocks_of_same_shape"}]
    q["unit_contexts"] = [{"observation_id": C("H19"), "unit": "EUR", "scope": "block"}]
    return recorded


def explicit_plan(blocks, start=19):
    recorded = adv.plan("sap_block_unlabelled"); q = recorded["plans"][0]
    q["record_blocks"] = [{"identity_row_id": R(start + 5 * i), "rows": [{"row_id": R(start + 5 * i + k), "header_row_id": R(15 + k), "measure": m}
                                                                         for k, m in ((1, "ahk"), (2, "cum_depreciation"), (3, "book_value"))], "role": "data"} for i in range(blocks)]
    q["unit_contexts"] = [{"observation_id": C(f"H{start + 5 * i}"), "unit": "EUR", "scope": "block"} for i in range(blocks)]
    return recorded


def check(tmp_path, source, recorded, truth):
    _, facts = adv.run(tmp_path, source, recorded)
    bad = []
    for fact in facts:
        cell = fact["pointer"]["cell"]; expected = truth.get(cell)
        got = (fact["field"], fact["asset_identity"].get("class_code") or fact["asset_identity"].get("asset_or_class"), Decimal(fact["amount"]["normalized_base"]))
        if expected is None or got != (expected[0], expected[1], Decimal(str(expected[2]))):
            bad.append((cell, got, expected))
    return facts, bad


def soft(ws, truth):
    ws["C16"] = "Zugang (lfd.)"; ws["E16"] = "Umbuchung (lfd.)"


def swap_labels(recorded):
    for header in recorded["plans"][0]["field_headers"]:
        if header["observation_id"] == C("C16"):
            header["semantic_field"] = "umb"
        elif header["observation_id"] == C("E16"):
            header["semantic_field"] = "zugang"
    return recorded


def test_g2_stacked_soft_swap_explicit(tmp_path):
    source, truth = many(tmp_path / "s.xlsx", blocks=3, mutate=soft)
    _, bad = check(tmp_path, source, swap_labels(explicit_plan(3)), truth)
    assert not bad, bad[:3]


def test_g2_stacked_soft_swap_pattern(tmp_path):
    source, truth = many(tmp_path / "s.xlsx", mutate=soft)
    _, bad = check(tmp_path, source, swap_labels(v11_plan()), truth)
    assert not bad, bad[:3]


def test_g2_stacked_soft_correct_pattern_binds_everything(tmp_path):
    source, truth = many(tmp_path / "s.xlsx", mutate=soft)
    facts, bad = check(tmp_path, source, v11_plan(), truth)
    assert not bad and len(facts) == len(truth) == 390


@pytest.mark.parametrize("swap", [False, True])
def test_g2_two_row_spiegel_soft_texts(tmp_path, swap):
    source = reqa.build_av(tmp_path / "av.xlsx", movement_texts=("Anfangsbestand", "Zugang (lfd.)", "Abg.", "Umbuchung (lfd.)", "Endbestand"))
    recorded = reqa.av_plan()
    if swap:
        for axis in recorded["plans"][0]["movement_axes"]:
            if axis["observation_id"] == "@cell:AV!D5":
                axis["movement_axis"] = "transfer"
            elif axis["observation_id"] == "@cell:AV!F5":
                axis["movement_axis"] = "addition"
    _, facts = adv.run(tmp_path, source, recorded)
    assert not adv.wrong(facts, reqa.AV_TRUTH)


# --- coverage items ---

def test_unmarked_subtotal_block_does_not_strip_identity_from_later_blocks(tmp_path):
    row = 19 + 10 * 5
    def mutate(ws, truth):
        ws[f"E{row}"] = None; ws[f"F{row}"] = "Summe Bilanzposition"; ws[f"C{row}"] = "2101200"
        for k in (1, 2, 3):
            for column in "BCDEFGH":
                truth.pop(f"{column}{row + k}", None)
    source, truth = many(tmp_path / "s.xlsx", mutate=mutate)
    facts, bad = check(tmp_path, source, v11_plan(), truth)
    assert not bad and len(facts) == len(truth)


def test_prior_year_grid_below_costs_only_its_own_blocks(tmp_path):
    def mutate(ws, truth):
        start = 19 + 30 * 5 + 3
        ws[f"A{start}"] = "Berichtsdatum: 31.12.2024   Anlagengitter  Vorjahr"
        for offset, header in enumerate(adv.fixtures.SAP_HEADER):
            for column, text in header.items():
                ws[f"{column}{start + 2 + offset}"] = text
        row = start + 6
        for index in range(6):
            identity, *measures = adv.fixtures.SAP_BLOCKS[index % 2]; identity = {**identity, "E": str(1000 + index)}
            for column, text in identity.items():
                ws[f"{column}{row}"] = text
            for k, values in enumerate(measures, start=1):
                for column, value in zip("BCDEFGH", values):
                    if value is not None:
                        ws[f"{column}{row + k}"] = value * 0.9 if isinstance(value, float) and value else value
            row += 5
    source, truth = many(tmp_path / "s.xlsx", mutate=mutate)
    facts, bad = check(tmp_path, source, v11_plan(), truth)
    assert not bad and len(facts) == len(truth) == 390


@pytest.mark.parametrize("unit_text", ["TEUR", "T€", "Tsd. EUR"])
def test_teur_in_an_identity_row_scales_that_block(tmp_path, unit_text):
    row = 19 + 13 * 5
    def mutate(ws, truth):
        ws[f"H{row}"] = unit_text
        for k in (1, 2, 3):
            for column in "BCDEFGH":
                if f"{column}{row + k}" in truth:
                    field, klass, value = truth[f"{column}{row + k}"]; truth[f"{column}{row + k}"] = (field, klass, value * 1000)
    source, truth = many(tmp_path / "s.xlsx", mutate=mutate)
    facts, bad = check(tmp_path, source, v11_plan(), truth)
    assert not bad and len(facts) == 390
    assert {f["amount"]["unit"] for f in facts if row < int(f["pointer"]["cell"][1:]) <= row + 3} == {"TEUR"}


def test_local_abgaenge_afa_word_order_is_afa_abgang(tmp_path):
    source = adv.register(tmp_path, lambda ws: ws.__setitem__("O13", "Abgänge AfA"))
    facts = json.loads(Path(run_input(source, tmp_path / "o", RunConfig()).evidence_path).read_text())["facts"]
    by_cell = {f["pointer"]["cell"]: f["field"] for f in facts}
    assert by_cell.get("O16") in (None, "afa_abgang") and "abgang" not in {by_cell.get(f"O{r}") for r in (14, 15, 16)}
    assert not adv.wrong(facts, adv.reg_truth())


# --- live v11 f05 plans (Claude, schema v11): the identity row is ordinal 0 of the pattern ---

F05 = ROOT / "hackathon-local" / "ElabHackathon" / "A3-Hackathon-Corpus" / ("f05_sap_class_" "gitter") / ("f05_sap_class_" "gitter_02.xlsx")
SAP_TEXT_FIELDS = {"AHK GJ-Beg": "ahk_open", "Zugang": "zugang", "Abgang": "abgang", "Umbuchung": "umb", "aktuelle AHK": "ahk_close", "AfA GJ-Beg": "kum_afa_open",
                   "AfA des Jahres": "total_afa_gj", "AfA Abgang": "afa_abgang", "AfA Umbuchung": "afa_umb", "Zuschreibungen": "zuschreibung",
                   "kumulierte AfA": "kum_afa_close", "Buchwrt GJ-Beg": "bw_prev", "lfd Buchwert": "bw_close"}


def sheet_truth(path):
    """Truth read from the sheet itself: class records (class code in E, no subtotal star in A), fields from the header text."""
    ws = load_workbook(path, data_only=True)["Anlagengitter"]
    truth = {}
    for row in range(19, ws.max_row + 1):
        klass = ws[f"E{row}"].value
        if not (isinstance(klass, str) and klass.strip().isdigit()) or ws[f"A{row}"].value or isinstance(ws[f"B{row + 1}"].value, str):
            continue
        for k in (1, 2, 3):
            for column in "BCDEFGH":
                field = SAP_TEXT_FIELDS.get(str(ws[f"{column}{15 + k}"].value or "").strip())
                value = ws[f"{column}{row + k}"].value
                if field and isinstance(value, (int, float)):
                    truth[f"{column}{row + k}"] = (field, klass.strip(), Decimal(str(value)))
    return truth


@pytest.mark.skipif(not F05.exists(), reason="private corpus unavailable")
def test_live_v11_f05_pattern_with_identity_ordinal_binds_with_zero_wrong():
    from carthagent.acquire_xlsx import acquire_xlsx
    from carthagent.mapper import MapperConfig, OpenAICompatibleMapper
    from carthagent.pipeline import _deduplicate_mapped
    from carthagent.policies import DEFAULT_BOUNDS, DEFAULT_POLICIES
    from carthagent.semantics import interpret
    records = [json.loads((ROOT / "tests" / "fixtures" / "option_b_plans" / f"live_v11_f05_req{n}.json").read_text()) for n in (1, 2)]
    document_id = records[0]["document_id"]
    observations = acquire_xlsx(F05, document_id, "a" * 64, DEFAULT_BOUNDS).observations
    mapped = []
    for record in records:
        mapper = OpenAICompatibleMapper(MapperConfig(True, "http://127.0.0.1:1/v1/responses", "replay", 1, 1)); mapper.last_verification = []
        mapped += mapper._validate(document_id, observations, {"plans": copy.deepcopy(record["plans"])})
    facts = interpret(observations, (), DEFAULT_POLICIES, _deduplicate_mapped(tuple(mapped)))[1]
    truth = sheet_truth(F05)
    got = {f.physical.occurrence.split("!")[1]: (f.business.field, f.business.class_code, f.amount.normalized_value, f.business.period, f.business.unit) for f in facts}
    assert got, "the live pattern binds"
    assert not [(cell, value) for cell, value in got.items() if truth.get(cell) != value[:3] or value[3:] != ("2025", "EUR")]
    assert got["B20"][:2] == ("ahk_open", "1000") and got["B21"][0] == "kum_afa_open" and got["H22"][0] == "bw_close"
