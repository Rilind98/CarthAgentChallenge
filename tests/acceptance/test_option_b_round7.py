"""Round 7: review rounds 4-6 should-fixes (R7-1..3) and QA round 6 items (R7-4..8)."""

from __future__ import annotations

from collections import Counter
from decimal import Decimal
import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from carthagent.mapper import MapperConfig
from carthagent.pipeline import RunConfig, run_input

ROOT = Path(__file__).resolve().parents[2]


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


adv = _load(Path(__file__).with_name("test_option_b_adversarial.py"), "r7_adv")


def local_facts(tmp_path, source, out="local"):
    result = run_input(source, tmp_path / out, RunConfig())
    assert result.evidence_path is not None, result.diagnostics
    return json.loads(Path(result.evidence_path).read_text())["facts"]


# --- R7-1 K1 groups by business key without field: two records behind one key abstain entirely ---

def test_r7_1_a_field_present_in_only_one_of_two_colliding_records_abstains(tmp_path):
    source = adv.fixtures.flat_sap_register(tmp_path / "reg.xlsx")
    workbook = load_workbook(source); ws = workbook.active; ws["B13"] = "Bemerkung"
    for row in (14, 15, 16):
        ws[f"B{row}"] = None
    ws["G15"] = "n/a"; workbook.save(source)
    facts = local_facts(tmp_path, source)
    shared = [f for f in facts if f["asset_identity"].get("asset_id") == "100001002"]
    assert not shared, [(f["pointer"]["cell"], f["field"]) for f in shared]
    assert [f for f in facts if f["asset_identity"].get("asset_id") == "100001001"], "the unshared record still binds"


def test_r7_1_one_multi_row_record_is_one_record(tmp_path):
    """A SAP block (identity row + three measure rows) is ONE record: its facts never collide with each other."""
    _, facts = adv.run(tmp_path, adv.fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx"), adv.plan("sap_block_unlabelled"))
    assert len(facts) == 26


# --- R7-2 a sheet/page-scoped plan context applies only to its own sheet (XLSX) or page (PDF) ---

def _pdf_word(identifier, text, page, box):
    from carthagent.models import SourceObservation
    return SourceObservation(f"d:{identifier}", "d", "a" * 64, "word", f"page:{page + 1}:{identifier}", "pdf_native", text, box, {"page_index": page})


def test_r7_2_a_period_cited_on_page_one_never_dates_page_two():
    from carthagent.mapper import _flat_context
    title = _pdf_word("t", "Geschäftsjahr 2025", 0, (50, 40, 200, 52))
    first, second = _pdf_word("a1", "1.000,00", 0, (300, 200, 350, 212)), _pdf_word("a2", "2.000,00", 1, (300, 200, 350, 212))
    by_id = {o.id: o for o in (title, first, second)}
    periods = [{"observation_id": title.id, "temporal_role": "current_period", "scope": "sheet"}]
    assert _flat_context(periods, first, by_id, []) is not None, "control: the cited page takes the period"
    assert _flat_context(periods, second, by_id, []) is None, "page 2 never takes a period cited on page 1"


def test_r7_2_a_sheet_context_never_crosses_sheets():
    from carthagent.mapper import _flat_context
    from carthagent.models import SourceObservation
    cell = lambda ref, text: SourceObservation(f"d:xlsx:{ref}", "d", "a" * 64, "cell", ref, "xlsx_native", text)  # noqa: E731
    title, here, there = cell("A!A1", "Geschäftsjahr 2025"), cell("A!C5", "10"), cell("B!C5", "10")
    by_id = {o.id: o for o in (title, here, there)}
    periods = [{"observation_id": title.id, "temporal_role": "current_period", "scope": "sheet"}]
    assert _flat_context(periods, here, by_id, []) is not None and _flat_context(periods, there, by_id, []) is None


# --- R7-3 a run into an output directory written by an older build never hits an immutable record collision ---

def test_r7_3_versions_moved_past_9f94a3b():
    from carthagent.exporter import TemplateExportAdapter
    from carthagent.pipeline import STRUCTURE_VERSION
    from carthagent.semantics import SEMANTICS_VERSION
    assert (SEMANTICS_VERSION, STRUCTURE_VERSION, TemplateExportAdapter.IMPLEMENTATION_VERSION) != ("14", "8", "2")
    assert int(SEMANTICS_VERSION) > 14 and int(STRUCTURE_VERSION) > 8 and int(TemplateExportAdapter.IMPLEMENTATION_VERSION) > 2


def test_r7_3_rerun_after_an_older_build_with_other_evidence_does_not_collide(tmp_path, monkeypatch):
    """An older build wrote the same abstentions/facts with other evidence (simulated: evidence ids in another order).
    The current build must still publish into the same output directory: record ids are content-addressed."""
    from dataclasses import replace
    from carthagent import structure
    source = adv.fixtures.flat_sap_register(tmp_path / "reg.xlsx")
    original = structure.ContextResolver.resolve
    def older(self, amount, observations):
        context = original(self, amount, observations)
        return replace(context, evidence_ids=tuple(reversed(context.evidence_ids)))
    from carthagent import pipeline
    fingerprint = pipeline._implementation_fingerprint
    monkeypatch.setattr(structure.ContextResolver, "resolve", older)
    monkeypatch.setattr(pipeline, "_implementation_fingerprint", lambda: {**fingerprint(), "semantics": "14", "structure": "8"})  # the older build's markers
    from carthagent import semantics
    current = semantics.SEMANTICS_VERSION
    monkeypatch.setattr(semantics, "SEMANTICS_VERSION", "14")
    first = run_input(source, tmp_path / "out", RunConfig())
    assert first.evidence_path is not None, first.diagnostics
    monkeypatch.setattr(structure.ContextResolver, "resolve", original)
    monkeypatch.setattr(pipeline, "_implementation_fingerprint", fingerprint); monkeypatch.setattr(semantics, "SEMANTICS_VERSION", current)
    second = run_input(source, tmp_path / "out", RunConfig())
    assert second.evidence_path is not None and not [d for d in second.diagnostics if "collision" in d], second.diagnostics
    assert json.loads(Path(second.evidence_path).read_text())["facts"]


# --- R7-5 a title line is a reporting line only with a reporting word; stamp words veto; a date alone never dates a sheet ---

@pytest.mark.parametrize(("title", "period"), [
    ("Stand 15.03.2026", None), ("Ausdruck vom 26.04.2026", None), ("Gedruckt am 26.04.2026", None), ("Datum: 26.04.2026", None),
    ("Druck 26.04.2026", None), ("Export 26.04.2026", None), ("erstellt am 26.04.2026", None), ("Abfrage 01.02.2026", None),
    ("generated 2026-04-26", None), ("printed 26.04.2026", None), ("26.04.2026", None),
    ("Berichtsdatum 31.12.2025", "2025"), ("Jahresabschluss per 30.06.2025", "2025"), ("Geschäftsjahr 2025", "2025"),
])
def test_r7_5_stamps_never_date_a_sheet(tmp_path, title, period):
    source = adv.fixtures.flat_sap_register(tmp_path / "reg.xlsx")
    workbook = load_workbook(source); ws = workbook["Sheet1"]; ws["A4"] = title; ws["A6"] = None; workbook.save(source)
    facts = local_facts(tmp_path, source)
    periods = Counter(f["asset_identity"].get("period") for f in facts)
    if period is None:
        assert not facts, periods
    else:
        assert facts and set(periods) == {period}, periods


# --- R7-7 a merged unit cell applies to every column of its merged range, never beyond ---

@pytest.mark.parametrize("variant", ["merged C6:E6", "separate C6 only"])
def test_r7_7_merged_unit_cell_covers_its_merged_columns(tmp_path, variant):
    source = adv.fixtures.two_year_two_sheets(tmp_path / "s.xlsx")
    workbook = load_workbook(source); ws = workbook["Tabelle1"]; ws["D6"] = None; ws["E6"] = None
    if variant.startswith("merged"):
        ws.merge_cells("C6:E6")
    workbook.save(source)
    recorded = adv.plan("two_year_two_sheets")
    recorded["plans"] = [plan for plan in recorded["plans"] if plan["region_observation_id"].endswith("Tabelle1")]
    plan = recorded["plans"][0]; plan.pop("record_blocks", None); plan.pop("column_spans", None); plan["block_pattern"] = []
    columns = [c for c in "CDEFHJKLMOP" if c not in "DE"]
    plan["unit_contexts"] = [{"observation_id": f"@cell:Tabelle1!{c}6", "unit": "EUR", "scope": "column"} for c in columns]
    _, facts = adv.run(tmp_path, source, recorded)
    facts = [f for f in facts if f["pointer"]["sheet"] == "Tabelle1"]
    truth = {key: value for key, value in adv.spiegel_truth().items() if key[0] == "Tabelle1"}
    assert not adv.wrong(facts, truth)
    bound_columns = {f["pointer"]["cell"][0] for f in facts}
    if variant.startswith("merged"):
        assert {"D", "E"} <= bound_columns and len(facts) == 33, (sorted(bound_columns), len(facts))
    else:
        assert not {"D", "E"} & bound_columns, "an unmerged unit cell never reaches the neighbouring columns"


# --- R7-8 an untitled sheet inherits the workbook's single title-line period only with a total tie-out and the same entity ---

CLASSES = (("1000", "EDV-Software", 1000, 0, 0), ("3000", "Maschinen", 2500, 100, -100))


def _with_class_summary(path, totals=(3500, 100, -100), entity="7307"):
    source = adv.fixtures.flat_sap_register(path)  # Sheet1: "Jahresabschluss per 30.06.2025", "BK: 7307", SUMMEN 3500 / 100 / -100
    workbook = load_workbook(source); ws = workbook.create_sheet("Sheet2")
    ws.append(["Buchungskreis", "Anlagenklasse", "Bezeichnung", "Währ-ung", "AHK GJ-Beg", "Zugang", "Abgang"])
    for code, name, *values in CLASSES:
        ws.append([entity, code, name, "EUR", *values])
    ws.append([f"Buchungskreis {entity}", None, None, "EUR", *totals])
    workbook.save(source); return source


@pytest.mark.parametrize(("totals", "entity", "inherits"), [
    ((3500, 100, -100), "7307", True),
    ((3300, 100, -100), "7307", False),  # an untitled prior-year sheet: its totals do not tie out to the dated sheet
    ((3500, 100, -100), "7310", False),  # the totals tie out but the company code differs
])
def test_r7_8_untitled_sheet_inherits_the_period_only_with_a_tie_out(tmp_path, totals, entity, inherits):
    source = _with_class_summary(tmp_path / "reg.xlsx", totals, entity)
    if not inherits and totals[0] != 3500:
        workbook = load_workbook(source); ws = workbook["Sheet2"]; ws["E3"] = 2300; workbook.save(source)  # rows still sum to the printed total
    facts = local_facts(tmp_path, source)
    sheet_two = [f for f in facts if f["pointer"]["sheet"] == "Sheet2"]
    assert [f for f in facts if f["pointer"]["sheet"] == "Sheet1"], "control: the dated sheet binds"
    if inherits:
        assert sheet_two and {f["asset_identity"].get("period") for f in sheet_two} == {"2025"}
        got = {(f["pointer"]["cell"], f["field"], Decimal(f["amount"]["normalized_base"])) for f in sheet_two}
        assert ("F3", "zugang", Decimal("100")) in got and ("G3", "abgang", Decimal("-100")) in got  # "AHK GJ-Beg" has no local field (pre-existing)
    else:
        assert not [f for f in sheet_two if f["asset_identity"].get("period")], sheet_two[:2]


CORPUS = ROOT / "hackathon-local" / "ElabHackathon" / "A3-Hackathon-Corpus"
F09 = CORPUS / "f09_excel_register" / "f09_excel_register_01.xlsx"
F09_SHEET2_FIELDS = {"F": "zugang", "G": "abgang", "H": "umb", "N": "kum_afa_open", "O": "total_afa_gj", "P": "afa_abgang", "Q": "afa_umb", "S": "zuschreibung", "T": "kum_afa_close", "U": "bw_prev", "V": "bw_close", "E": "ahk_open", "L": "ahk_close"}


@pytest.mark.skipif(not F09.exists(), reason="private corpus unavailable")
def test_r7_8_f09_sheet2_inherits_2025_with_zero_wrong(tmp_path):
    stub = _load(ROOT / "tests" / "fixtures" / "option_b_stub.py", "r7_stub")
    plans = [plan for n in (1, 2) for plan in json.loads((ROOT / "tests" / "fixtures" / "option_b_plans" / f"live_r5_f09_excel_register_01_req{n}.json").read_text())["plans"]]
    with stub.recorded_plan_server({"plans": plans}) as endpoint:
        result = run_input(F09, tmp_path / "out", RunConfig(mapper=MapperConfig(True, endpoint, "replay", 2, 256)))
    facts = [f for f in json.loads(Path(result.evidence_path).read_text())["facts"] if f["pointer"]["sheet"] == "Sheet2"]
    ws = load_workbook(F09, data_only=True)["Sheet2"]
    assert facts and {f["asset_identity"].get("period") for f in facts} == {"2025"}
    wrong = [(f["pointer"]["cell"], f["field"]) for f in facts
             if F09_SHEET2_FIELDS.get(f["pointer"]["cell"][0]) != f["field"] or Decimal(str(ws[f["pointer"]["cell"]].value)).quantize(Decimal("0.01")) != Decimal(f["amount"]["normalized_base"]).quantize(Decimal("0.01"))
             or str(ws["B" + f["pointer"]["cell"][1:]].value) != f["asset_identity"].get("class_code")]
    assert not wrong, wrong[:5]
    assert not [f for f in facts if int(f["pointer"]["cell"][1:]) >= 15], "the company total rows are not classes"
