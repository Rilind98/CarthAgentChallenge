"""Fix round 5: class-closing backstops (K1 key collision, K2 period/table end) and specific items S1-S6.

Ported from qa_probes/test_round3_xlsx.py, test_round3_azure.py and the review repros (cr_exp E6/E7/E8, L1).
Truth comes from fixture constants; assertions are absolute unless a probe names a pre-existing local path.
"""
from __future__ import annotations

import collections
import copy
from decimal import Decimal
import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from carthagent.pipeline import RunConfig, run_input

ROOT = Path(__file__).resolve().parents[2]


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


adv = _load(Path(__file__).with_name("test_option_b_adversarial.py"), "r5_adversarial")
reqa = _load(Path(__file__).with_name("test_option_b_reqa.py"), "r5_reqa")


def ident(*pairs):
    return [{"observation_id": f"@cell:Sheet1!{c}", "identity_role": r} for c, r in pairs]


def register(tmp_path, mutate, plan=None, out="mapped"):
    source = adv.register(tmp_path, mutate)
    if plan is None:
        result = run_input(source, tmp_path / out, RunConfig())
        return json.loads(Path(result.evidence_path).read_text())["facts"]
    return adv.run(tmp_path, source, plan, out)[1]


def collisions(facts):
    keys = collections.Counter((f["asset_identity"].get("entity"), f["asset_identity"].get("asset_id"), f["asset_identity"].get("asset_sub_id"), f["asset_identity"].get("class_code"),
                                f["asset_identity"].get("asset_or_class"), f["field"], f["asset_identity"].get("period")) for f in facts)
    return [k for k, n in keys.items() if n > 1]


def reg_plan(ids=None, periods=None, header_rows=None):
    recorded = adv.plan("flat_sap_register"); q = recorded["plans"][0]
    if ids is not None:
        q["identity_headers"] = ids
    if periods is not None:
        q["period_contexts"] = periods
    if header_rows is not None:
        q["header_rows"] = header_rows
    return recorded


# --- K1 business-key collision ---

def _kennung(ws, values=True):
    ws["B13"] = "Kennung"; ws["B15"] = "K1" if values else None; ws["B16"] = "K2" if values else None


@pytest.mark.parametrize("values", [True, False])
def test_k1_kennung_column_never_collapses_two_records_local_only(tmp_path, values):
    facts = register(tmp_path, lambda ws: _kennung(ws, values))
    assert not collisions(facts) and not adv.wrong(facts, adv.reg_truth())


def test_k1_s3_empty_header_descriptive_label_never_declares_an_id_column(tmp_path):
    """Review R-B1 (E7): the model labels the blank cell above "Kennung" class_name; no colliding pair may bind."""
    def mutate(ws):
        _kennung(ws); ws["A12"] = "Stammdaten"; ws["B12"] = " "; ws["F12"] = "Werte"
    recorded = reg_plan([h for h in ident(("A13", "asset_id"), ("D13", "source_name"))] + [{"observation_id": "@cell:Sheet1!B12", "identity_role": "class_name"}],
                        header_rows=["@row:Sheet1:12", "@row:Sheet1:13"])
    facts = register(tmp_path, mutate, recorded)
    assert not collisions(facts) and not adv.wrong(facts, adv.reg_truth())
    assert not [f for f in facts if f["pointer"]["cell"][1:] in ("15", "16")], "the two records sharing Anlage 100001002 stay abstained"


def test_k1_exporter_refuses_two_facts_for_one_destination_cell():
    from carthagent.exporter import TemplateExportAdapter
    from carthagent.models import BusinessIdentity, PhysicalIdentity, ValidatedFact
    from carthagent.numeric import parse_amount
    def fact(identifier, cell, value):
        business = BusinessIdentity(asset_or_class="A-1", asset_id="A-1", granularity="asset", row_kind="data", period="2025", field="ahk_open", currency="EUR", unit="EUR")
        return ValidatedFact(f"d:fact:{identifier}", "d", PhysicalIdentity("a" * 64, "d", f"S!{cell}"), business, parse_amount(value), f"d:x:S!{cell}",
                             {"kind": "xlsx", "sheet": "S", "cell": cell, "observation_id": f"d:x:S!{cell}"}, (), {})
    bindings = TemplateExportAdapter().plan_bindings("s", (fact("1", "B2", "10"), fact("2", "B3", "11"), ))
    assert [b.status for b in bindings] == ["destination_collision", "destination_collision"] and all(b.destination.cell is None for b in bindings)


# --- K2 + S4: a table ends at its grand total / at a one-cell title naming another year ---

PRIOR = {14: ("Software", (800, 200, 0, 0, 1000), (500, 100, 0, 600), (400, 300)), 15: ("Maschinen", (2800, 200, 0, 0, 3000), (700, 300, 0, 1000), (2000, 2100))}


def build_two(path, label_total=True):
    reqa.build_av(path)
    workbook = load_workbook(path); ws = workbook["AV"]
    if label_total:
        ws["B12"] = "Gesamt Anlagevermögen"
    for cols, index in ((reqa.AHK_COLS, 1), (reqa.AFA_COLS, 2), (reqa.BV_COLS, 3)):
        for position, column in enumerate(cols):
            ws[f"{column}12"] = sum(record[index][position] for record in reqa.RECS.values())
    ws["B13"] = "Entwicklung Vorjahr 2024"
    for row, (name, ahk, afa, bv) in PRIOR.items():
        ws[f"B{row}"] = name
        for cols, values in ((reqa.AHK_COLS, ahk), (reqa.AFA_COLS, afa), (reqa.BV_COLS, bv)):
            for column, value in zip(cols, values):
                ws[f"{column}{row}"] = value
    workbook.save(path); return path


@pytest.mark.parametrize(("label_total", "declare_prior", "total_role"), [(True, True, False), (True, True, True), (False, True, False), (True, False, False), (False, False, False)])
def test_k2_s4_rows_below_the_grand_total_and_the_prior_year_title_never_bind(tmp_path, label_total, declare_prior, total_role):
    source = build_two(tmp_path / "two.xlsx", label_total)
    recorded = reqa.av_plan(); q = recorded["plans"][0]
    R = lambda n: f"@row:AV:{n}"  # noqa: E731
    if declare_prior:
        q["record_blocks"] += [{"identity_row_id": R(r), "rows": [{"row_id": R(r), "header_row_id": R(5), "measure": None}], "role": "data"} for r in PRIOR]
    if total_role:
        q["record_blocks"].append({"identity_row_id": R(12), "rows": [{"row_id": R(12), "header_row_id": R(5), "measure": None}], "role": "total"})
    for name, plan in (("local", {"plans": []}), ("mapped", recorded)):
        _, facts = adv.run(tmp_path, source, plan, name)
        assert not [f["pointer"]["cell"] for f in facts if f["pointer"]["cell"][1:] in ("12", "13", "14", "15")], name
        assert not adv.wrong(facts, reqa.AV_TRUTH), name


def test_k2_two_year_sheets_keep_their_facts(tmp_path):
    _, facts = adv.run(tmp_path, adv.fixtures.two_year_two_sheets(tmp_path / "s.xlsx"), adv.plan("two_year_two_sheets"))
    assert len(facts) == 66 and not adv.wrong(facts, adv.spiegel_truth())


# --- S1 sheet period never from a column header ---

@pytest.mark.parametrize("text", ["Restbuchwert zum 31.12.2024", "Buchwert per 31.12.2024", "AfA GJ-Beg 2024", "Buchwrt GJ-Beg 31.12.2024", "Buchwrt GJ-Beg"])
def test_s1_column_header_is_never_the_sheet_period(tmp_path, text):
    def mutate(ws):
        ws["A4"] = "Jahresabschluss"; ws["T13"] = text
    """The mapper never takes the sheet period from a header-band cell. (The LOCAL rules still read a lone dated column
    header as that column's period; that pre-existing local reading is reported separately, so the assertion is on what
    the mapper adds over a local-only run.)"""
    (tmp_path / "local").mkdir(); (tmp_path / "mapped").mkdir()
    local = register(tmp_path / "local", mutate)
    facts = register(tmp_path / "mapped", mutate, reg_plan(periods=[{"observation_id": "@cell:Sheet1!T13", "temporal_role": "closing", "scope": "sheet"}]))
    identity = lambda f: (f["pointer"]["cell"], f["field"], json.dumps(f["asset_identity"], sort_keys=True), f["amount"]["normalized_base"])  # noqa: E731
    added = [f for f in facts if identity(f) not in {identity(g) for g in local}]
    assert not [f for f in added if f["asset_identity"].get("period") == "2024"] and not adv.wrong(added, adv.reg_truth())


def test_s1_gj_is_a_period_word_only_with_a_year():
    from carthagent.mapper import PERIOD_WORDS
    assert PERIOD_WORDS.search("GJ 2025") and PERIOD_WORDS.search("Geschäftsjahr 2025")
    assert not PERIOD_WORDS.search("AfA GJ-Beg 2024") and not PERIOD_WORDS.search("Buchwrt GJ-Ende")


# --- S2 Azure digit box vs column ---

azure = _load(ROOT / "tests" / "unit" / "test_azure_layout.py", "r5_azure_layout")
stub = azure.stub  # noqa: F401
documents = azure.documents  # noqa: F401
AZURE_TRUTH = {("000101", "ahk_open"): "1000.00", ("000101", "zugang"): "120.00", ("000101", "abgang"): "20.00", ("000101", "ahk_close"): "1100.00", ("000101", "bw_close"): "800.00",
               ("000202", "ahk_open"): "2000.00", ("000202", "zugang"): "340.00", ("000202", "abgang"): "40.00", ("000202", "ahk_close"): "2300.00", ("000202", "bw_close"): "1600.00"}


def _azure_facts(result):
    from carthagent.policies import DEFAULT_POLICIES
    from carthagent.semantics import interpret
    return interpret(result.observations, result.coverage, DEFAULT_POLICIES)[1]


def _azure_wrong(facts, truth):
    return [(f.business.asset_or_class, f.business.field, str(abs(f.amount.source_signed_value))) for f in facts if truth.get((f.business.asset_or_class, f.business.field)) != str(abs(f.amount.source_signed_value))]


def _word(payload, text):
    return next(w for w in payload["analyzeResult"]["pages"][0]["words"] if w["content"] == text)


@pytest.mark.parametrize("widen", [False, True])
def test_s2_digits_straddling_two_columns_never_bind(stub, documents, monkeypatch, widen):
    monkeypatch.setenv(azure.KEY_ENV, "k"); payload = azure.recorded("scanned_p1"); r = payload["analyzeResult"]
    c12 = next(c for c in r["tables"][0]["cells"] if (c["rowIndex"], c["columnIndex"]) == (1, 2))
    r["pages"][0]["words"].remove(_word(payload, "1.000,00")); c12["content"] = ""; c12["spans"] = []
    if widen:
        c13 = next(c for c in r["tables"][0]["cells"] if (c["rowIndex"], c["columnIndex"]) == (1, 3))
        poly = c13["boundingRegions"][0]["polygon"]; poly[0] = poly[6] = 4.5
        p12 = c12["boundingRegions"][0]["polygon"]; p12[2] = p12[4] = 4.49
    word = _word(payload, "120,00"); y = word["polygon"][1]; h = word["polygon"][5] - y
    word["polygon"] = [4.55, y, 5.35, y, 5.35, y + h, 4.55, y + h]
    azure.LayoutStub.reset(payload)
    facts = _azure_facts(azure.acquire(documents / "azure_scanned.pdf", stub))
    truth = dict(AZURE_TRUTH); truth.pop(("000101", "ahk_open"))
    assert not _azure_wrong(facts, truth)


def test_k1_two_stacked_azure_tables_never_bind_one_key_twice(stub, documents, monkeypatch):
    held = _load(ROOT / "tests" / "unit" / "test_azure_review_held.py", "r5_held")
    monkeypatch.setenv(azure.KEY_ENV, "k"); payload = held._with_second_table(azure.recorded("scanned_p1"), header_text={}, drop_header=False)
    r = payload["analyzeResult"]; lower = r["tables"][1]
    for cell in lower["cells"]:
        if cell["rowIndex"] == 1 and cell["columnIndex"] == 6 and cell.get("spans"):
            offset = cell["spans"][0]["offset"]
            word = next(w for w in r["pages"][0]["words"] if w["span"]["offset"] == offset)
            word["content"] = cell["content"] = "900,00"
    azure.LayoutStub.reset(payload)
    facts = _azure_facts(azure.acquire(documents / "azure_scanned.pdf", stub))
    keys = collections.Counter((f.business.asset_or_class, f.business.field, f.business.period) for f in facts)
    assert not [k for k, n in keys.items() if n > 1] and not _azure_wrong(facts, AZURE_TRUTH)


# --- S5 local unit reader ---

@pytest.mark.parametrize(("unit_text", "expected"), [("in Tsd. EUR", "TEUR"), ("TEUR", "TEUR"), ("Mio. EUR", None), ("EUR", "EUR")])
def test_s5_local_unit_reader_shares_the_scaled_unit_rules(tmp_path, unit_text, expected):
    source = tmp_path / "unit.xlsx"; workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"
    sheet.append([f"Anlagenspiegel Geschäftsjahr 2025 {unit_text}"]); sheet.append(["Asset ID", "Opening cost 2025"]); sheet.append(["A-1", 12]); workbook.save(source)
    facts = json.loads(Path(run_input(source, tmp_path / "out", RunConfig()).evidence_path).read_text())["facts"]
    assert [f["amount"]["unit"] for f in facts] == ([expected] if expected else [])


# --- S6 Lfd. is a running counter ---

@pytest.mark.parametrize("header", ["Lfd.", "lfd. Nr."])
def test_s6_running_counter_is_never_a_sub_number(tmp_path, header):
    def mutate(ws):
        ws["B13"] = header
        for row, value in ((14, "1"), (15, "2"), (16, "3")):
            ws[f"B{row}"] = value
    facts = register(tmp_path, mutate)
    assert not [f for f in facts if f["asset_identity"].get("asset_sub_id") in {"1", "2", "3"}]
    assert not collisions(facts) and not adv.wrong(facts, adv.reg_truth())


# --- coverage C1-C4 (QA round 3 coverage probes) ---

NAMES = {
    "check-in": ("Check-in Terminal Halle 3", "CAD-Software Lizenz", "CAD-Software Lizenz Nachrüstung"),
    "differenzial": ("Finanzbuchhaltung Software", "Differenzialgetriebe Prüfstand", "CAD-Software Lizenz Nachrüstung"),
    "lt-fahrzeug": ("Finanzbuchhaltung Software", "CAD-Software Lizenz", "Lt.-Fahrzeug Anhänger"),
    "kontrolle": ("Kontrollsystem Zutritt", "Brandschutzkontrolle Anlage", "Laut-Sprecheranlage"),
    "total-station": ("Total Station Leica TS16", "CAD-Software Lizenz", "CAD-Software Lizenz Nachrüstung"),
    "gesamt-name": ("Gesamtsystem ERP", "Gesamt Lizenzpaket Office", "CAD-Software Lizenz Nachrüstung"),
}


@pytest.mark.parametrize("case", list(NAMES))
def test_c1_c2_asset_names_with_marker_or_total_words_still_bind(tmp_path, case):
    def mutate(ws):
        for row, name in zip((14, 15, 16), NAMES[case]):
            ws[f"D{row}"] = name
    facts = register(tmp_path, mutate, reg_plan())
    assert not adv.wrong(facts, adv.reg_truth()) and len(facts) == 39, len(facts)


@pytest.mark.parametrize(("header", "values"), [("ID", (1, 2, 3)), ("ID", ("1", "2", "3")), ("Kennung", (1, 2, 3))])
def test_c3_counter_column_does_not_hide_the_register(tmp_path, header, values):
    def mutate(ws):
        ws["V13"] = header
        for row, value in zip((14, 15, 16), values):
            ws[f"V{row}"] = value
    facts = register(tmp_path, mutate, reg_plan())
    assert not adv.wrong(facts, adv.reg_truth()) and len(facts) == 39, len(facts)


@pytest.mark.parametrize("title", ["Jahresabschluss Geschaftsjahr 2025", "Geschäftsjahr 2025 Anlagen"])
def test_c4_period_words_survive_lost_or_decomposed_umlauts(title):
    from carthagent.structure import PERIOD_WORDS
    assert PERIOD_WORDS.search(title)
