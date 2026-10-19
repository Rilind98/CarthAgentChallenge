"""Round 8: review round 7 should-fixes (R7-S1..S3) and QA round 7 items."""

from __future__ import annotations

from decimal import Decimal
import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from carthagent.mapper import MapperConfig
from carthagent.pipeline import RunConfig, run_input

ROOT = Path(__file__).resolve().parents[2]


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); return module


adv = _load(Path(__file__).with_name("test_option_b_adversarial.py"), "r8_adv")


def local_facts(tmp_path, source, out="local"):
    result = run_input(source, tmp_path / out, RunConfig())
    assert result.evidence_path is not None, result.diagnostics
    return json.loads(Path(result.evidence_path).read_text())["facts"]


# --- R8-1 an asset qualifier ("AfA-pflichtig", "ohne AfA") never turns an AHK movement into depreciation (local, end to end) ---

ROWS = [["100001", "Fräse", 1000, 100, 0, 1100], ["100002", "Presse", 2000, 0, 300, 1700], ["100003", "Bagger", 500, 250, 0, 750]]


@pytest.mark.parametrize("header", ["Zugänge", "Zugänge AfA-pflichtig", "Zugänge (AfA-fähig)", "Zugänge (ohne AfA)", "Zugänge nicht abschreibbar", "Zugänge (davon abschreibungsfähig)"])
def test_r8_1_qualified_addition_columns_stay_additions(tmp_path, header):
    source = tmp_path / "c.xlsx"; workbook = Workbook(); ws = workbook.active; ws.title = "Data"
    ws.append(["Anlagenverzeichnis Geschäftsjahr 2025 EUR"]); ws.append(["Anlage", "Bezeichnung", "AHK 01.01.2025", header, "Abgänge", "AHK 31.12.2025"])
    for row in ROWS:
        ws.append(row)
    workbook.save(source)
    fields = ["ahk_open", "zugang", "abgang", "ahk_close"]
    truth = {f"{'CDEF'[i]}{3 + o}": (field, ROWS[o][0], "2025", Decimal(ROWS[o][2 + i])) for o in range(3) for i, field in enumerate(fields)}
    facts = local_facts(tmp_path, source)
    wrong = [(f["pointer"]["cell"], f["field"]) for f in facts if truth.get(f["pointer"]["cell"]) != (f["field"], f["asset_identity"].get("asset_or_class"), f["asset_identity"].get("period"), Decimal(f["amount"]["normalized_base"]))]
    assert not wrong
    assert [f for f in facts if f["pointer"]["cell"][0] == "D"], "the additions column binds as zugang"


# --- R8-2 a record without a company code collides with a qualified record of the same projected row in the exporter ---

def test_r8_2_blank_company_code_never_binds_a_second_class_row():
    from carthagent.exporter import TemplateExportAdapter
    from carthagent.models import BusinessIdentity, PhysicalIdentity, ValidatedFact
    from carthagent.numeric import parse_amount
    def fact(identifier, cell, entity, position, field="ahk_open"):
        business = BusinessIdentity(entity=entity, asset_or_class="1000", class_code="1000", balance_position=position, granularity="class", row_kind="data", period="2025", field=field, currency="EUR", unit="EUR")
        return ValidatedFact(f"d:fact:{identifier}", "d", PhysicalIdentity("a" * 64, "d", f"S!{cell}"), business, parse_amount("10"), f"d:x:S!{cell}",
                             {"kind": "xlsx", "sheet": "S", "cell": cell, "observation_id": f"d:x:S!{cell}"}, (), {})
    blank = TemplateExportAdapter().plan_bindings("s", (fact("1", "B20", "7150", "2101200"), fact("2", "B25", None, "2101300")))
    assert [b.status for b in blank] == ["destination_collision", "destination_collision"] and all(b.destination.cell is None for b in blank)
    control = TemplateExportAdapter().plan_bindings("s", (fact("1", "B20", "7150", "2101200"), fact("2", "B25", "7307", "2101300")))
    assert [b.status for b in control] == ["bound", "bound"], "two printed company codes are two rows"


# --- R8-3 a rerun with another mapper outcome into the same output directory never hits a coverage collision ---

def test_r8_3_rerun_after_a_failed_mapper_republishes_without_collision(tmp_path):
    stub = _load(ROOT / "tests" / "fixtures" / "option_b_stub.py", "r8_stub")
    source = adv.fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx")
    failed = run_input(source, tmp_path / "out", RunConfig(mapper=MapperConfig(True, "http://127.0.0.1:1/v1/responses", "down", 1, 1)))
    assert failed.evidence_path is not None, failed.diagnostics
    first = json.loads(Path(failed.evidence_path).read_text())
    with stub.recorded_plan_server(stub.load_plan("sap_block_unlabelled")) as endpoint:
        second = run_input(source, tmp_path / "out", RunConfig(mapper=MapperConfig(True, endpoint, "local-stub", 2, 256)))
    assert second.evidence_path is not None and not [d for d in second.diagnostics if "collision" in d], second.diagnostics
    evidence = json.loads(Path(second.evidence_path).read_text())
    assert len(evidence["facts"]) == 26 and len(evidence["facts"]) > len(first["facts"])
    linked = {item["linked_record_id"] for item in evidence["coverage"] if item.get("linked_record_id")}
    assert {f["id"] for f in evidence["facts"]} <= linked, "coverage reflects the second run"


# --- R8-4 period inheritance refuses prior-year/parallel sheets and needs two non-zero tie-outs ---

def _inh(path, second_name, second_title=None, headers=("Anlage", "Bezeichnung", "Zugänge", "AHK Ende", "Währung"), zugang=0.0):
    workbook = Workbook(); first = workbook.active; first.title = "GJ"; second = workbook.create_sheet(second_name)
    first["A1"] = "Anlagenspiegel Geschäftsjahr 2025"; first["A2"] = "BK 7307"
    if second_title:
        second["A1"] = second_title
    for ws, header in ((first, ("Anlage", "Bezeichnung", "Zugänge", "AHK Ende", "Währung")), (second, headers)):
        for index, text in enumerate(header):
            ws.cell(4, index + 1, text)
        for row, (asset, name, value, add) in enumerate((("100001", "Grundstück A", 1000.0, zugang), ("100002", "Grundstück B", 2000.0, zugang))):
            for index, cell in enumerate((asset, name, add, value, "EUR")):
                ws.cell(5 + row, index + 1, cell)
        ws["A7"] = "Summe BK 7307"; ws["C7"] = 2 * zugang; ws["D7"] = 3000.0
    workbook.save(path); return path


@pytest.mark.parametrize(("second", "title", "headers", "zugang"), [
    ("Vorjahr", None, ("Anlagennummer", "Text", "Zugänge", "AHK Ende", "Währung"), 50.0),       # prior-year sheet name
    ("Tabelle2", "Vergleich", ("Anlagennummer", "Text", "Zugänge", "AHK Ende", "Währung"), 50.0),  # comparison title
    ("Tabelle2", None, ("Anlage", "Bezeichnung", "Zugänge", "AHK Ende", "Währung"), 50.0),        # repeats the dated header band
    ("Tabelle2", None, ("Anlagennummer", "Text", "Zugänge", "AHK Ende", "Währung"), 0.0),          # only one non-zero tie (AHK)
])
def test_r8_4_inheritance_refuses_prior_year_parallel_and_single_ties(tmp_path, second, title, headers, zugang):
    facts = local_facts(tmp_path, _inh(tmp_path / "inh.xlsx", second, title, headers, zugang))
    assert [f for f in facts if f["pointer"]["sheet"] == "GJ"], "control: the dated sheet binds"
    assert not [f for f in facts if f["pointer"]["sheet"] == second and f["asset_identity"].get("period")]


def test_r8_4_detail_sheet_with_two_non_zero_ties_still_inherits(tmp_path):
    facts = local_facts(tmp_path, _inh(tmp_path / "inh.xlsx", "Tabelle2", None, ("Anlagennummer", "Text", "Zugänge", "AHK Ende", "Währung"), 50.0))
    second = [f for f in facts if f["pointer"]["sheet"] == "Tabelle2"]
    assert second and {f["asset_identity"].get("period") for f in second} == {"2025"}


# --- R8-6 a header naming several movements never lends weight to a single-movement label (swap stays caught) ---

@pytest.mark.parametrize(("c16", "e16"), [("Zugang/Umbuchung", "Umbuchung/Zugang"), ("Zugang", "Umbuchung")])
def test_r8_6_swapped_movement_labels_under_combined_headers_bind_nothing_wrong(tmp_path, c16, e16):
    source = adv.fixtures.sap_block_unlabelled(tmp_path / "sap.xlsx")
    workbook = load_workbook(source); workbook["Anlagengi" "tter"]["C16"] = c16; workbook["Anlagengi" "tter"]["E16"] = e16; workbook.save(source)
    recorded = adv.plan("sap_block_unlabelled")
    for item in recorded["plans"][0]["field_headers"]:
        if item["observation_id"].endswith("!C16"):
            item["semantic_field"] = "umb"
        elif item["observation_id"].endswith("!E16"):
            item["semantic_field"] = "zugang"
    _, facts = adv.run(tmp_path, source, recorded)
    assert not adv.wrong(facts, adv.sap_truth())
