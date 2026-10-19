"""Split-request harness (ported from QA probes test_reqa_m3.py / test_round2_split.py).

Dynamic loopback stubs write an HONEST plan from whatever rows each request carries; the pipeline splits a long
SAP sheet into several requests. Truth comes from build_option_b constants.
"""
from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import importlib.util
import json
from pathlib import Path
from threading import Thread

from openpyxl import Workbook

from carthagent.mapper import MapperConfig
from carthagent.pipeline import RunConfig, run_input
from carthagent.policies import DEFAULT_BOUNDS

_spec = importlib.util.spec_from_file_location("split_fixtures", Path(__file__).with_name("build_option_b.py"))
fixtures = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(fixtures)
SAP_FIELDS = (("ahk_open", "zugang", "abgang", "umb", None, None, "ahk_close"),
              ("kum_afa_open", "total_afa_gj", "afa_abgang", "afa_umb", None, "zuschreibung", "kum_afa_close"),
              ("bw_prev", None, None, None, None, None, "bw_close"))
SAP_FIELD_LABELS = {16: dict(zip("BCDEH", ("ahk_open", "zugang", "abgang", "umb", "ahk_close"))),
                    17: dict(zip("BCDEGH", ("kum_afa_open", "total_afa_gj", "afa_abgang", "afa_umb", "zuschreibung", "kum_afa_close"))),
                    18: dict(zip("BH", ("bw_prev", "bw_close")))}


def empty_plan(region):
    return {"region_observation_id": region, "table_type": "class_grid_matrix", "header_rows": [], "record_blocks": [], "column_spans": [],
            "field_headers": [], "identity_headers": [], "measure_axes": [], "movement_axes": [], "row_roles": [], "period_contexts": [], "unit_contexts": []}


def respond(handler, plan):
    text = json.dumps({"plans": [plan]})
    body = json.dumps({"status": "completed", "output": [{"type": "message", "content": [{"type": "output_text", "text": text}]}]}).encode()
    handler.send_response(200); handler.send_header("Content-Type", "application/json"); handler.send_header("Content-Length", str(len(body))); handler.end_headers(); handler.wfile.write(body)


def payload_of(handler):
    request = json.loads(handler.rfile.read(int(handler.headers["Content-Length"])))
    observations = json.loads(request["input"][1]["content"][0]["text"])["observations"]
    rows = {o["row"]: o["id"] for o in observations if o["kind"] == "logical_row"}
    cells = {(o["row"], o["column"]): o for o in observations if o["kind"] in {"cell", "empty_cell"}}
    region = next(o["id"] for o in observations if o["kind"] == "table_region")
    return rows, cells, region


def build_two_tables(path, a_blocks, abbreviated, b_blocks=70):
    """Table A (header rows 15-18) then table B with its own band and the Zugang/Umbuchung columns swapped."""
    workbook = Workbook(); ws = workbook.active; ws.title = "G"
    ws["A2"] = "Berichtsdatum: 31.12.2025"
    truth: dict[tuple[str, str], tuple] = {}

    def table(header_top, blocks, swap, code0):
        header = [dict(row) for row in fixtures.SAP_HEADER]
        if swap:
            header[1]["C"], header[1]["E"] = header[1]["E"], header[1]["C"]; header[2]["C"], header[2]["E"] = header[2]["E"], header[2]["C"]
            if abbreviated:
                header[1]["C"], header[1]["E"], header[2]["C"], header[2]["E"] = "Umb.", "Zug.", "AfA Umb.", "AfA lfd. J."
        for offset, cells in enumerate(header):
            for column, text in cells.items():
                ws[f"{column}{header_top + offset}"] = text
        row = header_top + 4
        for index in range(blocks):
            identity, *measures = fixtures.SAP_BLOCKS[index % 2]; identity = dict(identity); identity["E"] = str(code0 + index)
            measures = [list(m) for m in measures]
            if index % 2 == 0:
                measures[0][3] = 777.0; measures[0][6] = measures[0][6] + 777.0
            names = [list(n) for n in SAP_FIELDS]
            if swap:
                for k in (0, 1):
                    measures[k][1], measures[k][3] = measures[k][3], measures[k][1]; names[k][1], names[k][3] = names[k][3], names[k][1]
            for column, text in identity.items():
                ws[f"{column}{row}"] = text
            for k, (values, labels) in enumerate(zip(measures, names), start=1):
                for column, value, field in zip("BCDEFGH", values, labels):
                    if value is not None:
                        ws[f"{column}{row + k}"] = value
                    if field:
                        truth[("G", f"{column}{row + k}")] = (field, identity["E"], "2025", Decimal(str(value)))
            row += 5
        return row

    end_a = table(15, a_blocks, False, 1000)
    table(end_a + 2, b_blocks, True, 5000)
    workbook.save(path)
    return path, end_a + 2, truth


class NearestBand(BaseHTTPRequestHandler):
    """Govern each record of table B by the nearest VISIBLE 4-row header band above it; label header cells by their words."""
    bands: tuple[int, ...] = (15,)
    pieces: list = []
    pick = staticmethod(max)  # which visible band governs table B (min = deliberately the OTHER table's band)

    def do_POST(self):
        rows, cells, region = payload_of(self)
        visible = [band for band in type(self).bands if all(x in rows for x in range(band, band + 4))]
        type(self).pieces.append({"rows": (min(rows), max(rows)), "bands": visible})
        plan = empty_plan(region)
        visible = [type(self).pick(visible)] if visible else []
        for band in visible:
            plan["header_rows"] += [rows[x] for x in range(band, band + 4)]
            for offset, labels in SAP_FIELD_LABELS.items():
                for column, field in labels.items():
                    cell = cells[(band + offset - 15, ord(column) - 64)]
                    field = {"Zugang": "zugang", "Umbuchung": "umb", "AfA des Jahres": "total_afa_gj", "AfA Umbuchung": "afa_umb"}.get(cell.get("text"), field)
                    plan["field_headers"].append({"observation_id": cell["id"], "semantic_field": field, "parent_header_ids": []})
            for column, role in zip("BCDE", ("entity", "balance_position", "account", "class_code")):
                plan["identity_headers"].append({"observation_id": cells[(band, ord(column) - 64)]["id"], "identity_role": role})
            for row in sorted(rows):
                if row < type(self).bands[-1] + 4 or (row - type(self).bands[-1] - 4) % 5 or not all(x in rows for x in (row, row + 1, row + 2, row + 3)):
                    continue
                plan["record_blocks"].append({"identity_row_id": rows[row], "rows": [{"row_id": rows[row + k], "header_row_id": rows[band + k], "measure": m}
                                              for k, m in ((1, "ahk"), (2, "cum_depreciation"), (3, "book_value"))], "role": "data"})
            plan["period_contexts"] = [{"observation_id": cells[(2, 1)]["id"], "temporal_role": "current_period", "scope": "sheet"}] if (2, 1) in cells else []
            plan["unit_contexts"] = [{"observation_id": cells[(int(b["identity_row_id"].rsplit(":", 1)[1]), 8)]["id"], "unit": "EUR", "scope": "block"} for b in plan["record_blocks"]]
        respond(self, plan)

    def log_message(self, *args):
        return


def serve(handler, source, out, chunk=120):
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler); thread = Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        bounds = replace(DEFAULT_BOUNDS, row_chunk_size=chunk, row_overlap=20, max_mapper_requests=64)
        return run_input(source, out, RunConfig(bounds=bounds, mapper=MapperConfig(True, f"http://127.0.0.1:{server.server_port}/v1/responses", "local-stub", 2, 256)))
    finally:
        server.shutdown(); thread.join()
