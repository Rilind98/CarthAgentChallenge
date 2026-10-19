"""Synthetic layout fixtures for option B (model-declared structure verified by code).

Shapes follow the corpus families f05 (SAP class grid), f08 (two-year Spiegel) and
f09 (flat SAP register). Every number is synthetic and roll-forward consistent.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

SAP_HEADER = (
    {"B": "BuKr GsBe", "C": "BilanzPos", "D": "BestKtoAHK", "E": "Klasse", "H": "Währg"},
    {"B": "AHK GJ-Beg", "C": "Zugang", "D": "Abgang", "E": "Umbuchung", "F": "Nachaktivg", "G": "Invest.Förderg", "H": "aktuelle AHK"},
    {"B": "AfA GJ-Beg", "C": "AfA des Jahres", "D": "AfA Abgang", "E": "AfA Umbuchung", "F": "AfA Nachaktivg", "G": "Zuschreibungen", "H": "kumulierte AfA"},
    {"B": "Buchwrt GJ-Beg", "H": "lfd Buchwert"},
)
# identity row cells, then AHK / AfA / book-value measure rows (columns B..H; None = blank)
SAP_BLOCKS = (
    ({"B": "7150", "C": "2101200", "D": "0027", "E": "1000", "F": "EDV-Software", "H": "EUR"},
     (100000.00, 5000.50, 0, 0, 0, 0, 105000.50), (-80000.00, -10000.25, 0, 0, 0, 0, -90000.25), (20000.00, None, None, None, None, None, 15000.25)),
    ({"B": "7150", "C": "2101300", "D": "0011", "E": "2000", "F": "Grund und Boden", "H": "EUR"},
     (300000.00, 0, -20000.00, 0, 0, 0, 280000.00), (-10000.00, -2000.00, 1000.00, 0, 0, 0, -11000.00), (290000.00, None, None, None, None, None, 269000.00)),
)
SAP_SUBTOTAL = ({"A": "*", "B": "7150", "H": "EUR"},
                (400000.00, 5000.50, -20000.00, 0, 0, 0, 385000.50), (-90000.00, -12000.25, 1000.00, 0, 0, 0, -101000.25), (310000.00, None, None, None, None, None, 284000.25))


def sap_block_unlabelled(path: Path, *, class_cell: object = None) -> Path:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Anlagengitter"
    preamble = {1: ("Sortierkriterien", "Aufst.", "Abst.", "Zwischensumme"), 2: ("Buchungskreis", "X", None, "X"), 3: ("Geschäftsbereich", "X", None, "X"),
                4: ("Bilanzposition", "X", None, "X"), 5: ("Bestandskonto AHK", "X", None, "X"), 6: ("Anlagenklasse", "X"), 8: ("Datenstatistik", "Anzahl"),
                9: ("Übergebene Sätze", 2), 10: ("Ermittelte Summenzeilen", 1),
                12: ("Berichtsdatum: 31.12.2025   Anlagengitter  Bereich 01 Handelsbilanz",),
                13: ("Erstellungsdatum: 26.04.2026   Gitter HGB §268 Abs.2 - 13 Spalten (breit, unvollständig)",)}
    for row, values in preamble.items():
        for column, value in zip("ABCD", values):
            if value is not None:
                sheet[f"{column}{row}"] = value
    for offset, header in enumerate(SAP_HEADER):
        for column, text in header.items():
            sheet[f"{column}{15 + offset}"] = text
    row = 19
    for identity, *measures in (*SAP_BLOCKS, SAP_SUBTOTAL):
        for column, text in identity.items():
            sheet[f"{column}{row}"] = text
        for index, values in enumerate(measures, start=1):
            for column, value in zip("BCDEFGH", values):
                if value is not None:
                    sheet[f"{column}{row + index}"] = value
        row += 5
    if class_cell is not None:
        sheet["E19"] = class_cell
    workbook.save(path); return path


SPIEGEL_ROWS = {  # label row -> (number, name, AHK open, Zugänge, Abgänge, Umbuchungen, AHK close, AfA open, AfA Zugänge, AfA Abgänge, AfA close, BW close, BW prior)
    2025: {10: ("2.", "Entgeltlich erworbene Konzessionen und Lizenzen", 1000.00, 200.00, 0, 0, 1200.00, 600.00, 100.00, 0, 700.00, 500.00, 400.00),
           15: ("1.", "Grundstücke und Bauten", 5000.00, 0, 0, 0, 5000.00, 0, 0, 0, 0, 5000.00, 5000.00),
           16: ("2.", "Technische Anlagen und Maschinen", 3000.00, 500.00, 100.00, 50.00, 3450.00, 1000.00, 300.00, 80.00, 1220.00, 2230.00, 2000.00)},
    2024: {10: ("2.", "Entgeltlich erworbene Konzessionen und Lizenzen", 900.00, 100.00, 0, 0, 1000.00, 500.00, 100.00, 0, 600.00, 400.00, 400.00),
           15: ("1.", "Grundstücke und Bauten", 5000.00, 0, 0, 0, 5000.00, 0, 0, 0, 0, 5000.00, 5000.00),
           16: ("2.", "Technische Anlagen und Maschinen", 2800.00, 200.00, 0, 0, 3000.00, 700.00, 300.00, 0, 1000.00, 2000.00, 2100.00)},
}
SPIEGEL_VALUE_COLUMNS = ("C", "D", "E", "F", "H", "J", "K", "L", "M", "O", "P")


def spiegel_column(column: str, year: int) -> str:
    """Tabelle1 (2025) starts its amounts in column C; Tabelle2 (2024) is shifted one column left."""
    return column if year == 2025 else chr(ord(column) - 1)


def two_year_two_sheets(path: Path) -> Path:
    workbook = Workbook(); first = workbook.active; first.title = "Tabelle1"
    for sheet, year in ((first, 2025), (workbook.create_sheet("Tabelle2"), 2024)):
        col = lambda name: spiegel_column(name, year)  # noqa: E731
        label = col("B")
        sheet[f"{label}1"] = "Entwicklung des Anlagevermögens"
        sheet[f"{label}2"] = f"Wirtschaftsjahr vom 01.01.{year} bis 31.12.{year}"
        for start, end, text in (("C", "H", "Anschaffungs- und Herstellungskosten"), ("J", "M", "kumulierte Abschreibungen"), ("O", "P", "Buchwerte")):
            sheet[f"{col(start)}4"] = text; sheet.merge_cells(f"{col(start)}4:{col(end)}4")
        sheet[f"{col('S')}4"] = "lt. Bilanz"; sheet[f"{col('T')}4"] = "Differenz"
        for column, value in (("C", datetime(year, 1, 1)), ("D", "Zugänge"), ("E", "Abgänge"), ("F", "Umbuchungen"), ("H", datetime(year, 12, 31)),
                              ("J", datetime(year, 1, 1)), ("K", "Zugänge"), ("L", "Abgänge"), ("M", datetime(year, 12, 31)),
                              ("O", datetime(year, 12, 31)), ("P", datetime(year - 1, 12, 31)), ("S", "(aus Bilanz)")):
            sheet[f"{col(column)}5"] = value
        for column in (*SPIEGEL_VALUE_COLUMNS, "S", "T"):
            sheet[f"{col(column)}6"] = "EUR"
        for row, numeral, text in ((8, "I.", "Immaterielle Vermögensgegenstände"), (13, "II.", "Sachanlagen")):
            if year == 2025:
                sheet[f"A{row}"] = numeral; sheet[f"B{row}"] = text
            else:
                sheet[f"A{row}"] = f"{numeral} {text}"
        for row, (index, name, *values) in SPIEGEL_ROWS[year].items():
            if year == 2025:
                sheet[f"A{row}"] = index; sheet[f"B{row}"] = name
            else:
                sheet[f"A{row}"] = f"{index} {name}"
            for column, value in zip(SPIEGEL_VALUE_COLUMNS, values):
                sheet[f"{col(column)}{row}"] = value
            sheet[f"{col('S')}{row}"] = values[-2]; sheet[f"{col('T')}{row}"] = 0
        for subtotal, members in ((11, (10,)), (17, (15, 16))):
            for index, column in enumerate(SPIEGEL_VALUE_COLUMNS):
                sheet[f"{col(column)}{subtotal}"] = sum(SPIEGEL_ROWS[year][member][2 + index] for member in members)
        if year == 2025:
            # Unlabelled grand total, then a reconciliation section below the table (shape of the real f08, rows 26-39).
            for index, column in enumerate(SPIEGEL_VALUE_COLUMNS):
                sheet[f"{column}19"] = sum(values[2 + index] for values in SPIEGEL_ROWS[year].values())
            for row, label, cells in RECONCILIATION:
                sheet[f"B{row}"] = label
                for column, value in cells:
                    sheet[f"{column}{row}"] = value
    workbook.save(path); return path


RECONCILIATION = (  # (row, label, [(column, value)]) - never records: they restate or check totals
    (21, "Abstimmung Vorjahreswerte", ()),
    (22, "AHK 01.01.2025 lt. Anlagengitter Vorjahr (AV 2)", (("C", 9000.00), ("T", 0))),
    (23, "kum. Abschreibungen 01.01.2025 lt. Anlagengitter Vorjahr (AV 2)", (("J", 1600.00), ("T", 0))),
    (25, "Abstimmung aktuelle Werte Anlagegitter", ()),
    (26, "Zugänge lt. Anlagegitter", (("D", 700.00), ("T", 0))),
    (27, "Abgänge lt. Anlagegitter", (("E", 100.00), ("T", 0))),
    (28, "Abschreibungen lt. Anlagegitter", (("K", 400.00), ("T", 0))),
    (30, "Bearbeitungsstand 21.02.2026: Werte lt. vorläufiger Bilanz", ()),
    (31, "Differenz Zugänge: GWG-Sammelposten im Anlagegitter nicht enthalten", (("D", 12.50),)),
)


REGISTER_HEADER = ("Anlage", "Unter-nummer", "Aktivierung am", "Anlagenbezeichnung", "Währ-ung", "AHK GJ-Beg", "Zugang", "Abgang", "Umbuchung", "Nachaktivg",
                   "Invest.Förderg", "aktuelle AHK", "AfA GJ-Beg", "AfA des Jahres", "AfA Abgang", "AfA Umbuchung", "AfA Nachaktivg", "Zuschreibungen",
                   "kumulierte AfA", "Buchwrt GJ-Beg", "lfd Buchwert")
REGISTER_ROWS = (
    ("100001001", "0", datetime(2018, 7, 2), "Finanzbuchhaltung Software", "EUR", 1000.00, 0, 0, 0, 0, 0, 1000.00, -999.00, 0, 0, 0, 0, 0, -999.00, 1.00, 1.00),
    ("100001002", "0", datetime(2021, 3, 19), "CAD-Software Lizenz", "EUR", 2000.00, 100.00, 0, 0, 0, 0, 2100.00, -1500.00, -200.00, 0, 0, 0, 0, -1700.00, 500.00, 400.00),
    ("100001002", "1", datetime(2022, 11, 18), "CAD-Software Lizenz Nachrüstung", "EUR", 500.00, 0, -100.00, 0, 0, 0, 400.00, -300.00, -50.00, 60.00, 0, 0, 0, -290.00, 200.00, 110.00),
)


def flat_sap_register(path: Path) -> Path:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Sheet1"
    sheet["A4"] = "Jahresabschluss per 30.06.2025"; sheet["A6"] = "BK: 7307"; sheet["A9"] = "SUMMEN"
    for index in range(5, len(REGISTER_HEADER)):
        sheet.cell(9, index + 1, sum(row[index] for row in REGISTER_ROWS))
    for index, text in enumerate(REGISTER_HEADER):
        sheet.cell(13, index + 1, text)
    for offset, values in enumerate(REGISTER_ROWS):
        for index, value in enumerate(values):
            sheet.cell(14 + offset, index + 1, value)
    workbook.save(path); return path


def generate(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    return {name: builder(directory / name) for name, builder in (
        ("sap_block_unlabelled.xlsx", sap_block_unlabelled), ("two_year_two_sheets.xlsx", two_year_two_sheets), ("flat_sap_register.xlsx", flat_sap_register))}
