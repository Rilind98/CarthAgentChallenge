"""Independent adversarial probes against the public CLI, artifacts, and ledger.

These tests deliberately avoid answer keys and generator imports.  Network probes
use loopback servers only.  Production modules are never patched or edited.
"""
from __future__ import annotations

from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
from threading import Thread
import zipfile

from openpyxl import Workbook
import pytest
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
PRIVATE = ROOT / "hackathon-local"
TEMPLATE = PRIVATE / "templates" / "Arbeitsversion-merged-layout-20.xlsx"
CORPUS = PRIVATE / "ElabHackathon" / "A3-Hackathon-Corpus"
V3_OUTPUT = PRIVATE / "corpus-xlsx-checkpoint-v3"


def run_cli(
    source: Path,
    output: Path,
    *,
    config: Path | None = None,
    template: Path | None = None,
    cwd: Path = ROOT,
    env: dict[str, str] | None = None,
    timeout: int = 90,
) -> tuple[subprocess.CompletedProcess[str], list[dict[str, object]]]:
    command = [sys.executable, "-m", "carthagent.cli", "run", str(source), "--output", str(output)]
    if config is not None:
        command.extend(("--config", str(config)))
    if template is not None:
        command.extend(("--template", str(template)))
    completed = subprocess.run(command, cwd=cwd, capture_output=True, text=True, timeout=timeout, env=env)
    records = [json.loads(line) for line in completed.stdout.splitlines() if line.startswith("{")]
    assert records, (completed.returncode, completed.stdout, completed.stderr)
    return completed, records


def evidence(record: dict[str, object]) -> dict[str, object]:
    path = record.get("evidence_path")
    assert isinstance(path, str) and path
    return json.loads(Path(path).read_text())


def simple_xlsx(path: Path, headers: list[object], rows: list[list[object]], *, title: str = "Schedule") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_persisted_ledger_graph_closure_invalidation_and_resume(tmp_path: Path) -> None:
    source = simple_xlsx(tmp_path / "source.xlsx", ["Asset ID", "Opening cost EUR 2025"], [["A-1", 12]])
    output = tmp_path / "out"
    first_completed, first_records = run_cli(source, output)
    assert first_completed.returncode == 0
    first = first_records[0]
    first_data = evidence(first)
    assert first_data["facts"], first_data["diagnostics"]

    database = output / ".carthagent" / "ledger.sqlite3"
    connection = sqlite3.connect(database)
    record_ids = {row[0] for row in connection.execute("SELECT id FROM records")}
    record_kinds = {row[0]: row[1] for row in connection.execute("SELECT id, kind FROM records")}
    dependencies = list(connection.execute("SELECT record_id, dependency_id, axis FROM dependencies"))
    assert dependencies
    assert all(record_id in record_ids and dependency_id in record_ids for record_id, dependency_id, _ in dependencies)

    fact_id = first_data["facts"][0]["id"]
    observation_id = first_data["facts"][0]["pointer"]["observation_id"]
    fact_dependencies = {dependency for record, dependency, _ in dependencies if record == fact_id}
    assert observation_id in fact_dependencies
    assert any(record == first_data["snapshot"]["id"] and dependency == fact_id for record, dependency, _ in dependencies)
    assert any("binding" in record and dependency == fact_id for record, dependency, _ in dependencies)

    downstream = {
        row[0]
        for row in connection.execute(
            """
            WITH RECURSIVE downstream(id) AS (
              SELECT record_id FROM dependencies WHERE dependency_id = ?
              UNION
              SELECT d.record_id FROM dependencies d JOIN downstream x ON d.dependency_id = x.id
            )
            SELECT id FROM downstream
            """,
            (observation_id,),
        )
    }
    assert fact_id in downstream
    assert any(record_kinds.get(value) == "snapshot" for value in downstream)
    assert any(record_kinds.get(value) == "output_binding" for value in downstream)
    assert any(record_kinds.get(value) == "export" for value in downstream)
    for record_id in downstream:
        connection.execute(
            "INSERT INTO events(record_id,event,payload) VALUES(?,?,?)",
            (record_id, "invalidated", json.dumps({"dependency_id": observation_id, "axis": "semantic"})),
        )
    connection.commit()
    connection.close()

    rebuilt_completed, rebuilt_records = run_cli(source, output)
    assert rebuilt_completed.returncode == 0
    assert rebuilt_records[0]["diagnostics"] != ["resumed_verified_publication"]
    resumed_completed, resumed_records = run_cli(source, output)
    assert resumed_completed.returncode == 0
    assert resumed_records[0]["diagnostics"] == ["resumed_verified_publication"]


def test_cross_version_v3_publication_and_cache_are_not_reused(tmp_path: Path) -> None:
    source = CORPUS / "f00_demo" / "f00_demo_01.xlsx"
    old_directory = next(iter(V3_OUTPUT.glob("f00_demo_01-*")), None)
    if not source.exists() or old_directory is None or not (V3_OUTPUT / ".carthagent" / "ledger.sqlite3").exists():
        pytest.skip("private v3 checkpoint unavailable")
    copied = tmp_path / "v3-output"
    shutil.copytree(V3_OUTPUT, copied)
    old_manifest = json.loads((copied / old_directory.name / "manifest.v1.json").read_text())
    completed, records = run_cli(source, copied, timeout=180)
    assert completed.returncode == 0, records
    current = records[0]
    assert current["diagnostics"] != ["resumed_verified_publication"]
    assert current["run_id"] != old_manifest["run_id"]
    current_manifest = json.loads(Path(current["manifest_path"]).read_text())
    assert current_manifest["run_id"] == current["run_id"]
    assert current_manifest["readback_verified"] is True


def test_template_export_never_creates_destination_locator_files(tmp_path: Path) -> None:
    if not TEMPLATE.exists():
        pytest.skip("private layout-20 template unavailable")
    source = simple_xlsx(
        tmp_path / "source.xlsx",
        ["Asset ID", "Name", "AHK 01.01 EUR", "AHK 31.12 EUR", "Buchwert Vorjahr EUR"],
        [["A-1", "Machine", 100, 100, 100]],
    )
    repo_before = {item.name for item in ROOT.iterdir() if item.is_file()}
    sandbox_before = {item.name for item in tmp_path.iterdir() if item.is_file()}
    completed, records = run_cli(source, tmp_path / "out", template=TEMPLATE, cwd=tmp_path)
    assert completed.returncode == 0 and records[0]["status"] in {"complete", "partial"}
    repo_after = {item.name for item in ROOT.iterdir() if item.is_file()}
    sandbox_after = {item.name for item in tmp_path.iterdir() if item.is_file()}
    assert not {name for name in repo_after - repo_before if "!" in name}
    assert not {name for name in sandbox_after - sandbox_before if "!" in name}


def test_template_clears_all_unavailable_active_row_closing_formulas(tmp_path: Path) -> None:
    if not TEMPLATE.exists():
        pytest.skip("private layout-20 template unavailable")
    source = simple_xlsx(
        tmp_path / "opening-only.xlsx",
        ["Asset ID", "Name", "AHK 01.01.2025 EUR"],
        [["A-1", "Machine", 100]],
    )
    completed, records = run_cli(source, tmp_path / "out", template=TEMPLATE)
    assert completed.returncode == 0
    from openpyxl import load_workbook

    sheet = load_workbook(records[0]["review_path"], data_only=False)["Anlagenverzeichnis"]
    row = next(index for index in range(2, 502) if str(sheet.cell(index, 1).value) == "A-1")
    for column in (27, 34, 35, 36):  # AA, AH, AI, AJ
        assert sheet.cell(row, column).value is None, sheet.cell(row, column).coordinate


def test_register_capacity_never_silently_drops_source_facts(tmp_path: Path) -> None:
    if not TEMPLATE.exists():
        pytest.skip("private layout-20 template unavailable")
    rows = [[f"A-{index:04d}", f"Asset {index}", index] for index in range(1, 502)]
    source = simple_xlsx(tmp_path / "register-overflow.xlsx", ["Asset ID", "Name", "Opening cost EUR 2025"], rows)
    completed, records = run_cli(source, tmp_path / "out", template=TEMPLATE, timeout=180)
    assert completed.returncode == 0
    data = evidence(records[0])
    facts = [item for item in data["facts"] if item["field"] == "ahk_open"]
    assert len(facts) == 501
    bindings = [item for item in data["output_bindings"] if item["record_id"] in {fact["id"] for fact in facts}]
    assert len(bindings) == 501
    statuses = Counter(item["status"] for item in bindings)
    if statuses.get("projection_overflow"):
        assert records[0]["status"] == "partial"
        assert statuses["projection_overflow"] == 1
    else:
        assert statuses == {"bound": 501}


def test_two_concurrent_cli_runs_publish_one_verified_pair(tmp_path: Path) -> None:
    source = simple_xlsx(tmp_path / "source.xlsx", ["Asset ID", "Opening cost EUR"], [["A-1", 12]])
    output = tmp_path / "out"
    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [executor.submit(run_cli, source, output) for _ in range(2)]
        results = [future.result() for future in futures]
    assert all(completed.returncode == 0 for completed, _ in results), [items for _, items in results]
    records = [items[0] for _, items in results]
    assert len({record["output_dir"] for record in records}) == 1
    manifest_path = Path(records[0]["manifest_path"])
    manifest = json.loads(manifest_path.read_text())
    directory = manifest_path.parent
    assert manifest["readback_verified"] is True
    assert (directory / ".complete").read_text().strip() == manifest["snapshot_digest"]
    import hashlib

    for name, digest in manifest["files"].items():
        assert hashlib.sha256((directory / name).read_bytes()).hexdigest() == digest


def test_ooxml_compression_bomb_is_bounded_before_parsing(tmp_path: Path) -> None:
    source = tmp_path / "compressed.xlsx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 250_000)
    assert source.stat().st_size < 100_000
    config = tmp_path / "bounds.toml"
    config.write_text("[bounds]\nmax_input_bytes = 100000\n")
    completed, records = run_cli(source, tmp_path / "out", config=config)
    assert completed.returncode != 0
    assert records[0]["status"] == "failed"
    assert any("archive" in str(item) and "limit" in str(item) for item in records[0]["diagnostics"])
    assert records[0]["output_dir"] is None


def test_timezone_does_not_change_facts_or_snapshot(tmp_path: Path) -> None:
    source = simple_xlsx(
        tmp_path / "source.xlsx",
        ["Asset ID", "Opening cost EUR", "As of"],
        [["A-1", 12, date(2024, 2, 29)]],
    )
    payloads = []
    for zone in ("UTC", "Europe/Berlin"):
        environment = dict(os.environ)
        environment["TZ"] = zone
        completed, records = run_cli(source, tmp_path / zone.replace("/", "-"), env=environment)
        assert completed.returncode == 0
        payloads.append(evidence(records[0]))
    assert payloads[0]["facts"] == payloads[1]["facts"]
    assert payloads[0]["snapshot"]["digest"] == payloads[1]["snapshot"]["digest"]


class MapperPlanHandler(BaseHTTPRequestHandler):
    calls = 0
    authorizations: list[str | None] = []

    def do_POST(self) -> None:
        type(self).calls += 1
        type(self).authorizations.append(self.headers.get("Authorization"))
        request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        payload = json.loads(request["input"][1]["content"][0]["text"])
        observations = payload["observations"]
        by_id = {item["id"]: item for item in observations}
        region = next(item for item in observations if item["kind"] == "table_region")
        rows = [by_id[item] for item in region["member_ids"] if item in by_id]
        numeric = next(item for item in observations if item.get("numeric_value") and item.get("column") == 1)
        row = next(item for item in rows if numeric["id"] in item["member_ids"])
        asset_header = next(item for item in observations if item.get("text") == "Asset ID")
        name_header = next(item for item in observations if item.get("text") == "Name")
        unit_context = next(item for item in observations if item.get("text") == "Currency EUR")
        plan = {
            "plans": [
                {
                    "region_observation_id": region["id"],
                    "field_headers": [
                        {"observation_id": asset_header["id"], "semantic_field": "ahk_open", "parent_header_ids": []}
                    ],
                    "identity_headers": [{"observation_id": name_header["id"], "identity_role": "asset_id"}],
                    "row_roles": [{"observation_id": row["id"], "row_role": "data"}],
                    "period_contexts": [],
                    "unit_contexts": [{"observation_id": unit_context["id"], "unit": "EUR"}],
                }
            ]
        }
        body = json.dumps(
            {
                "status": "completed",
                "output": [{"type": "message", "content": [{"type": "output_text", "text": json.dumps(plan)}]}],
            }
        ).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


class RedirectTargetHandler(BaseHTTPRequestHandler):
    calls = 0
    authorizations: list[str | None] = []

    def do_POST(self) -> None:
        type(self).calls += 1
        type(self).authorizations.append(self.headers.get("Authorization"))
        self.send_response(204)
        self.end_headers()

    def log_message(self, format: str, *args: object) -> None:
        return


class RedirectHandler(BaseHTTPRequestHandler):
    calls = 0
    authorizations: list[str | None] = []
    target = ""

    def do_POST(self) -> None:
        type(self).calls += 1
        type(self).authorizations.append(self.headers.get("Authorization"))
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(307)
        self.send_header("Location", type(self).target)
        self.send_header("Content-Length", "0")
        self.end_headers()

    def log_message(self, format: str, *args: object) -> None:
        return


class UnavailableHandler(BaseHTTPRequestHandler):
    calls = 0

    def do_POST(self) -> None:
        type(self).calls += 1
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(500)
        self.send_header("Content-Length", "0")
        self.end_headers()

    def log_message(self, format: str, *args: object) -> None:
        return


class local_server:
    def __init__(self, handler: type[BaseHTTPRequestHandler]):
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        self.thread = Thread(target=self.server.serve_forever, daemon=True)

    def __enter__(self) -> str:
        self.thread.start()
        return f"http://127.0.0.1:{self.server.server_port}/v1/responses"

    def __exit__(self, exc_type, exc, traceback) -> None:
        self.server.shutdown()
        self.thread.join()


def mapper_config(path: Path, endpoint: str) -> Path:
    path.write_text(
        "\n".join(
            (
                "[mapper]",
                "enabled = true",
                f'endpoint = "{endpoint}"',
                'model = "local-adversarial-stub"',
                "timeout_seconds = 2",
                "max_output_tokens = 2048",
                "reasoning_effort = \"low\"",
                "",
            )
        )
    )
    return path


def test_numeric_identity_and_metadata_cannot_become_model_proposed_money_or_zero(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Schedule"
    sheet.append(["Currency EUR"])
    sheet.append(["Asset ID", "Name", "Purchase date", "Useful life"])
    sheet.append(["000101", "Machine", date(2025, 1, 31), 10])
    source = tmp_path / "identity-only.xlsx"
    workbook.save(source)
    MapperPlanHandler.calls = 0
    MapperPlanHandler.authorizations = []
    with local_server(MapperPlanHandler) as endpoint:
        config = mapper_config(tmp_path / "mapper.toml", endpoint)
        environment = {key: value for key, value in os.environ.items() if key != "OPENAI_API_KEY"}
        environment["OPENAI_API_KEY"] = "dummy-must-not-be-sent"
        completed, records = run_cli(source, tmp_path / "out", config=config, env=environment)
    assert completed.returncode == 0
    data = evidence(records[0])
    assert MapperPlanHandler.calls >= 1
    assert MapperPlanHandler.authorizations == [None] * MapperPlanHandler.calls
    assert data["facts"] == []
    assert not any(item.get("amount", {}).get("normalized_base") == "0" for item in data["facts"])


def test_mapper_redirect_is_not_followed_and_dummy_key_is_not_forwarded(tmp_path: Path) -> None:
    source = simple_xlsx(tmp_path / "mystery.xlsx", ["Asset ID", "Mystery EUR"], [["A-1", 123]])
    RedirectHandler.calls = 0
    RedirectHandler.authorizations = []
    RedirectTargetHandler.calls = 0
    RedirectTargetHandler.authorizations = []
    with local_server(RedirectTargetHandler) as target:
        RedirectHandler.target = target
        with local_server(RedirectHandler) as endpoint:
            config = mapper_config(tmp_path / "mapper.toml", endpoint)
            environment = {key: value for key, value in os.environ.items() if key != "OPENAI_API_KEY"}
            environment["OPENAI_API_KEY"] = "dummy-must-not-be-sent"
            completed, records = run_cli(source, tmp_path / "out", config=config, env=environment)
    assert completed.returncode == 0
    assert RedirectHandler.calls >= 1
    assert RedirectHandler.authorizations == [None] * RedirectHandler.calls
    assert RedirectTargetHandler.calls == 0 and RedirectTargetHandler.authorizations == []
    assert any("mapper_unavailable" in str(item) for item in records[0]["diagnostics"])
    assert "dummy-must-not-be-sent" not in completed.stdout + completed.stderr


def test_mapper_outage_is_additive_and_cannot_delete_local_facts(tmp_path: Path) -> None:
    source = simple_xlsx(
        tmp_path / "partly-mapped.xlsx",
        ["Asset ID", "Opening cost EUR 2025", "Mystery EUR"],
        [["A-1", 100, 999]],
    )
    disabled_completed, disabled_records = run_cli(source, tmp_path / "disabled")
    assert disabled_completed.returncode == 0
    disabled = evidence(disabled_records[0])
    assert any(item["field"] == "ahk_open" for item in disabled["facts"])

    UnavailableHandler.calls = 0
    with local_server(UnavailableHandler) as endpoint:
        config = mapper_config(tmp_path / "mapper.toml", endpoint)
        enabled_completed, enabled_records = run_cli(source, tmp_path / "unavailable", config=config)
    assert enabled_completed.returncode == 0
    enabled = evidence(enabled_records[0])
    assert UnavailableHandler.calls == 3, "transient retries must remain bounded"
    assert enabled["facts"] == disabled["facts"]
    assert any("mapper_unavailable" in str(item) for item in enabled_records[0]["diagnostics"])


def test_template_duplicate_asset_ids_across_entities_never_overwrite(tmp_path: Path) -> None:
    if not TEMPLATE.exists():
        pytest.skip("private layout-20 template unavailable")
    workbook = Workbook()
    first = workbook.active
    first.title = "EntityA"
    first.append(["Asset ID", "Opening cost EUR 2025"])
    first.append(["A-1", 100])
    second = workbook.create_sheet("EntityB")
    second.append(["Asset ID", "Opening cost EUR 2025"])
    second.append(["A-1", 200])
    second.append(["A-2", 300])
    source = tmp_path / "duplicate-entities.xlsx"
    workbook.save(source)

    completed, records = run_cli(source, tmp_path / "out", template=TEMPLATE)
    assert completed.returncode == 0
    data = evidence(records[0])
    opening = [item for item in data["facts"] if item["field"] == "ahk_open"]
    assert len(opening) == 3
    by_fact = {item["record_id"]: item for item in data["output_bindings"]}
    duplicate = [item for item in opening if item["asset_identity"]["asset_id"] == "A-1"]
    unrelated = next(item for item in opening if item["asset_identity"]["asset_id"] == "A-2")
    assert {item["pointer"]["sheet"] for item in duplicate} == {"EntityA", "EntityB"}
    duplicate_bindings = [by_fact[item["id"]] for item in duplicate]
    if all(item["status"] == "bound" for item in duplicate_bindings):
        assert len({item["destination"]["cell"] for item in duplicate_bindings}) == 2
    else:
        assert {item["status"] for item in duplicate_bindings} == {"ambiguous_destination_identity"}
    assert by_fact[unrelated["id"]]["status"] == "bound"


def test_duplicate_same_row_and_field_is_not_accepted_twice(tmp_path: Path) -> None:
    source = simple_xlsx(
        tmp_path / "duplicates.xlsx",
        ["Asset ID", "Additions EUR", "Additions EUR"],
        [["A-1", 100, 100]],
    )
    completed, records = run_cli(source, tmp_path / "out")
    assert completed.returncode == 0
    data = evidence(records[0])
    additions = [item for item in data["facts"] if item["field"] == "zugang"]
    assert additions == [], "two indistinguishable physical occurrences for one business slot require abstention"
    assert data["abstentions"]


def test_raw_signed_amount_is_preserved_without_blanket_absolute_value(tmp_path: Path) -> None:
    source = simple_xlsx(
        tmp_path / "signed.xlsx",
        ["Asset ID", "Disposals EUR 2025"],
        [["A-1", "184.782,83-"]],
    )
    completed, records = run_cli(source, tmp_path / "out")
    assert completed.returncode == 0
    data = evidence(records[0])
    facts = [item for item in data["facts"] if item["field"] == "abgang"]
    assert len(facts) == 1, data["abstentions"]
    amount = facts[0]["amount"]
    assert amount["raw"] == "184.782,83-"
    assert Decimal(amount["source_signed"]) == Decimal("-184782.83")
    assert Decimal(amount["normalized_base"]) == Decimal("-184782.83")


@pytest.mark.parametrize("header", ["Disposals (-) EUR", "Disposals -J EUR"])
def test_header_polarity_or_ocr_noise_never_changes_positive_source_sign(tmp_path: Path, header: str) -> None:
    source = simple_xlsx(tmp_path / "positive-source.xlsx", ["Asset ID", f"{header} 2025"], [["A-1", 25]])
    completed, records = run_cli(source, tmp_path / "out")
    assert completed.returncode == 0
    data = evidence(records[0])
    facts = [item for item in data["facts"] if item["field"] == "abgang"]
    assert len(facts) == 1, data["abstentions"]
    amount = facts[0]["amount"]
    assert amount["raw"] == "25"
    assert Decimal(amount["source_signed"]) == Decimal("25")
    assert Decimal(amount["normalized_base"]) == Decimal("25")


def make_two_table_pdf(path: Path) -> Path:
    page = canvas.Canvas(str(path), pagesize=(612, 792))
    page.setFont("Helvetica", 10)
    page.drawString(40, 750, "Asset schedule 2025 — EUR")
    page.drawString(40, 730, "Asset ID")
    page.drawString(220, 730, "Opening cost EUR")
    page.drawString(40, 710, "EUR-1")
    page.drawString(220, 710, "2")
    page.drawString(40, 610, "Asset schedule 2025 — TEUR")
    page.drawString(40, 590, "Asset ID")
    page.drawString(220, 590, "Opening cost TEUR")
    page.drawString(40, 570, "TEUR-1")
    page.drawString(220, 570, "3")
    page.save()
    return path


def test_pdf_multiple_tables_never_apply_page_wide_teur_scale(tmp_path: Path) -> None:
    source = make_two_table_pdf(tmp_path / "two-tables.pdf")
    completed, records = run_cli(source, tmp_path / "out")
    assert completed.returncode == 0
    data = evidence(records[0])
    by_asset: dict[str, list[dict[str, object]]] = {}
    for fact in data["facts"]:
        identity = fact["asset_identity"].get("asset_id") or fact["asset_identity"].get("asset_or_class")
        if identity:
            by_asset.setdefault(str(identity), []).append(fact)
    for fact in by_asset.get("EUR-1", []):
        assert Decimal(fact["amount"]["normalized_base"]) == Decimal("2")
        assert fact["amount"]["unit"] == "EUR"
    for fact in by_asset.get("TEUR-1", []):
        assert Decimal(fact["amount"]["normalized_base"]) == Decimal("3000")
        assert fact["amount"]["unit"] == "TEUR"
    assert not any(Decimal(fact["amount"]["normalized_base"]) == Decimal("2000") for fact in by_asset.get("EUR-1", []))
    assert not any(Decimal(fact["amount"]["normalized_base"]) == Decimal("3") for fact in by_asset.get("TEUR-1", []))

